# -*- coding: utf-8 -*-
"""
Generador del Excel maestro PK-16 (criticidad y dimensionamiento de mantenimiento).

PowerPoint final = resultado oficial. Excel = datos + cálculos + criterios + metodología + trazabilidad.
Todas las hojas de cálculo leen sus parámetros de 12_SUPUESTOS mediante nombres definidos (P_*).
"""
from __future__ import annotations

import datetime as dt

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import pk16_datos as D
import pk16_oficial as O
from pk16_xlsx_estilo import (BORDER, hl, CENTER, F_BASE, F_BIG, F_BOLD, F_HEAD, F_INPUT, F_LINK, F_NOTE, F_SECTION,
                              F_SUB, FILL, NAVY, WRAP, class_format, define, header_row, note, put, section,
                              status_format, title, widths)

SHEETS = [
    ("00_INICIO", "Navegación, principio de trabajo y leyenda"),
    ("01_RESUMEN_EJECUTIVO", "Cifras oficiales del PPT reproducidas por el modelo y estado de validación"),
    ("02_MAESTRO_CRITICIDAD", "Base única de activos: todos los registros, estado, clasificación y fuente"),
    ("03_SCORE_CRITICIDAD", "Score por activo: 5 criterios, clase, reglas y motivo (fórmulas)"),
    ("04_FTE_CRITICIDAD", "FTE base por clase (oficial), sensibilidad de ratios y origen de A/B/C/D"),
    ("05_FTE_POR_GRUPOS", "Reparto presentado, reproducción de la fuente y reconstrucción por grupos"),
    ("06_CAPACIDAD_PERSONAL", "Waterfall de necesidad, plantilla, capacidad real, GAP y conciliación"),
    ("07_SABADOS_FORMACION", "Formación y descanso compensatorio (libranzas)"),
    ("08_PM_AVISOS", "Preventivo, avisos, correctivo, TPM: unidades y conversión a FTE"),
    ("09_MEB_ECO", "Nuevos productos: oficial 75 críticos, cruce SAP y riesgo de doble conteo"),
    ("10_AVERIAS", "Averías y paro 2025→actual: resumen, cruce y discrepancias"),
    ("11_TRAZABILIDAD", "KPI → valor oficial → diapositiva → cálculo → fuente → estado"),
    ("12_SUPUESTOS", "Parámetros editables (única hoja de entrada)"),
    ("13_METODOLOGIA", "Método, reglas, jerarquía de ubicaciones y referencias"),
    ("14_FUENTES", "Inventario de fuentes, confianza y uso"),
    ("15_PENDIENTES_VALIDACION", "Incertidumbres abiertas, impacto y responsable"),
    ("16_CONTROLES", "Controles automáticos OK / REVISAR / ERROR"),
    ("17_DIFERENCIAS", "Diferencias detectadas entre resultado presentado y fuente"),
    ("18_REG_AVERIAS", "Registros de avería SAP 2025→actual con su fila de origen y activo asignado"),
    ("19_CRUCE_MONTAJE", "Fichero de montaje (264 filas) cruzado con SAP"),
    ("20_CRUCE_CUELLOS", "Cuellos de botella (129) cruzados con SAP"),
]


class Ctx:
    """Estado compartido del generador."""

    def __init__(self, M, S):
        self.M, self.S = M, S
        self.inv = M["inv"]
        self.wb = Workbook()
        self.ws = {}
        first = True
        for name, _ in SHEETS:
            if first:
                ws = self.wb.active
                ws.title = name
                first = False
            else:
                ws = self.wb.create_sheet(name)
            self.ws[name] = ws
        self.n_score = len(S)
        self.score_first = 6
        self.score_last = 5 + len(S)
        self.kpi = {}  # referencias de celdas clave para resumen/trazabilidad/controles


# ======================================================================================
# 12_SUPUESTOS
# ======================================================================================
def hoja_supuestos(c: Ctx):
    ws = c.ws["12_SUPUESTOS"]
    title(ws, "12 · SUPUESTOS Y PARÁMETROS EDITABLES",
          "Única hoja de entrada. El resto del libro toma estos valores por nombre (P_…). Cambie un valor y todo se recalcula.")
    widths(ws, {"A": 16, "B": 46, "C": 14, "D": 16, "E": 16, "F": 52, "G": 11, "H": 26, "I": 60})
    put(ws, "A5", "Leyenda:", F_BOLD, border=False)
    for ref, f, t in (("B5", "oficial", "Valor oficial presentado (PPT)"), ("C5", "dato", "Dato de fuente"),
                      ("D5", "supuesto", "Supuesto / regla"), ("E5", "pendiente", "Pendiente / no disponible")):
        put(ws, ref, t, F_BASE, f, align=CENTER)
    header_row(ws, 7, 1, ["Nombre (P_…)", "Parámetro", "Valor", "Unidad", "Tipo", "Fuente (archivo · hoja · celda)",
                          "Diapositiva", "Estado validación", "Observación"])
    ws.freeze_panes = "C8"
    r = [8]
    tipo_fill = {"Oficial PPT": "oficial", "Dato fuente": "dato", "Supuesto": "supuesto", "Regla modelo": "supuesto",
                 "Pendiente": "pendiente", "Selector": "supuesto"}

    def sec(txt):
        section(ws, r[0], txt, 9)
        r[0] += 1

    def p(name, label, val, unit, tipo, fuente, diapo="", estado="", obs="", fmt=None):
        row = r[0]
        put(ws, f"A{row}", "P_" + name, F_BOLD)
        put(ws, f"B{row}", label, F_BASE, align=WRAP)
        cell = put(ws, f"C{row}", val, F_INPUT, tipo_fill[tipo], fmt=fmt, align=Alignment(horizontal="right"))
        put(ws, f"D{row}", unit)
        put(ws, f"E{row}", tipo, fill_=tipo_fill[tipo])
        put(ws, f"F{row}", fuente, align=WRAP)
        put(ws, f"G{row}", diapo, align=CENTER)
        put(ws, f"H{row}", estado)
        put(ws, f"I{row}", obs, align=WRAP)
        define(c.wb, "P_" + name, "12_SUPUESTOS", f"$C${row}")
        c.kpi["P_" + name] = f"'12_SUPUESTOS'!$C${row}"
        r[0] += 1
        return cell

    ppt = O.PPT
    sec("1. Parque y distribución oficial A/B/C/D")
    p("PARQUE_OF", "Parque oficial (activos)", 2043, "activos", "Oficial PPT", f"{ppt}", "4 · 8", "OFICIAL – suma de familias (ver 04)",
      "Suma de familias del contexto global (430+92+264+58+290+74+355+386+94). No procede de un inventario fila a fila.", "int")
    p("N_A", "Activos clase A", 418, "activos", "Oficial PPT", f"{ppt} (tabla diap. 4)", "4", "RECONSTRUIDO con observaciones (ver 04)",
      "= filas con criticidad A en MAESTRO_FABRICA_FINAL 01_Maestro_Fabrica (129 CB + 55 A2 + 185 montaje + 43 MEB montaje + 6 MEB oficial).", "int")
    p("N_B", "Activos clase B", 275, "activos", "Oficial PPT", f"{ppt} (tabla diap. 4)", "4", "TRAZABILIDAD INCOMPLETA", "La fuente da 119 (B - Alta carga). Ver 17_DIFERENCIAS.", "int")
    p("N_C", "Activos clase C", 1200, "activos", "Oficial PPT", f"{ppt} (tabla diap. 4)", "4", "TRAZABILIDAD INCOMPLETA", "Valor redondo declarado. La fuente da 189 (153 C + 36 C/Pendiente).", "int")
    p("N_D", "Activos clase D", 150, "activos", "Oficial PPT", f"{ppt} (tabla diap. 4)", "4", "TRAZABILIDAD INCOMPLETA", "La fuente da 920 D + 397 no clasificados.", "int")

    sec("2. Ratios de dimensionamiento por criticidad (criterio interno de cartera, no norma universal)")
    for k, v, mn in (("A", 12, 8), ("B", 20, 12), ("C", 35, 20), ("D", 60, 35)):
        p(f"RATIO_{k}", f"Ratio oficial clase {k} (activos por mantenedor)", v, "activos/FTE", "Oficial PPT", f"{ppt} tabla diap. 4", "4",
          "OFICIAL", f"Extremo superior del rango {mn}–{v} de la diapositiva 3 (menor dotación).", "int")
    for k, v, mx in (("A", 8, 12), ("B", 12, 20), ("C", 20, 35), ("D", 35, 60)):
        p(f"RMIN_{k}", f"Ratio mínimo del rango clase {k} (sensibilidad)", v, "activos/FTE", "Oficial PPT", f"{ppt} diap. 3 'RATIOS CRITICIDAD'", "3",
          "OFICIAL (rango)", f"Rango {v}–{mx}. Solo se usa en el análisis de sensibilidad.", "int")

    sec("3. Puntos equivalentes por clase (visualizar y comparar cartera; NO son horas ni sustituyen al FTE)")
    for k, v in (("A", 4), ("B", 2.5), ("C", 1.5), ("D", 1)):
        p(f"PT_{k}", f"Puntos equivalentes clase {k}", v, "puntos/activo", "Dato fuente",
          "MAESTRO_CRITICIDAD_MAQUINAS.xlsx · Criterios · C4:C7", "", "DATO FUENTE", "", "0.0")

    sec("4. Método de score (diapositiva 4) – máximos por criterio y umbrales de clase")
    for k, lbl, v in (("MAX_C1", "Máximo C1 Impacto en producción", 4), ("MAX_C2", "Máximo C2 Seguridad / calidad / cliente / trazabilidad", 2),
                      ("MAX_C3", "Máximo C3 Unicidad / redundancia / cuello de botella", 2), ("MAX_C4", "Máximo C4 Averías, recurrencia y paro", 1.5),
                      ("MAX_C5", "Máximo C5 Obsolescencia, repuestos y soporte", 0.5)):
        p(k, lbl, v, "puntos", "Oficial PPT", f"{ppt} diap. 4 'DE DÓNDE SALE LA CLASE'", "4", "OFICIAL", "", "0.0")
    p("UMB_A", "Umbral clase A (score ≥)", 8, "puntos", "Oficial PPT", f"{ppt} diap. 4 'Rangos'", "4", "OFICIAL", "A ≥ 8", "0.0")
    p("UMB_B", "Umbral clase B (score ≥)", 6, "puntos", "Oficial PPT", f"{ppt} diap. 4 'Rangos'", "4", "OFICIAL", "B 6–7,9", "0.0")
    p("UMB_C", "Umbral clase C (score ≥)", 3.5, "puntos", "Oficial PPT", f"{ppt} diap. 4 'Rangos'", "4", "OFICIAL", "C 3,5–5,9 · D < 3,5", "0.0")

    sec("5. Plantilla nominal y horas productivas")
    p("PLANTILLA", "Plantilla nominal actual de mantenimiento", 80, "personas", "Oficial PPT", f"{ppt} diap. 5 y 7 '~80 personas' · DIMENSIONAMIENTO 01_Supuestos_Fuentes B6", "5 · 7",
      "PENDIENTE DE VALIDACIÓN", "Dato facilitado por el usuario ('~80'). No hay fichero de organización/turnos en el repositorio.", "int")
    p("FTE_PERSONA", "FTE contractual por persona nominal", 1.0, "FTE/persona", "Supuesto", "Supuesto del modelo", "", "SUPUESTO",
      "Una persona nominal = 1 FTE contractual antes de pérdidas. Ajustar si hay jornadas parciales.", "0.00")
    p("HORAS_PROD", "Horas productivas anuales por FTE", 1450, "h/FTE/año", "Supuesto", "DIMENSIONAMIENTO_MANTENIMIENTO_DATOS_FINAL.xlsx · 01_Supuestos_Fuentes · B8", "", "SUPUESTO – validar con RRHH",
      "Capacidad efectiva estimada, no jornada teórica.", "int")

    sec("6. Formación")
    p("FORM_H", "Horas de formación por persona y año", 55, "h/persona/año", "Supuesto", "DIMENSIONAMIENTO … 01_Supuestos_Fuentes · B7", "5", "SUPUESTO – validar plan de formación", "", "int")

    sec("7. Fines de semana, festivos y descanso compensatorio (libranzas)")
    p("PERS_SAB", "Personas por jornada especial (sábado)", 10, "personas/jornada", "Supuesto", "DIMENSIONAMIENTO … 01_Supuestos_Fuentes · B9", "5", "SUPUESTO – validar calendario", "", "int")
    p("SABADOS", "Jornadas especiales al año (sábados)", 45, "jornadas/año", "Supuesto", "DIMENSIONAMIENTO … 01_Supuestos_Fuentes · B10", "5", "SUPUESTO – validar calendario", "", "int")
    p("FESTIVOS", "Festivos trabajados al año", 0, "jornadas/año", "Pendiente", "DATO NO DISPONIBLE", "", "PENDIENTE DE VALIDACIÓN",
      "No hay calendario de festivos con cobertura. Con 0 no altera el resultado oficial.", "int")
    p("JORNADAS", "Jornadas efectivas por FTE y año", 220, "días/FTE/año", "Supuesto", "DIMENSIONAMIENTO … 01_Supuestos_Fuentes · B11", "", "SUPUESTO",
      "Base para convertir jornadas de libranza a FTE.", "int")
    p("H_COBERTURA", "Horas anuales reales de cobertura fin de semana (método horario)", 0, "h/año", "Pendiente", "DATO NO DISPONIBLE", "", "PENDIENTE DE VALIDACIÓN",
      "Si se informa (>0), 07 calcula el método horario alternativo.", "int")

    sec("8. Carga planificada: preventivo, avisos y TPM")
    p("PUNTOS_PM", "Puntos de preventivo al año", 50000, "puntos/año", "Dato fuente", f"{O.PPT_CTX} diap. 1 · MAESTRO_FABRICA 07_Contexto_Global fila 14", "", "DATO GLOBAL – sin detalle por plan",
      "La diapositiva 8 del PPT final muestra '15.000 puntos de preventivo yoY' (ver 17).", "int")
    p("AVISOS", "Avisos de intervención en máquina al año", 10000, "avisos/año", "Dato fuente", f"{O.PPT_CTX} diap. 1 (2024)", "", "DATO GLOBAL", "", "int")
    p("PRODUCT", "Productividad PM + avisos por FTE", 6300, "unidades/FTE/año", "Supuesto", "DIMENSIONAMIENTO … 01_Supuestos_Fuentes · B13", "5", "HIPÓTESIS – sin horas reales",
      "Equivale a ~13,8 min por unidad (ver 08). Puntos y avisos no consumen el mismo tiempo.", "int")
    p("PLANES_PM", "Planes de mantenimiento preventivo", 9000, "planes", "Dato fuente", f"{O.PPT_CTX} diap. 1", "", "DATO GLOBAL", "Informativo.", "int")
    p("TPM", "Intervenciones TPM al año", 3000, "intervenciones/año", "Dato fuente", f"{O.PPT_CTX} diap. 1 (2024)", "", "DATO GLOBAL",
      "Informado como actividad operativa. No se suma como FTE (incluido en la actividad del grupo).", "int")
    p("TPM_SUMAR", "¿Sumar TPM como FTE adicional? (0 = no, criterio oficial)", 0, "0/1", "Selector", "Criterio oficial: 'TPM excluido' (DIMENSIONAMIENTO 00_Resumen H11)", "", "OFICIAL",
      "Solo para análisis: 1 suma TPM × P_H_TPM ÷ P_HORAS_PROD.", "0")
    p("H_TPM", "Horas por intervención TPM (solo si P_TPM_SUMAR = 1)", 1, "h/intervención", "Supuesto", "MAESTRO_INTEGRAL 07_Preventivo_TPM_FTE (escenario medio)", "", "ESCENARIO", "", "0.0")

    sec("9. Cobertura organizativa (vacaciones, absentismo, turnos, especialidades, simultaneidad)")
    p("COB_MIN", "Factor de cobertura mínimo", 0.15, "%", "Oficial PPT", f"{ppt} diap. 5 'Cobertura 15–25%'", "5", "OFICIAL – sujeto a validación organizativa", "", "pct")
    p("COB_MAX", "Factor de cobertura máximo", 0.25, "%", "Oficial PPT", f"{ppt} diap. 5 'Cobertura 15–25%'", "5", "OFICIAL – sujeto a validación organizativa", "", "pct")
    cell = p("COB_BASE", "Base sobre la que se aplica la cobertura (1 = Subtotal · 2 = Base + PM · 3 = Solo base)", 1, "selector", "Selector",
             f"1: método diap. 2 ('Subtotal FTE × factor') → 115–125 (diap. 5 y 7) · 2: DIMENSIONAMIENTO 04_Calculo_FTE B10 → 114–124 · 3: waterfall diap. 5 (+13/+21) → 113–121",
             "2 · 5 · 7", "PENDIENTE DE VALIDACIÓN (ver 06 y 17)", "No reincluye formación ni libranzas como conceptos: solo cambia la base del %.", "0")
    dv = DataValidation(type="list", formula1='"1,2,3"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(cell)

    sec("10. Criterios de score: umbrales de averías (C4), obsolescencia (C5) y reglas")
    p("PER_INI", "Inicio del periodo principal de averías", dt.date(2025, 1, 1), "fecha", "Oficial PPT", f"{ppt} diap. 4 'Averías y paro 2025→actual'", "4", "OFICIAL", "", "dd/mm/yyyy")
    fin = c.M["av"]["FECHA"].max()
    p("PER_FIN", "Último registro de avería disponible", fin.date() if pd.notna(fin) else None, "fecha", "Dato fuente", "SAP.EXPORT.AVERIAS.LISTADO.xlsx · AVERIAS · col. B (máximo)", "", "DATO FUENTE", "", "dd/mm/yyyy")
    p("C4_H15", "C4 = 1,5 si horas de paro ≥", D.PARAM["C4_H15"], "h", "Regla modelo", "Regla del modelo (periodo ~19 meses)", "", "PENDIENTE DE VALIDACIÓN", "o nº averías ≥ P_C4_R15, o parada máxima ≥ P_C4_MAX15", "0")
    p("C4_R15", "C4 = 1,5 si nº de averías ≥", D.PARAM["C4_R15"], "avisos", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "", "0")
    p("C4_MAX15", "C4 = 1,5 si parada máxima ≥", D.PARAM["C4_MAX15"], "min", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "", "0")
    p("C4_H10", "C4 = 1 si horas de paro ≥", D.PARAM["C4_H10"], "h", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "o nº averías ≥ P_C4_R10", "0")
    p("C4_R10", "C4 = 1 si nº de averías ≥", D.PARAM["C4_R10"], "avisos", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "", "0")
    p("C4_H05", "C4 = 0,5 si horas de paro ≥", D.PARAM["C4_H05"], "h", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "o nº averías ≥ P_C4_R05", "0")
    p("C4_R05", "C4 = 0,5 si nº de averías ≥", D.PARAM["C4_R05"], "avisos", "Regla modelo", "Regla del modelo", "", "PENDIENTE DE VALIDACIÓN", "", "0")
    p("C5_Y05", "C5 = 0,5 si año de puesta en servicio ≤", D.PARAM["C5_Y05"], "año", "Regla modelo", "SAP Fe.puesta serv. (1900 = desconocido)", "", "PENDIENTE DE VALIDACIÓN",
      "Proxy de obsolescencia: no hay datos de PLC/CNC, repuestos ni soporte.", "0")
    p("C5_Y025", "C5 = 0,25 si año de puesta en servicio ≤", D.PARAM["C5_Y025"], "año", "Regla modelo", "SAP Fe.puesta serv.", "", "PENDIENTE DE VALIDACIÓN", "", "0")
    p("C5_UNK", "C5 cuando el año no está disponible", D.PARAM["C5_UNK"], "puntos", "Regla modelo", "Regla del modelo (valor medio, no penaliza ni premia)", "", "PENDIENTE DE VALIDACIÓN", "", "0.00")
    p("C1_SAPA", "C1 mínimo si el Indicador ABC de SAP es 'A'", D.PARAM["C1_SAPA"], "puntos", "Regla modelo", "LISTADO_EQUIPOS_SAP · Indicador ABC", "", "PENDIENTE DE VALIDACIÓN", "", "0")

    sec("11. Tolerancias de los controles")
    p("TOL_RED", "Tolerancia por redondeo de presentación", 0.5, "FTE / activos", "Regla modelo", "Regla del modelo", "", "", "Diferencias ≤ 0,5 se consideran redondeo.", "0.00")
    p("TOL_EXACT", "Tolerancia de cálculo exacto", 0.01, "FTE", "Regla modelo", "Regla del modelo", "", "", "", "0.00")

    r[0] += 1
    sec("12. Tipologías de criticidad: valores base de C1, C2 y C3 (editables)")
    hr = r[0]
    header_row(ws, hr, 1, ["Código", "Tipología", "C1 base", "C2 base", "C3 base", "Datos suficientes (S/N)", "Activos en parque", "", "Criterio aplicado"])
    r[0] += 1
    t0 = r[0]
    cnt = c.S["TIPOLOGIA"].value_counts()
    for code, name, c1, c2, c3, suf, desc in D.TIPOLOGIAS:
        row = r[0]
        put(ws, f"A{row}", code, F_BOLD)
        put(ws, f"B{row}", name, align=WRAP)
        for col, v in (("C", c1), ("D", c2), ("E", c3)):
            put(ws, f"{col}{row}", v, F_INPUT, "supuesto", fmt="0.0")
        put(ws, f"F{row}", suf, F_INPUT, "supuesto", align=CENTER)
        put(ws, f"G{row}", f"=COUNTIF(S_TIPO,A{row})", fmt="int")
        put(ws, f"H{row}", "")
        put(ws, f"I{row}", desc, align=WRAP)
        ws.row_dimensions[row].height = 30
        r[0] += 1
    define(c.wb, "T_TIPO", "12_SUPUESTOS", f"$A${t0}:$F${r[0] - 1}")
    note(ws, f"A{r[0] + 1}", "Reglas que prevalecen sobre la tabla (03): cuello de botella confirmado → C1 = 4, C3 = 2 y clase A · montaje principal → C1 = 4, C2 = 2, C3 = 2 y clase A · "
         "activo crítico MEB ECO confirmado → C1 = 4, C3 = 2 y clase A · indicio de seguridad/calidad en la descripción → C2 = 2 · equipo equivalente en el grupo → C3 = 0,5.")
    note(ws, f"A{r[0] + 2}", "Datos suficientes = N, o activo con grupo propuesto → la clase nunca puede ser D (queda 'C provisional').")


# ======================================================================================
# 03_SCORE_CRITICIDAD
# ======================================================================================
SCORE_COLS = [
    # (col, cabecera, ancho, campo S | None si fórmula, formato)
    ("A", "ID activo", 22, "ID", None), ("B", "Descripción", 40, "DESC", None), ("C", "Grupo mantenimiento", 24, "GRUPO_MANT", None),
    ("D", "Producto", 16, "PRODUCTO", None), ("E", "Programa", 12, "PROGRAMA", None), ("F", "Proceso", 16, "PROCESO", None),
    ("G", "Línea / sección", 26, "LINEA", None), ("H", "Tecnología", 20, "TECNOLOGIA", None), ("I", "Tipología", 9, "TIPOLOGIA", None),
    ("J", "Nombre tipología", 30, "TIPOLOGIA_NOMBRE", None), ("K", "Cuello botella", 8, "CB", None), ("L", "Montaje principal", 9, "MONTAJE_PPAL", None),
    ("M", "MEB crítico", 8, "MEB_CRIT", None), ("N", "SAP ABC", 7, "SAP_ABC", None), ("O", "Indicio seg./calidad", 9, "IND_SEGCAL", None),
    ("P", "Equipos equivalentes en grupo", 10, "EQUIV_N", "int"), ("Q", "Averías 2025→ (nº)", 10, "AV_N", "int"),
    ("R", "Horas paro 2025→ (h)", 10, "AV_H", "h"), ("S", "Parada máx. (min)", 10, "AV_MAX", "int"), ("T", "Año puesta servicio", 9, "ANIO", "0"),
    ("U", "Estado registro", 26, "ESTADO", None),
    ("V", "C1 Producción (0–4)", 10, None, "0.00"), ("W", "C2 Seg./calidad (0–2)", 10, None, "0.00"), ("X", "C3 Unicidad / CB (0–2)", 10, None, "0.00"),
    ("Y", "C4 Averías (0–1,5)", 10, None, "0.00"), ("Z", "C5 Obsolescencia (0–0,5)", 11, None, "0.00"), ("AA", "SCORE (0–10)", 9, None, "0.00"),
    ("AB", "Clase por score", 8, None, None), ("AC", "Regla forzada", 30, "REGLA_FORZADA", None), ("AD", "Datos suficientes", 9, None, None),
    ("AE", "CLASE FINAL", 8, None, None), ("AF", "Provisional", 9, None, None), ("AG", "Clase (etiqueta)", 13, None, None),
    ("AH", "Motivo de criticidad", 90, None, None), ("AI", "Puntos equivalentes", 9, None, "0.0"), ("AJ", "Ratio aplicado", 8, None, "0"),
    ("AK", "FTE equivalente", 9, None, "0.000"), ("AL", "Estado validación", 26, "VALIDACION", None), ("AM", "Fuente de evidencias", 50, "EVID", None),
]


def _evidencias(r):
    ev = [f"Inventario: {r['FUENTE']} fila {r['FILA']}"]
    if r["CB"] == "S":
        ev.append("CB: CUELLOS_DE_BOTELLA.xlsx (" + r["CB_DETALLE"][:60] + ")")
    if r["MONTAJE_PPAL"] == "S" and _s(r.get("MONT_FILA")):
        ev.append("Montaje: MAESTRO_FABRICA_FINAL 04_Montaje_Completo fila " + _s(r.get("MONT_FILA")))
    if r["AV_N"]:
        ev.append("Averías: " + _s(r["AV_FUENTE"]))
    if r["MEB_CRIT"] == "S":
        ev.append("MEB: MAESTRO_FABRICA_FINAL 02_MEB_ECO_Oficial fila 7")
    return " · ".join(ev)


def _s(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    return str(x)


def hoja_score(c: Ctx):
    ws = c.ws["03_SCORE_CRITICIDAD"]
    title(ws, "03 · SCORE DE CRITICIDAD POR ACTIVO (5 criterios · máx. 10 puntos)",
          "Una fila por activo del parque principal. Columnas azules = datos y evidencias; columnas blancas = fórmulas que leen 12_SUPUESTOS.")
    put(ws, "C3", "SCORE = C1 + C2 + C3 + C4 + C5 · Clase: A ≥ P_UMB_A · B ≥ P_UMB_B · C ≥ P_UMB_C · D resto (solo con datos suficientes)", F_BOLD, border=False)
    for col, h, w, f, fmt in SCORE_COLS:
        ws.column_dimensions[col].width = w
    header_row(ws, 5, 1, [h for _, h, _, _, _ in SCORE_COLS], height=45)
    S = c.S.sort_values(["GRUPO_MANT", "LINEA", "ID"]).reset_index(drop=True)
    S["EVID"] = S.apply(_evidencias, axis=1)
    c.S_sorted = S
    for i, rec in S.iterrows():
        r = 6 + i
        for col, h, w, f, fmt in SCORE_COLS:
            if f is None:
                continue
            v = rec[f]
            if f == "ANIO":
                v = int(v) if (v is not None and not pd.isna(v)) else None
            elif isinstance(v, float) and pd.isna(v):
                v = None
            cell = ws[f"{col}{r}"]
            cell.value = v
            cell.font = F_BASE
            cell.fill = FILL["dato"]
            if fmt:
                cell.number_format = fmt if fmt not in ("int", "h") else ("#,##0" if fmt == "int" else "#,##0.0")
        frz = f'OR(K{r}="S",L{r}="S",M{r}="S")'
        ws[f"V{r}"] = f'=IF({frz},4,MAX(VLOOKUP($I{r},T_TIPO,3,FALSE),IF(N{r}="A",P_C1_SAPA,0)))'
        ws[f"W{r}"] = f'=IF(OR(L{r}="S",O{r}="S"),2,VLOOKUP($I{r},T_TIPO,4,FALSE))'
        ws[f"X{r}"] = f'=IF({frz},2,IF(AND(P{r}>=1,VLOOKUP($I{r},T_TIPO,5,FALSE)>0.5),0.5,VLOOKUP($I{r},T_TIPO,5,FALSE)))'
        ws[f"Y{r}"] = (f'=IF(OR(R{r}>=P_C4_H15,Q{r}>=P_C4_R15,S{r}>=P_C4_MAX15),1.5,'
                       f'IF(OR(R{r}>=P_C4_H10,Q{r}>=P_C4_R10),1,IF(OR(R{r}>=P_C4_H05,Q{r}>=P_C4_R05),0.5,0)))')
        ws[f"Z{r}"] = f'=IF(T{r}="",P_C5_UNK,IF(T{r}<=P_C5_Y05,0.5,IF(T{r}<=P_C5_Y025,0.25,0)))'
        ws[f"AA{r}"] = f"=SUM(V{r}:Z{r})"
        ws[f"AB{r}"] = f'=IF(AA{r}>=P_UMB_A,"A",IF(AA{r}>=P_UMB_B,"B",IF(AA{r}>=P_UMB_C,"C","D")))'
        ws[f"AD{r}"] = f'=IF(ISNUMBER(SEARCH("propuesta",U{r})),"N",VLOOKUP($I{r},T_TIPO,6,FALSE))'
        ws[f"AE{r}"] = f'=IF(AC{r}<>"",LEFT(AC{r},1),IF(AND(AB{r}="D",AD{r}<>"S"),"C",AB{r}))'
        ws[f"AF{r}"] = f'=IF(AC{r}<>"","N",IF(AD{r}<>"S","S","N"))'
        ws[f"AG{r}"] = f'=AE{r}&IF(AF{r}="S"," provisional","")'
        ws[f"AH{r}"] = (
            f'=AG{r}&": "&IF(K{r}="S","cuello de botella confirmado; ","")'
            f'&IF(L{r}="S","estación de montaje en línea serie (regla montaje principal); ","")'
            f'&IF(M{r}="S","activo crítico MEB ECO confirmado; ","")'
            f'&IF({frz},"",IF(V{r}>=4,"puede parar línea/proceso; ",IF(V{r}>=3,"pérdida importante de proceso; ",IF(V{r}>=2,"parada de estación con alternativa; ",IF(V{r}>=1,"impacto productivo menor; ","sin impacto productivo directo; ")))))'
            f'&IF(W{r}>=2,"riesgo de seguridad, calidad o trazabilidad; ","")'
            f'&IF(X{r}>=2,"sin redundancia; ",IF(AND(X{r}<=0.5,P{r}>=1),"existe equipo equivalente en el grupo; ",""))'
            f'&IF(Y{r}>=1,"averías recurrentes ("&Q{r}&" avisos, "&TEXT(R{r},"0")&" h de paro desde 2025); ",IF(Y{r}>0,"averías moderadas ("&Q{r}&" avisos); ",""))'
            f'&IF(Z{r}>=0.5,"obsolescencia por antigüedad; ",IF(T{r}="","año no disponible; ",""))'
            f'&IF(AF{r}="S","PENDIENTE: información insuficiente o grupo propuesto – validar con Producción/Mantenimiento","")')
        ws[f"AI{r}"] = f'=IF(AE{r}="A",P_PT_A,IF(AE{r}="B",P_PT_B,IF(AE{r}="C",P_PT_C,P_PT_D)))'
        ws[f"AJ{r}"] = f'=IF(AE{r}="A",P_RATIO_A,IF(AE{r}="B",P_RATIO_B,IF(AE{r}="C",P_RATIO_C,P_RATIO_D)))'
        ws[f"AK{r}"] = f"=1/AJ{r}"
        for col in ("V", "W", "X", "Y", "Z", "AA", "AI", "AJ", "AK"):
            ws[f"{col}{r}"].number_format = "0.00" if col not in ("AJ", "AK") else ("0" if col == "AJ" else "0.000")
            ws[f"{col}{r}"].font = F_BASE
        ws[f"AA{r}"].font = F_BOLD
        for col in ("AB", "AD", "AE", "AF", "AG", "AH"):
            ws[f"{col}{r}"].font = F_BASE
        ws[f"AE{r}"].font = F_BOLD
    last = 5 + len(S)
    c.score_last = last
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:AM{last}"
    class_format(ws, f"AE6:AE{last}")
    class_format(ws, f"AB6:AB{last}")
    class_format(ws, f"AG6:AG{last}")
    status_format(ws, f"AL6:AL{last}")
    for name, col in (("S_ID", "A"), ("S_GRUPO", "C"), ("S_PRODUCTO", "D"), ("S_PROGRAMA", "E"), ("S_PROCESO", "F"),
                      ("S_LINEA", "G"), ("S_TECNO", "H"), ("S_TIPO", "I"), ("S_CB", "K"), ("S_MONT", "L"), ("S_MEB", "M"),
                      ("S_AVH", "R"), ("S_ESTADO", "U"), ("S_C1", "V"), ("S_SCORE", "AA"), ("S_CLASE_SCORE", "AB"), ("S_FORZ", "AC"),
                      ("S_SUF", "AD"), ("S_CLASE", "AE"), ("S_PROV", "AF"), ("S_ETIQ", "AG"), ("S_MOTIVO", "AH"), ("S_PUNTOS", "AI"),
                      ("S_FTE", "AK")):
        define(c.wb, name, "03_SCORE_CRITICIDAD", f"${col}$6:${col}${last}")


# ======================================================================================
# 02_MAESTRO_CRITICIDAD
# ======================================================================================
MAESTRO_COLS = [
    ("ID", "ID", 22), ("EQUIPO", "Equipo SAP", 22), ("INVENTARIO", "Inventario", 18), ("PLANO", "Plano / Campo clasif. SAP", 16),
    ("UBIC", "Ubicación técnica", 26), ("NIVEL", "Nivel jerárquico", 8), ("UBIC_PADRE", "Ubicación padre", 24), ("EQUIPO_PADRE", "Equipo padre (por ubicación)", 22),
    ("UBIC_RAIZ", "Ubicación raíz aplicada", 20), ("TIPO_COINC", "Tipo coincidencia", 14), ("ORIGEN_REGLA", "Origen de la regla de clasificación", 34),
    ("DESC", "Descripción", 40), ("DESC_UBIC", "Sección SAP (desc. ubicación)", 30), ("TP", "Tipo objeto SAP", 10), ("TIPO_ACTIVO", "Tipo de activo", 24),
    ("TECNOLOGIA", "Tecnología", 22), ("FABRICANTE", "Marca / fabricante", 18), ("MODELO", "Modelo", 18), ("ANIO", "Año puesta servicio", 9),
    ("ESTADO", "Estado del registro", 34), ("EN_PARQUE", "En parque principal", 9), ("MOTIVO_ESTADO", "Motivo del estado", 48),
    ("PROGRAMA", "Programa", 12), ("PRODUCTO", "Producto", 16), ("LINEA", "Línea", 28), ("PROCESO", "Proceso", 18), ("DESC_UBIC2", "Sección", 30),
    ("GRUPO_MANT", "Grupo mantenimiento", 24), ("NUEVO_PRODUCTO", "Nuevo producto", 8), ("CB", "Cuello botella", 8), ("CB_DETALLE", "Detalle cuello de botella", 40),
    ("MONTAJE_PPAL", "Montaje principal", 9), ("MEB_CRIT", "MEB crítico", 8), ("SAP_ABC", "SAP ABC", 7),
    ("CRIT_PREV_MC", "Criticidad previa MAESTRO_CRITICIDAD", 18), ("CRIT_PREV_INT", "Criticidad previa MAESTRO_INTEGRAL", 22),
    ("AV_N", "Averías 2025→ (nº)", 9), ("AV_H", "Horas paro 2025→", 9),
    ("=SCORE", "Score", 8), ("=CLASE", "Criticidad", 9), ("=MOTIVO", "Motivo de criticidad", 70), ("=PUNTOS", "Puntos equivalentes", 9), ("=FTE", "FTE", 8),
    ("DEP_PREVIA", "Depuración SAP previa", 34), ("FILA_DEP", "Registro en depuración", 22), ("FUENTE", "Fuente", 34), ("HOJA", "Hoja origen", 18),
    ("FILA", "Fila origen", 8), ("VALIDACION", "Estado de validación", 28), ("OBS", "Observación", 50),
]


def _obs(r):
    o = []
    if r["FAMILIA_PPT"]:
        o.append("Familia contexto PPT: " + r["FAMILIA_PPT"])
    if r.get("CONTROL_JERARQUIA") == "OK":
        o.append("Regla jerárquica = depuración previa")
    if "Ambig" in _s(r.get("AV_FUENTE")):
        o.append("Averías asignadas por descripción repetida")
    return " · ".join(o)


def hoja_maestro(c: Ctx):
    ws = c.ws["02_MAESTRO_CRITICIDAD"]
    title(ws, "02 · MAESTRO DE CRITICIDAD – BASE ÚNICA DE ACTIVOS",
          "Todos los registros (SAP completo + fichero de montaje + cuellos sin equivalencia). Nada se borra: los excluidos conservan estado y motivo.")
    inv = c.inv.copy()
    inv["DESC_UBIC2"] = inv["DESC_UBIC"]
    inv["OBS"] = inv.apply(_obs, axis=1)
    orden = {"S": 0, "N": 1}
    inv = inv.sort_values(by=["EN_PARQUE", "GRUPO_MANT", "LINEA", "ID"], key=lambda s: s.map(orden) if s.name == "EN_PARQUE" else s).reset_index(drop=True)
    header_row(ws, 5, 1, [h for _, h, _ in MAESTRO_COLS], [w for _, _, w in MAESTRO_COLS], height=45)
    col_of = {f: get_column_letter(j + 1) for j, (f, _, _) in enumerate(MAESTRO_COLS)}
    a, u = col_of["ID"], col_of["EN_PARQUE"]
    look = {"=SCORE": "S_SCORE", "=CLASE": "S_ETIQ", "=MOTIVO": "S_MOTIVO", "=PUNTOS": "S_PUNTOS", "=FTE": "S_FTE"}
    c.maestro_row = {}
    for i, rec in inv.iterrows():
        r = 6 + i
        c.maestro_row[rec["ID"]] = r
        for j, (f, h, w) in enumerate(MAESTRO_COLS):
            col = get_column_letter(j + 1)
            cell = ws[f"{col}{r}"]
            if f.startswith("="):
                rng = look[f]
                if f == "=FTE":
                    cell.value = f'=IF(${u}{r}="S",INDEX({rng},MATCH(${a}{r},S_ID,0)),0)'
                    cell.number_format = "0.000"
                else:
                    cell.value = f'=IF(${u}{r}="S",INDEX({rng},MATCH(${a}{r},S_ID,0)),"—")'
                    if f in ("=SCORE", "=PUNTOS"):
                        cell.number_format = "0.00"
                cell.font = F_BASE
                continue
            v = rec.get(f)
            if f == "ANIO":
                v = int(v) if (v is not None and not pd.isna(v)) else None
            elif f == "NIVEL":
                v = int(v) if (v is not None and not pd.isna(v)) else None
            elif isinstance(v, float) and pd.isna(v):
                v = None
            elif isinstance(v, tuple):
                v = str(v)
            cell.value = v
            cell.font = F_BASE
        if rec["EN_PARQUE"] != "S":
            for j in range(len(MAESTRO_COLS)):
                ws.cell(row=r, column=j + 1).fill = FILL["pendiente"] if "Pendiente" in rec["ESTADO"] else FILL["total"]
    last = 5 + len(inv)
    c.maestro_last = last
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:{get_column_letter(len(MAESTRO_COLS))}{last}"
    class_format(ws, f"{col_of['=CLASE']}6:{col_of['=CLASE']}{last}")
    status_format(ws, f"{col_of['VALIDACION']}6:{col_of['VALIDACION']}{last}")
    for name, f in (("M_ID", "ID"), ("M_ESTADO", "ESTADO"), ("M_PARQUE", "EN_PARQUE"), ("M_GRUPO", "GRUPO_MANT"), ("M_PROGRAMA", "PROGRAMA"),
                    ("M_TIPO", "TIPO_ACTIVO"), ("M_CB", "CB"), ("M_FTE", "=FTE"), ("M_EQUIPO", "EQUIPO"), ("M_FAMILIA", "OBS"),
                    ("M_FUENTE", "FUENTE"), ("M_TP", "TP"), ("M_DEP", "DEP_PREVIA"), ("M_LINEA", "LINEA")):
        col = col_of[f]
        define(c.wb, name, "02_MAESTRO_CRITICIDAD", f"${col}$6:${col}${last}")
    c.maestro_cols = col_of


# ======================================================================================
# Estadísticas de origen (MAESTRO_FABRICA_FINAL 01_Maestro_Fabrica)
# ======================================================================================
def stats_fabrica(c: Ctx):
    f = c.M["fab"]
    def n(bloque, crit=None, startswith=None):
        x = f[f["Bloque"] == bloque]
        if crit is not None:
            x = x[x["Criticidad"].fillna("") == crit]
        if startswith is not None:
            x = x[x["Criticidad"].fillna("").str.startswith(startswith)]
        return len(x), (int(x["FILA_ORIGEN"].min()) if len(x) else None), (int(x["FILA_ORIGEN"].max()) if len(x) else None)
    st = {
        "CB": n("Cuello botella"), "A2": n("Histórico paros", "A2 - Alto paro"), "MONT_A1": n("Montaje detalle", "A1 - Montaje"),
        "MEB_MONT": n("Montaje detalle", "A1 - MEB ECO montaje candidato"), "MEB_OF": n("MEB ECO oficial"),
        "B": n("Histórico paros", "B - Alta carga"), "C": n("Histórico paros", "C - Media"),
        "C_PEND": n("Montaje detalle", startswith="C/Pendiente"), "D": n("Histórico paros", "D - Baja/ruta"),
    }
    # Solapes: máquinas de cuellos que también figuran en el bloque de histórico de paros
    cb_dest = set(c.M["cb"]["ACTIVO_DESTINO"].dropna().str.upper().str.strip())
    hist = f[f["Bloque"] == "Histórico paros"]
    hist_names = hist["Producto / Familia / Activo"].fillna("").str.upper().str.strip()
    st["SOLAPE_CB_HIST"] = int(hist_names.isin(cb_dest).sum())
    st["SOLAPE_CB_A2"] = int(hist_names[hist["Criticidad"] == "A2 - Alto paro"].isin(cb_dest).sum())
    cbm = c.M["cb"]
    st["CB_UNICOS"] = int(cbm.loc[cbm["ID_ASIGNADO"] != "", "ID_ASIGNADO"].nunique())
    st["CB_SIN"] = int((cbm["ID_ASIGNADO"] == "").sum())
    c.fab = st
    return st


# ======================================================================================
# 04_FTE_CRITICIDAD
# ======================================================================================
def hoja_fte_criticidad(c: Ctx):
    ws = c.ws["04_FTE_CRITICIDAD"]
    st = stats_fabrica(c)
    title(ws, "04 · FTE BASE POR CRITICIDAD", "FTE de la clase = Número de activos de la clase ÷ Ratio aplicado · FTE base = Σ clases (diapositiva 4)")
    widths(ws, {"A": 34, "B": 13, "C": 14, "D": 13, "E": 22, "F": 12, "G": 13, "H": 11, "I": 11, "J": 12, "K": 12, "L": 24, "M": 50})
    section(ws, 5, "1. Cálculo oficial reproducido (diapositiva 4 · tabla 'CÁLCULO BASE POR CRITICIDAD')", 13)
    header_row(ws, 6, 1, ["Clase", "Activos", "Ratio aplicado (activos/FTE)", "FTE de la clase", "Fórmula", "Puntos eq. por activo",
                          "Puntos equivalentes", "% activos", "% FTE", "Valor PPT", "Diferencia", "Estado", "Lectura"])
    of = O.OF
    for i, k in enumerate("ABCD"):
        r = 7 + i
        put(ws, f"A{r}", f"{k} · " + {"A": "crítica", "B": "media-alta", "C": "carga media", "D": "baja"}[k], F_BOLD)
        put(ws, f"B{r}", f"=P_N_{k}", fmt="int")
        put(ws, f"C{r}", f"=P_RATIO_{k}", fmt="0")
        put(ws, f"D{r}", f"=B{r}/C{r}", F_BOLD, fmt="fte")
        put(ws, f"E{r}", f"= {k} ÷ ratio {k}")
        put(ws, f"F{r}", f"=P_PT_{k}", fmt="0.0")
        put(ws, f"G{r}", f"=B{r}*F{r}", fmt="#,##0.0")
        put(ws, f"H{r}", f"=B{r}/$B$11", fmt="pct")
        put(ws, f"I{r}", f"=D{r}/$D$11", fmt="pct")
        put(ws, f"J{r}", of[f"FTE_{k}"][0], fill_="oficial", fmt="fte")
        put(ws, f"K{r}", f"=D{r}-J{r}", fmt="0.000")
        put(ws, f"L{r}", f'=IF(ABS(K{r})<=P_TOL_EXACT,"OK","REVISAR")')
        put(ws, f"M{r}", {"A": "Críticos, cuellos de botella, montaje principal y MEB crítico.",
                          "B": "Impacto productivo elevado con alternativa parcial.",
                          "C": "Carga media; incluye los activos sin información suficiente (nunca D).",
                          "D": "Bajo impacto demostrado."}[k], align=WRAP)
        ws[f"A{r}"].fill = FILL[k]
        if k == "A":
            ws[f"A{r}"].font = F_HEAD
    r = 11
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    put(ws, f"B{r}", "=SUM(B7:B10)", F_BOLD, "total", "int")
    put(ws, f"C{r}", "", fill_="total")
    put(ws, f"D{r}", "=SUM(D7:D10)", F_BOLD, "total", "fte")
    put(ws, f"E{r}", "= Σ FTE A+B+C+D", fill_="total")
    put(ws, f"F{r}", "", fill_="total")
    put(ws, f"G{r}", "=SUM(G7:G10)", F_BOLD, "total", "#,##0.0")
    put(ws, f"H{r}", "=SUM(H7:H10)", fill_="total", fmt="pct")
    put(ws, f"I{r}", "=SUM(I7:I10)", fill_="total", fmt="pct")
    put(ws, f"J{r}", of["FTE_BASE"][0], F_BOLD, "oficial", "fte")
    put(ws, f"K{r}", "=D11-J11", fill_="total", fmt="0.000")
    put(ws, f"L{r}", '=IF(ABS(K11)<=P_TOL_EXACT,"OK","REVISAR")', F_BOLD)
    put(ws, f"M{r}", "Puntos equivalentes = peso relativo de cartera; no son horas ni sustituyen al FTE.", align=WRAP)
    status_format(ws, "L7:L11")
    define(c.wb, "K_FTE_BASE", "04_FTE_CRITICIDAD", "$D$11")
    define(c.wb, "K_PARQUE_CLASES", "04_FTE_CRITICIDAD", "$B$11")
    note(ws, "A12", "Los ratios son un criterio interno de dimensionamiento de cartera (carga de trabajo, criticidad, benchmarking y contexto operativo). "
         "SMRP y NASA NPR 8831.2F respaldan la metodología por carga de trabajo; no publican literalmente estos ratios A/B/C/D.")

    section(ws, 14, "2. Sensibilidad por rango de ratio (diapositiva 3: A 8–12 · B 12–20 · C 20–35 · D 35–60)", 13)
    header_row(ws, 15, 1, ["Clase", "Activos", "Ratio mínimo", "FTE (ratio mínimo · mayor cobertura)", "Ratio medio", "FTE (ratio medio)",
                          "Ratio oficial (máximo)", "FTE (oficial)", "Δ FTE mín. vs oficial", "", "", "", "Lectura"])
    for i, k in enumerate("ABCD"):
        r = 16 + i
        put(ws, f"A{r}", k, F_BOLD, k)
        put(ws, f"B{r}", f"=P_N_{k}", fmt="int")
        put(ws, f"C{r}", f"=P_RMIN_{k}", fmt="0")
        put(ws, f"D{r}", f"=B{r}/C{r}", fmt="fte")
        put(ws, f"E{r}", f"=(P_RMIN_{k}+P_RATIO_{k})/2", fmt="0.0")
        put(ws, f"F{r}", f"=B{r}/E{r}", fmt="fte")
        put(ws, f"G{r}", f"=P_RATIO_{k}", fmt="0")
        put(ws, f"H{r}", f"=B{r}/G{r}", fmt="fte")
        put(ws, f"I{r}", f"=D{r}-H{r}", fmt="fte")
    r = 20
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "BDFHI":
        put(ws, f"{col}{r}", f"=SUM({col}16:{col}19)", F_BOLD, "total", "fte" if col != "B" else "int")
    put(ws, "M16", "El resultado oficial usa el extremo superior de cada rango (menos FTE por activo). "
        "El escenario de mayor cobertura (ratio mínimo) multiplica la base.", align=WRAP)
    ws.merge_cells("M16:M20")
    define(c.wb, "K_FTE_RMIN", "04_FTE_CRITICIDAD", "$D$20")
    define(c.wb, "K_FTE_RMED", "04_FTE_CRITICIDAD", "$F$20")

    section(ws, 22, "3. ¿De dónde sale la distribución oficial A/B/C/D? (MAESTRO_FABRICA_FINAL.xlsx · 01_Maestro_Fabrica)", 13)
    header_row(ws, 23, 1, ["Componente", "Filas en fuente", "Filas origen", "Clase oficial", "Valor oficial", "Diferencia", "Estado", "", "", "", "", "", "Explicación"])
    rows = [
        ("Cuellos de botella (A1)", st["CB"], "A"), ("Histórico paros A2 - Alto paro", st["A2"], "A"),
        ("Montaje A1 - Montaje", st["MONT_A1"], "A"), ("Montaje A1 - MEB ECO montaje candidato", st["MEB_MONT"], "A"),
        ("MEB ECO oficial (filas agregadas)", st["MEB_OF"], "A"),
    ]
    r = 24
    for lbl, (n, f0, f1), k in rows:
        put(ws, f"A{r}", lbl)
        put(ws, f"B{r}", n, fill_="dato", fmt="int")
        put(ws, f"C{r}", f"{f0}–{f1}" if f0 else "")
        put(ws, f"D{r}", k, F_BOLD, k)
        r += 1
    put(ws, f"A{r}", "Total A reconstruido", F_BOLD, "total")
    put(ws, f"B{r}", f"=SUM(B24:B{r - 1})", F_BOLD, "total", "int")
    put(ws, f"D{r}", "A", F_BOLD, "A")
    put(ws, f"E{r}", "=P_N_A", fill_="oficial", fmt="int")
    put(ws, f"F{r}", f"=B{r}-E{r}", fmt="int")
    put(ws, f"G{r}", f'=IF(F{r}=0,"RECONSTRUIDO","REVISAR")', F_BOLD)
    put(ws, f"M{r}", f"A = recuento de FILAS con criticidad A. Mezcla granularidades: las 6 filas MEB oficial representan 75 activos; "
        f"los 129 cuellos corresponden a {st['CB_UNICOS']} máquinas SAP distintas + {st['CB_SIN']} entradas sin equivalencia (líneas/islas); "
        f"{st['SOLAPE_CB_A2']} máquinas A2 también son cuello y {st['SOLAPE_CB_HIST']} máquinas de cuellos figuran además en el bloque de paros.", align=WRAP)
    ws.row_dimensions[r].height = 60
    ra = r
    r += 1
    for lbl, val, k, expl in (
        ("Histórico paros B - Alta carga", st["B"][0], "B", "Oficial 275: +156 sin registro de origen."),
        ("Histórico paros C - Media + Montaje C/Pendiente", st["C"][0] + st["C_PEND"][0], "C", "Oficial 1.200 (valor redondo declarado)."),
        ("Histórico paros D - Baja/ruta + 397 no clasificados (2.043 − 1.646)", st["D"][0] + 397, "D", "Oficial 150. Coherente con la regla 'D solo con bajo impacto demostrado', pero sin registro activo a activo."),
    ):
        put(ws, f"A{r}", lbl)
        put(ws, f"B{r}", val, fill_="dato", fmt="int")
        put(ws, f"D{r}", k, F_BOLD, k)
        put(ws, f"E{r}", f"=P_N_{k}", fill_="oficial", fmt="int")
        put(ws, f"F{r}", f"=B{r}-E{r}", fmt="int")
        put(ws, f"G{r}", "TRAZABILIDAD INCOMPLETA", F_BOLD)
        put(ws, f"M{r}", expl, align=WRAP)
        r += 1
    put(ws, f"A{r}", "B + C + D oficiales = parque − A", F_BOLD, "total")
    put(ws, f"B{r}", "=P_PARQUE_OF-P_N_A", F_BOLD, "total", "int")
    put(ws, f"E{r}", "=P_N_B+P_N_C+P_N_D", fill_="oficial", fmt="int")
    put(ws, f"F{r}", f"=B{r}-E{r}", fmt="int")
    put(ws, f"G{r}", f'=IF(F{r}=0,"OK","REVISAR")', F_BOLD)
    put(ws, f"M{r}", "El total sí cuadra: el reparto B/C/D es una asignación declarada sobre los 1.625 activos no A.", align=WRAP)
    r += 1
    put(ws, f"A{r}", "Versión anterior (68,1 FTE en DIMENSIONAMIENTO 00_Resumen B10)", F_BOLD)
    put(ws, f"B{r}", f"=E{ra}/P_RATIO_A+B{ra + 1}/P_RATIO_B+B{ra + 2}/P_RATIO_C+B{ra + 3}/P_RATIO_D", F_BOLD, fmt="fte")
    put(ws, f"E{r}", 68.1, fill_="dato", fmt="fte1")
    put(ws, f"F{r}", f"=B{r}-E{r}", fmt="fte")
    put(ws, f"G{r}", f'=IF(ABS(F{r})<=0.05,"RECONSTRUIDO","REVISAR")', F_BOLD)
    put(ws, f"M{r}", "Con A=418 y B/C/D de la fuente (119 / 189 / 1.317) se obtiene exactamente el 68,1 que quedó sin actualizar en 00_Resumen del Excel de dimensionamiento.", align=WRAP)
    status_format(ws, f"G24:G{r}")
    rtop = r + 2

    section(ws, rtop, "4. Contraste técnico activo a activo (03_SCORE_CRITICIDAD) – NO sustituye al resultado oficial", 13)
    header_row(ws, rtop + 1, 1, ["Clase", "Activos (score)", "de ellos provisionales", "FTE (ratio oficial)", "Puntos equivalentes",
                                 "Oficial (activos)", "Diferencia activos", "Oficial (FTE)", "Diferencia FTE", "Estado", "", "", "Lectura"])
    for i, k in enumerate("ABCD"):
        r = rtop + 2 + i
        put(ws, f"A{r}", k, F_BOLD, k)
        put(ws, f"B{r}", f'=COUNTIF(S_CLASE,"{k}")', fmt="int")
        put(ws, f"C{r}", f'=COUNTIFS(S_CLASE,"{k}",S_PROV,"S")', fmt="int")
        put(ws, f"D{r}", f"=B{r}/P_RATIO_{k}", fmt="fte")
        put(ws, f"E{r}", f'=SUMIFS(S_PUNTOS,S_CLASE,"{k}")', fmt="#,##0.0")
        put(ws, f"F{r}", f"=P_N_{k}", fill_="oficial", fmt="int")
        put(ws, f"G{r}", f"=B{r}-F{r}", fmt="int")
        put(ws, f"H{r}", f"=D{7 + i}", fill_="oficial", fmt="fte")
        put(ws, f"I{r}", f"=D{r}-H{r}", fmt="fte")
        put(ws, f"J{r}", f'=IF(G{r}=0,"OK","REVISAR")', F_BOLD)
    r = rtop + 6
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "BCDEFGHI":
        put(ws, f"{col}{r}", f"=SUM({col}{rtop + 2}:{col}{r - 1})", F_BOLD, "total", "int" if col in "BCFG" else ("#,##0.0" if col == "E" else "fte"))
    put(ws, f"J{r}", f'=IF(G{r}=0,"OK","REVISAR")', F_BOLD)
    status_format(ws, f"J{rtop + 2}:J{r}")
    put(ws, f"M{rtop + 2}", "El inventario fila a fila (SAP depurado + montaje) no llega a 2.043 y su reparto por score difiere del declarado. "
        "Se muestra para trazabilidad y priorización; el resultado oficial sigue siendo el de la tabla 1.", align=WRAP)
    ws.merge_cells(f"M{rtop + 2}:M{r}")
    define(c.wb, "K_TEC_ACTIVOS", "04_FTE_CRITICIDAD", f"$B${r}")
    define(c.wb, "K_TEC_FTE", "04_FTE_CRITICIDAD", f"$D${r}")
    define(c.wb, "K_TEC_PROV", "04_FTE_CRITICIDAD", f"$C${r}")
    c.kpi["tec_rows"] = (rtop + 2, r)
    ws.freeze_panes = "B7"


# ======================================================================================
# 05_FTE_POR_GRUPOS
# ======================================================================================
def _tabla_dimension(c, ws, r0, titulo, rango, valores, ncol_label="Grupo"):
    section(ws, r0, titulo, 16)
    header_row(ws, r0 + 1, 1, [ncol_label, "Activos", "A", "B", "C", "D", "Provisionales", "Puntos equivalentes",
                               "FTE A", "FTE B", "FTE C", "FTE D", "FTE base", "% FTE técnico", "Fuente", "Estado"])
    r = r0 + 2
    first = r
    for v in valores:
        put(ws, f"A{r}", v, F_BOLD)
        put(ws, f"B{r}", f'=COUNTIF({rango},$A{r})', fmt="int")
        for j, k in enumerate("ABCD"):
            col = "CDEF"[j]
            put(ws, f"{col}{r}", f'=COUNTIFS({rango},$A{r},S_CLASE,"{k}")', fmt="int")
            put(ws, f"{'IJKL'[j]}{r}", f"={col}{r}/P_RATIO_{k}", fmt="fte")
        put(ws, f"G{r}", f'=COUNTIFS({rango},$A{r},S_PROV,"S")', fmt="int")
        put(ws, f"H{r}", f"=SUMIFS(S_PUNTOS,{rango},$A{r})", fmt="#,##0.0")
        put(ws, f"M{r}", f"=SUM(I{r}:L{r})", F_BOLD, fmt="fte")
        put(ws, f"N{r}", f"=M{r}/K_TEC_FTE", fmt="pct")
        put(ws, f"O{r}", "03_SCORE_CRITICIDAD (activo a activo)")
        put(ws, f"P{r}", f'=IF(G{r}>0,"PROVISIONAL – "&G{r}&" activos pendientes","PROPUESTA TÉCNICA – validar")')
        r += 1
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "BCDEFGHIJKLMN":
        put(ws, f"{col}{r}", f"=SUM({col}{first}:{col}{r - 1})", F_BOLD, "total",
            "int" if col in "BCDEFG" else ("pct" if col == "N" else ("#,##0.0" if col == "H" else "fte")))
    put(ws, f"O{r}", "Control: debe igualar el total técnico", fill_="total")
    put(ws, f"P{r}", f'=IF(AND(B{r}=K_TEC_ACTIVOS,ABS(M{r}-K_TEC_FTE)<=P_TOL_EXACT),"OK","ERROR")', F_BOLD)
    status_format(ws, f"P{first}:P{r}")
    return r


def hoja_fte_grupos(c: Ctx):
    ws = c.ws["05_FTE_POR_GRUPOS"]
    title(ws, "05 · FTE POR GRUPOS, PRODUCTO, PROCESO Y TECNOLOGÍA",
          "1) Lo presentado se conserva como referencia (pendiente). 2) Se reproduce su fuente. 3) Se reconstruye desde los activos sin forzar el cuadre.")
    widths(ws, {"A": 30, "B": 10, "C": 9, "D": 9, "E": 9, "F": 9, "G": 12, "H": 12, "I": 9, "J": 9, "K": 9, "L": 9, "M": 11, "N": 11, "O": 34, "P": 40})
    section(ws, 5, "1. Reparto presentado en la diapositiva 7 'FTE MANTENIMIENTO POR GRUPO' (referencia presentada · PENDIENTE)", 16)
    header_row(ws, 6, 1, ["Grupo (PPT)", "FTE presentado", "% del FTE base oficial", "Estado", "", "", "", "", "", "", "", "", "", "", "Fuente", "Observación"])
    r = 7
    for g, v in O.GRUPOS_S7:
        put(ws, f"A{r}", g, F_BOLD)
        put(ws, f"B{r}", v, fill_="oficial", fmt="fte")
        put(ws, f"C{r}", f"=B{r}/K_FTE_BASE", fmt="pct")
        put(ws, f"D{r}", "PENDIENTE DE VALIDACIÓN")
        put(ws, f"O{r}", f"{O.PPT} diap. 7 (tabla incrustada)")
        r += 1
    put(ws, f"A{r}", "Suma repartida", F_BOLD, "total")
    put(ws, f"B{r}", f"=SUM(B7:B{r - 1})", F_BOLD, "total", "fte")
    put(ws, f"C{r}", f"=B{r}/K_FTE_BASE", fill_="total", fmt="pct")
    rs = r
    r += 1
    put(ws, f"A{r}", "FTE base oficial (04)", F_BOLD)
    put(ws, f"B{r}", "=K_FTE_BASE", F_BOLD, fmt="fte")
    r += 1
    put(ws, f"A{r}", "SIN ASIGNAR (no se completa artificialmente)", F_BOLD, "pendiente")
    put(ws, f"B{r}", f"=B{r - 1}-B{rs}", F_BOLD, "pendiente", "fte")
    put(ws, f"C{r}", f"=B{r}/K_FTE_BASE", fill_="pendiente", fmt="pct")
    put(ws, f"D{r}", "TRAZABILIDAD INCOMPLETA")
    put(ws, f"P{r}", O.S7_NOTA, align=WRAP)
    status_format(ws, f"D7:D{r}")
    r += 2

    section(ws, r, "2. Reproducción de la fuente: DIMENSIONAMIENTO_MANTENIMIENTO_DATOS_FINAL.xlsx · 03_Secciones (filas 6–13)", 16)
    header_row(ws, r + 1, 1, ["Sección / producto", "A", "B", "C", "D", "A_MEB", "FTE fuente (fórmula)", "FTE presentado (diap. 7)",
                              "Diferencia", "Estado", "", "", "", "", "Fila fuente", "Observación"])
    r += 2
    r0 = r
    for g, a, b, cc, d, am, fila in O.SECCIONES_DIM:
        put(ws, f"A{r}", g, F_BOLD)
        for col, v in zip("BCDEF", (a, b, cc, d, am)):
            put(ws, f"{col}{r}", v, fill_="dato", fmt="int")
        put(ws, f"G{r}", f"=B{r}/P_RATIO_A+C{r}/P_RATIO_B+D{r}/P_RATIO_C+E{r}/P_RATIO_D+F{r}/P_RATIO_A", F_BOLD, fmt="fte")
        put(ws, f"H{r}", f"=INDEX($B$7:$B$14,MATCH(A{r},$A$7:$A$14,0))", fill_="oficial", fmt="fte")
        put(ws, f"I{r}", f"=G{r}-H{r}", fmt="fte")
        put(ws, f"J{r}", f'=IF(ABS(I{r})<=0.01,"OK","DIFERENCIA")', F_BOLD)
        put(ws, f"O{r}", f"03_Secciones fila {fila}")
        obs = ""
        if am:
            obs = "A y A_MEB tienen el mismo valor y ambos suman: posible DOBLE CONTEO del activo crítico MEB."
        put(ws, f"P{r}", obs, align=WRAP)
        r += 1
    put(ws, f"A{r}", "RESTO PLANTA / SOPORTE GLOBAL (cierre en fuente)", F_BOLD, "pendiente")
    put(ws, f"G{r}", f"=K_FTE_BASE-SUM(G{r0}:G{r - 1})", F_BOLD, "pendiente", "fte")
    put(ws, f"J{r}", "TRAZABILIDAD INCOMPLETA", F_BOLD)
    put(ws, f"O{r}", "03_Secciones fila 13")
    put(ws, f"P{r}", "La fuente cuadra el total por diferencia ('para cuadrar'). No es un reparto validado.", align=WRAP)
    status_format(ws, f"J{r0}:J{r}")
    r += 2

    S = c.S
    grupos = list(S.groupby("GRUPO_MANT")["FTE"].sum().sort_values(ascending=False).index)
    r = _tabla_dimension(c, ws, r, "3. Reconstrucción técnica por grupo de mantenimiento (desde los activos de 03)", "S_GRUPO", grupos) + 2
    c.kpi["grupos_tec"] = grupos
    r = _tabla_dimension(c, ws, r, "4. Por producto", "S_PRODUCTO", list(S["PRODUCTO"].value_counts().index), "Producto") + 2
    r = _tabla_dimension(c, ws, r, "5. Por programa (parque actual convencional vs nuevos productos MEB ECO)", "S_PROGRAMA",
                         list(S["PROGRAMA"].value_counts().index), "Programa") + 2
    rp = r - 2
    r = _tabla_dimension(c, ws, r, "6. Por proceso", "S_PROCESO", list(S["PROCESO"].value_counts().index), "Proceso") + 2
    r = _tabla_dimension(c, ws, r, "7. Por tecnología", "S_TECNO", list(S["TECNOLOGIA"].value_counts().index), "Tecnología") + 2
    r = _tabla_dimension(c, ws, r, "8. Por línea / sección", "S_LINEA", list(S["LINEA"].value_counts().index), "Línea / sección") + 2
    ws.freeze_panes = "B5"


# ======================================================================================
# 07_SABADOS_FORMACION
# ======================================================================================
def _vrow(ws, r, concepto, valor, unidad, formula_txt, fuente, fmt=None, estado=None, fill_=None, bold=False):
    put(ws, f"A{r}", concepto, F_BOLD if bold else F_BASE, align=WRAP)
    put(ws, f"B{r}", valor, F_BOLD if bold else F_BASE, fill_, fmt)
    put(ws, f"C{r}", unidad)
    put(ws, f"D{r}", formula_txt, align=WRAP)
    put(ws, f"E{r}", fuente, align=WRAP)
    put(ws, f"F{r}", estado if estado is not None else "")


def hoja_sabados_formacion(c: Ctx):
    ws = c.ws["07_SABADOS_FORMACION"]
    title(ws, "07 · FORMACIÓN Y DESCANSO COMPENSATORIO (LIBRANZAS)",
          "Ambos conceptos son pérdida de disponibilidad para atender el parque. 'Libranza' = descanso compensatorio por trabajar en fin de semana o festivo (no son vacaciones).")
    widths(ws, {"A": 46, "B": 14, "C": 16, "D": 46, "E": 44, "F": 30})
    section(ws, 5, "1. Formación: Plantilla nominal × Horas de formación por persona ÷ Horas productivas anuales por FTE", 6)
    header_row(ws, 6, 1, ["Concepto", "Valor", "Unidad", "Fórmula", "Fuente", "Estado"])
    _vrow(ws, 7, "Plantilla nominal", "=P_PLANTILLA", "personas", "P_PLANTILLA", "12_SUPUESTOS · diap. 5/7", "int", "PENDIENTE DE VALIDACIÓN")
    _vrow(ws, 8, "Horas de formación por persona y año", "=P_FORM_H", "h/persona/año", "P_FORM_H", "12_SUPUESTOS", "int", "SUPUESTO")
    _vrow(ws, 9, "Horas de formación totales", "=B7*B8", "h/año", "= B7 × B8", "Cálculo", "int")
    _vrow(ws, 10, "Horas productivas anuales por FTE", "=P_HORAS_PROD", "h/FTE/año", "P_HORAS_PROD", "12_SUPUESTOS", "int", "SUPUESTO")
    _vrow(ws, 11, "FTE FORMACIÓN – resultado exacto", "=B9/B10", "FTE", "= B7 × B8 ÷ B10", "Cálculo", "0.000", bold=True)
    _vrow(ws, 12, "FTE formación – valor presentado a Dirección", O.OF["W_FORM"][0], "FTE", "Diapositiva 5 'Formación estimación +3'", O.PPT, "0", fill_="oficial")
    _vrow(ws, 13, "Diferencia (exacto − presentado)", "=B11-B12", "FTE", "", "", "0.000")
    put(ws, "F13", '=IF(ABS(B13)<=P_TOL_RED,"OK (redondeo)","REVISAR")', F_BOLD)
    define(c.wb, "K_FORM", "07_SABADOS_FORMACION", "$B$11")

    section(ws, 15, "2. Fines de semana y festivos: Personas por jornada × Jornadas anuales ÷ Jornadas efectivas por FTE", 6)
    header_row(ws, 16, 1, ["Concepto", "Valor", "Unidad", "Fórmula", "Fuente", "Estado"])
    _vrow(ws, 17, "Personas necesarias por jornada especial (sábado)", "=P_PERS_SAB", "personas/jornada", "P_PERS_SAB", "12_SUPUESTOS", "int", "SUPUESTO")
    _vrow(ws, 18, "Sábados con cobertura al año", "=P_SABADOS", "jornadas/año", "P_SABADOS", "12_SUPUESTOS", "int", "SUPUESTO")
    _vrow(ws, 19, "Festivos trabajados al año", "=P_FESTIVOS", "jornadas/año", "P_FESTIVOS", "12_SUPUESTOS", "int", "PENDIENTE DE VALIDACIÓN", "pendiente")
    _vrow(ws, 20, "Jornadas de descanso compensatorio generadas", "=B17*(B18+B19)", "jornadas/año", "= B17 × (B18 + B19)", "Cálculo", "int")
    _vrow(ws, 21, "Jornadas efectivas por FTE y año", "=P_JORNADAS", "días/FTE/año", "P_JORNADAS", "12_SUPUESTOS", "int", "SUPUESTO")
    _vrow(ws, 22, "FTE libranzas – método por jornadas", "=B20/B21", "FTE", "= B20 ÷ B21", "Cálculo", "0.000")
    _vrow(ws, 23, "Horas anuales reales de cobertura (si existen)", "=P_H_COBERTURA", "h/año", "P_H_COBERTURA", "DATO NO DISPONIBLE", "int", "PENDIENTE DE VALIDACIÓN", "pendiente")
    _vrow(ws, 24, "FTE libranzas – método horario", '=IF(B23>0,B23/P_HORAS_PROD,"DATO NO DISPONIBLE")', "FTE", "= B23 ÷ P_HORAS_PROD", "Cálculo", "0.000")
    _vrow(ws, 25, "FTE LIBRANZAS APLICADO", "=IF(B23>0,B24,B22)", "FTE", "Horario si hay horas reales; si no, por jornadas", "Cálculo", "0.000", bold=True)
    _vrow(ws, 26, "FTE libranzas – valor presentado a Dirección", O.OF["W_SAB"][0], "FTE", "Diapositiva 5 'Sábados / libranzas +2'", O.PPT, "0", fill_="oficial")
    _vrow(ws, 27, "Diferencia (exacto − presentado)", "=B25-B26", "FTE", "", "", "0.000")
    put(ws, "F27", '=IF(ABS(B27)<=P_TOL_RED,"OK (redondeo)","REVISAR")', F_BOLD)
    define(c.wb, "K_SAB", "07_SABADOS_FORMACION", "$B$25")
    status_format(ws, "F7:F27")
    note(ws, "A29", "Formación y libranzas son capas independientes del waterfall: la cobertura organizativa (06) no debe volver a incluirlas.")
    note(ws, "A30", "Pendiente: calendario real de sábados/festivos, personas por jornada por especialidad y turnos. Sin esos datos se mantiene el supuesto del estudio.")


# ======================================================================================
# 08_PM_AVISOS
# ======================================================================================
def hoja_pm(c: Ctx):
    ws = c.ws["08_PM_AVISOS"]
    title(ws, "08 · PREVENTIVOS, AVISOS, CORRECTIVO Y TPM",
          "No se mezclan unidades: planes, puntos, avisos, órdenes, intervenciones y horas se muestran por separado. Se prioriza la hora real cuando existe.")
    widths(ws, {"A": 44, "B": 14, "C": 18, "D": 16, "E": 46, "F": 34, "G": 30})
    av = c.M["av"]
    av25 = av[(av["FECHA"] >= "2025-01-01") & (av["FECHA"] < "2026-01-01")]
    av26 = av[av["FECHA"] >= "2026-01-01"]
    logs = c.M["logs"]
    section(ws, 5, "1. Unidades disponibles (cada fila en su unidad)", 7)
    header_row(ws, 6, 1, ["Concepto", "Cantidad", "Unidad", "Periodo", "Fuente", "Uso en el cálculo oficial", "Estado"])
    rows = [
        ("Planes de mantenimiento preventivo", "=P_PLANES_PM", "planes", "vigentes", O.PPT_CTX + " diap. 1", "Informativo", "DATO GLOBAL"),
        ("Puntos de preventivo", "=P_PUNTOS_PM", "puntos", "año", O.PPT_CTX + " diap. 1", "Sumado (60.000 unidades)", "DATO GLOBAL – sin detalle por plan"),
        ("Avisos de intervención en máquina", "=P_AVISOS", "avisos", "2024", O.PPT_CTX + " diap. 1", "Sumado (60.000 unidades)", "DATO GLOBAL"),
        ("Intervenciones TPM", "=P_TPM", "intervenciones", "2024", O.PPT_CTX + " diap. 1", "Informado, NO sumado", "OFICIAL: TPM excluido"),
        ("Horas de máquina parada (declaradas)", 10000, "h máquina", "declarado", O.PPT_CTX + " diap. 1", "Contexto", "DATO GLOBAL"),
        ("Puntos de preventivo 'yoY' (diapositiva 8)", O.OF["S8_PUNTOS_PM"][0], "puntos", "no indicado", O.PPT + " diap. 8", "No usado", "DIFERENCIA (ver 17)"),
        ("Averías SAP con orden – 2025", len(av25), "registros", "ene–dic 2025", "SAP.EXPORT.AVERIAS.LISTADO.xlsx", "Criterio C4 (03)", "DATO FUENTE"),
        ("Horas de paro SAP – 2025", round(av25["MIN"].sum() / 60, 1), "h máquina", "ene–dic 2025", "SAP.EXPORT.AVERIAS.LISTADO.xlsx (min ÷ 60)", "Criterio C4 (03)", "DATO FUENTE"),
        ("Averías SAP con orden – 2026", len(av26), "registros", f"ene–{av26['FECHA'].max():%d/%m} 2026", "SAP.EXPORT.AVERIAS.LISTADO.xlsx", "Criterio C4 (03)", "DATO FUENTE"),
        ("Horas de paro SAP – 2026", round(av26["MIN"].sum() / 60, 1), "h máquina", f"ene–{av26['FECHA'].max():%d/%m} 2026", "SAP.EXPORT.AVERIAS.LISTADO.xlsx (min ÷ 60)", "Criterio C4 (03)", "DATO FUENTE"),
        ("Paros líneas de montaje L1 + L3 (2025)", round(logs["CP"].sum(), 0), "CP (min, unidad a validar)", "ene–dic 2025", "data.1.xlsx · data.2.xlsx", "Criterio C4 montaje", "PENDIENTE: unidad de paro"),
    ]
    for i, (a, b, cc, d, e, f, g) in enumerate(rows):
        r = 7 + i
        put(ws, f"A{r}", a)
        put(ws, f"B{r}", b, F_BASE, "dato" if not str(b).startswith("=") else None, "#,##0")
        put(ws, f"C{r}", cc)
        put(ws, f"D{r}", d)
        put(ws, f"E{r}", e, align=WRAP)
        put(ws, f"F{r}", f)
        put(ws, f"G{r}", g)
    status_format(ws, f"G7:G{6 + len(rows)}")
    r = 8 + len(rows)
    note(ws, f"A{r}", "Horas de paro = horas de MÁQUINA parada, no horas de mano de obra de mantenimiento. No se convierten directamente a FTE.")

    r += 2
    section(ws, r, "2. Cálculo oficial de la carga planificada (diapositiva 5 'Preventivos +10')", 7)
    header_row(ws, r + 1, 1, ["Concepto", "Valor", "Unidad", "", "Fórmula", "Fuente", "Estado"])
    b = r + 2
    vals = [
        ("Unidades de carga = puntos PM + avisos", "=P_PUNTOS_PM+P_AVISOS", "unidades/año", "= P_PUNTOS_PM + P_AVISOS", "DIMENSIONAMIENTO 01_Supuestos B12", "MEZCLA DE UNIDADES"),
        ("Productividad por FTE", "=P_PRODUCT", "unidades/FTE/año", "P_PRODUCT", "DIMENSIONAMIENTO 01_Supuestos B13", "HIPÓTESIS"),
        ("FTE PREVENTIVO + AVISOS – exacto", f"=B{b}/B{b + 1}", "FTE", "= unidades ÷ productividad", "Cálculo", ""),
        ("Tiempo implícito por unidad", f"=P_HORAS_PROD/B{b + 1}*60", "min/unidad", "= P_HORAS_PROD ÷ P_PRODUCT × 60", "Cálculo", "PENDIENTE DE VALIDACIÓN"),
        ("FTE TPM sumado (solo si P_TPM_SUMAR = 1)", "=IF(P_TPM_SUMAR=1,P_TPM*P_H_TPM/P_HORAS_PROD,0)", "FTE", "TPM × h/interv. ÷ horas productivas", "Criterio oficial: 0", ""),
        ("FTE CARGA PLANIFICADA APLICADA", f"=B{b + 2}+B{b + 4}", "FTE", "PM + avisos (+ TPM si se activa)", "Cálculo", ""),
        ("Valor presentado a Dirección", O.OF["W_PM"][0], "FTE", "", O.PPT + " diap. 5", ""),
        ("Diferencia (aplicado − presentado)", f"=B{b + 5}-B{b + 6}", "FTE", "", "", ""),
    ]
    for i, (a, v, u, fx, fu, es) in enumerate(vals):
        rr = b + i
        put(ws, f"A{rr}", a, F_BOLD if "APLICADA" in a or "exacto" in a else F_BASE)
        put(ws, f"B{rr}", v, F_BOLD if "APLICADA" in a else F_BASE, "oficial" if "presentado" in a else None,
            "#,##0" if i in (0, 1) else ("0.0" if i == 3 else "0.000"))
        put(ws, f"C{rr}", u)
        put(ws, f"D{rr}", "")
        put(ws, f"E{rr}", fx, align=WRAP)
        put(ws, f"F{rr}", fu, align=WRAP)
        put(ws, f"G{rr}", es)
    put(ws, f"G{b + 7}", f'=IF(ABS(B{b + 7})<=P_TOL_RED,"OK (redondeo)","REVISAR")', F_BOLD)
    status_format(ws, f"G{b}:G{b + 7}")
    define(c.wb, "K_PM", "08_PM_AVISOS", f"$B${b + 5}")
    define(c.wb, "K_PM_MIN_UNIDAD", "08_PM_AVISOS", f"$B${b + 3}")
    r = b + 9

    section(ws, r, "3. Separación de la carga de mantenimiento: qué está en cada capa", 7)
    header_row(ws, r + 1, 1, ["Tipo de carga", "Dato disponible", "Unidad", "¿Sumado como FTE?", "Capa del modelo", "Estado", "Observación"])
    tipos = [
        ("Preventivo", "50.000 puntos (global)", "puntos", "Sí", "Carga planificada (08)", "DATO GLOBAL", "Falta plan por máquina: frecuencia × duración = horas reales."),
        ("Predictivo", "DATO NO DISPONIBLE", "—", "No", "—", "PENDIENTE", "No hay inventario de rutas predictivas."),
        ("Correctivo (averías)", f"{len(av):,} órdenes / {av['MIN'].sum() / 60:,.0f} h máquina (2025→{av['FECHA'].max():%m/%Y})".replace(",", "."), "órdenes / h máquina", "Indirecto", "Base por criticidad (04) + criterio C4", "DATO FUENTE", "El ratio por clase absorbe la atención correctiva."),
        ("Avisos", "10.000 avisos (2024)", "avisos", "Sí", "Carga planificada (08)", "DATO GLOBAL", "Riesgo de solape con correctivo de la base: validar."),
        ("Inspecciones", "DATO NO DISPONIBLE", "—", "No", "—", "PENDIENTE", ""),
        ("Revisiones reglamentarias", "DATO NO DISPONIBLE", "—", "No", "—", "PENDIENTE", "Instalaciones de planta (compresores, AT/BT, contraincendios…) sin inventario fila a fila."),
        ("Mejoras / proyectos", "DATO NO DISPONIBLE", "—", "No", "—", "PENDIENTE", "No incluidas en el dimensionamiento presentado."),
        ("TPM", "3.000 intervenciones (2024)", "intervenciones", "No", "Informado (actividad operativa)", "OK", "TPM informado como actividad operativa. No sumado como FTE central adicional cuando ya está incluido en la actividad del grupo."),
        ("Asistencia a producción", "DATO NO DISPONIBLE", "—", "No", "Cobertura organizativa", "PENDIENTE", ""),
    ]
    for i, row in enumerate(tipos):
        rr = r + 2 + i
        for j, v in enumerate(row):
            put(ws, f"{'ABCDEFG'[j]}{rr}", v, align=WRAP)
    status_format(ws, f"F{r + 2}:F{r + 1 + len(tipos)}")
    r = r + 3 + len(tipos)

    section(ws, r, "4. Escenarios de contraste por horas (NO oficiales) – sustituir por horas reales de planes cuando existan", 7)
    header_row(ws, r + 1, 1, ["Escenario", "min por punto PM", "h por aviso", "Horas PM / año", "Horas avisos / año", "FTE (horas ÷ P_HORAS_PROD)", "Fuente"])
    for i, (esc, mpp, hpa) in enumerate((("Bajo", 10, 0.5), ("Medio", 15, 1.0), ("Alto", 20, 1.5))):
        rr = r + 2 + i
        put(ws, f"A{rr}", esc, F_BOLD)
        put(ws, f"B{rr}", mpp, F_INPUT, "supuesto", "0")
        put(ws, f"C{rr}", hpa, F_INPUT, "supuesto", "0.0")
        put(ws, f"D{rr}", f"=P_PUNTOS_PM*B{rr}/60", fmt="#,##0")
        put(ws, f"E{rr}", f"=P_AVISOS*C{rr}", fmt="#,##0")
        put(ws, f"F{rr}", f"=(D{rr}+E{rr})/P_HORAS_PROD", F_BOLD, fmt="0.00")
        put(ws, f"G{rr}", "MAESTRO_INTEGRAL 07_Preventivo_TPM_FTE (sin TPM)")
    note(ws, f"A{r + 6}", "El oficial (9,52 FTE) equivale a ~13,8 min por unidad: entre el escenario bajo y el medio. Validar con horas estándar de los planes SAP.")


# ======================================================================================
# 06_CAPACIDAD_PERSONAL
# ======================================================================================
def hoja_capacidad(c: Ctx):
    ws = c.ws["06_CAPACIDAD_PERSONAL"]
    title(ws, "06 · CAPACIDAD DE PERSONAL, NECESIDAD Y GAP",
          "Inventario → Criticidad → Carga → Disponibilidad real → Necesidad → GAP. El GAP es ORIENTATIVO mientras plantilla, turnos y especialidades no estén validados.")
    widths(ws, {"A": 7, "B": 46, "C": 13, "D": 11, "E": 13, "F": 11, "G": 11, "H": 22, "I": 44, "J": 52})
    section(ws, 5, "1. Waterfall de la necesidad (reconstrucción de las diapositivas 5 y 7)", 10)
    header_row(ws, 6, 1, ["Capa", "Concepto", "Valor exacto", "Redondeado", "Valor PPT", "Diapositiva", "Diferencia", "Estado", "Fórmula", "Fuente / observación"])
    of = O.OF
    rows = [
        ("A", "FTE base por criticidad", "=K_FTE_BASE", of["W_BASE"][0], "4 · 5", "= Σ activos clase ÷ ratio", "04_FTE_CRITICIDAD"),
        ("B", "FTE preventivo y carga planificada (PM + avisos)", "=K_PM", of["W_PM"][0], "5", "= (P_PUNTOS_PM + P_AVISOS) ÷ P_PRODUCT", "08_PM_AVISOS · TPM no sumado"),
        ("C", "FTE formación", "=K_FORM", of["W_FORM"][0], "5", "= P_PLANTILLA × P_FORM_H ÷ P_HORAS_PROD", "07_SABADOS_FORMACION"),
        ("D", "FTE fines de semana y libranzas", "=K_SAB", of["W_SAB"][0], "5", "= P_PERS_SAB × (P_SABADOS + P_FESTIVOS) ÷ P_JORNADAS", "07_SABADOS_FORMACION"),
        ("", "Subtotal antes de cobertura (A+B+C+D)", "=SUM(C7:C10)", None, "", "= A + B + C + D", "Necesidad mínima sin cobertura organizativa"),
        ("", "Base sobre la que se aplica la cobertura", "=CHOOSE(P_COB_BASE,C11,C7+C8,C7)", None, "2", 'ELEGIR(P_COB_BASE; subtotal; A+B; A)', "Selector P_COB_BASE en 12_SUPUESTOS"),
        ("E", "Cobertura organizativa mínima", "=C12*P_COB_MIN", of["W_COB_MIN"][0], "5", "= base × P_COB_MIN", "Vacaciones, absentismo, turnos, especialidades, simultaneidad"),
        ("E", "Cobertura organizativa máxima", "=C12*P_COB_MAX", of["W_COB_MAX"][0], "5", "= base × P_COB_MAX", "Sin reincluir formación ni libranzas"),
        ("F", "NECESIDAD TOTAL MÍNIMA", "=C11+C13", of["S7_REF_MIN"][0], "7 (5: 113)", "= Subtotal + E mín.", "Diap. 7 '115-125' · diap. 5 waterfall '113–121'"),
        ("F", "NECESIDAD TOTAL MÁXIMA", "=C11+C14", of["S7_REF_MAX"][0], "7 (5: 121)", "= Subtotal + E máx.", ""),
        ("G", "Plantilla nominal", "=P_PLANTILLA*P_FTE_PERSONA", of["PLANTILLA"][0], "5 · 7", "= P_PLANTILLA × P_FTE_PERSONA", "Personas ≠ FTE productivos"),
        ("H", "GAP MÍNIMO (orientativo)", "=C15-C17", of["S7_GAP"][0], "7 (5: +33)", "= F mín. − G", "Ver conciliación (sección 2)"),
        ("H", "GAP MÁXIMO (orientativo)", "=C16-C17", None, "—", "= F máx. − G", "No presentado en el PPT (DIMENSIONAMIENTO: 44)"),
        ("", "Cobertura actual = plantilla ÷ subtotal sin cobertura organizativa", "=C17/C11", of["COB_ACTUAL"][0], "5 · 7", "= G ÷ (A+B+C+D)", "'Cobertura actual estimada de un 80 %'"),
        ("", "Plantilla ÷ necesidad total mínima", "=C17/C15", None, "—", "= G ÷ F mín.", "Lectura complementaria"),
        ("", "Plantilla ÷ necesidad total máxima", "=C17/C16", None, "—", "= G ÷ F máx.", "Lectura complementaria"),
    ]
    for i, (capa, conc, f, ppt, diap, ftxt, fuente) in enumerate(rows):
        r = 7 + i
        bold = capa in ("F", "H") or conc.startswith("Subtotal")
        put(ws, f"A{r}", capa, F_BOLD, align=CENTER)
        put(ws, f"B{r}", conc, F_BOLD if bold else F_BASE, align=WRAP)
        pct = "Cobertura" in conc or "÷ necesidad" in conc
        put(ws, f"C{r}", f, F_BOLD if bold else F_BASE, "total" if bold else None, "pct" if pct else "0.00")
        put(ws, f"D{r}", f"=ROUND(C{r},2)" if pct else f"=ROUND(C{r},0)", fmt="pct" if pct else "0")
        put(ws, f"E{r}", ppt, fill_="oficial" if ppt is not None else "pendiente", fmt="pct" if pct else "0")
        put(ws, f"F{r}", diap, align=CENTER)
        if ppt is not None:
            put(ws, f"G{r}", f"=C{r}-E{r}", fmt="0.0%" if pct else "0.00")
            tol = "0.01" if pct else "P_TOL_RED"
            put(ws, f"H{r}", f'=IF(ABS(G{r})<={tol},"OK","REVISAR – ver 17")', F_BOLD)
        else:
            put(ws, f"G{r}", "")
            put(ws, f"H{r}", "")
        put(ws, f"I{r}", ftxt, align=WRAP)
        put(ws, f"J{r}", fuente, align=WRAP)
    status_format(ws, "H7:H22")
    for name, ref in (("K_SUBTOTAL", "$C$11"), ("K_COB_BASE_VAL", "$C$12"), ("K_COB_MIN", "$C$13"), ("K_COB_MAX", "$C$14"),
                      ("K_NEC_MIN", "$C$15"), ("K_NEC_MAX", "$C$16"), ("K_PLANT_FTE", "$C$17"), ("K_GAP_MIN", "$C$18"),
                      ("K_GAP_MAX", "$C$19"), ("K_COB_ACTUAL", "$C$20")):
        define(c.wb, name, "06_CAPACIDAD_PERSONAL", ref)

    section(ws, 24, "2. Conciliación de la cobertura organizativa: el PPT usa tres bases distintas (DIFERENCIA DETECTADA)", 10)
    header_row(ws, 25, 1, ["Sel.", "Base de cobertura", "Valor base", "Cob. mín.", "Cob. máx.", "Nec. mín.", "Nec. máx.", "GAP mín. / máx.", "Cifras del PPT que reproduce", "Fuente"])
    conc = [
        (1, "Subtotal A+B+C+D (método diap. 2)", "=C11", "115–125 (diap. 5 indicador y diap. 7) · 80 % cobertura", "Diap. 2 'Subtotal FTE × factor de cobertura'"),
        (2, "Base + PM (A+B)", "=C7+C8", "GAP +34 (diap. 7) · 114–124 (Excel)", "DIMENSIONAMIENTO 04_Calculo_FTE B10:B13"),
        (3, "Solo base por criticidad (A)", "=C7", "+13 / +21 y 113–121 · GAP +33 (diap. 5)", "Diap. 5 waterfall"),
    ]
    for i, (sel, lbl, f, rep, src) in enumerate(conc):
        r = 26 + i
        put(ws, f"A{r}", sel, F_BOLD, align=CENTER)
        put(ws, f"B{r}", lbl)
        put(ws, f"C{r}", f, fmt="0.00")
        put(ws, f"D{r}", f"=C{r}*P_COB_MIN", fmt="0.00")
        put(ws, f"E{r}", f"=C{r}*P_COB_MAX", fmt="0.00")
        put(ws, f"F{r}", f"=$C$11+D{r}", F_BOLD, fmt="0.00")
        put(ws, f"G{r}", f"=$C$11+E{r}", F_BOLD, fmt="0.00")
        put(ws, f"H{r}", f'=TEXT(F{r}-$C$17,"0.0")&" / "&TEXT(G{r}-$C$17,"0.0")')
        put(ws, f"I{r}", rep, align=WRAP)
        put(ws, f"J{r}", src, align=WRAP)
    ws.conditional_formatting.add("A26:J28", FormulaRule(formula=["$A26=P_COB_BASE"], fill=FILL["supuesto"]))
    note(ws, "B29", "Fila resaltada = base activa en 12_SUPUESTOS (P_COB_BASE). Ninguna base reproduce a la vez todas las cifras del PPT: queda registrado en 17_DIFERENCIAS.")

    section(ws, 31, "3. Plantilla nominal frente a capacidad real disponible", 10)
    header_row(ws, 32, 1, ["", "Concepto", "Valor", "Unidad", "", "", "", "Estado", "Fórmula", "Observación"])
    cap = [
        ("Personas nominales de mantenimiento", "=P_PLANTILLA", "personas", "PENDIENTE DE VALIDACIÓN", "P_PLANTILLA", "'~80 personas' (diap. 5 y 7). Sin fichero de organización."),
        ("FTE contractual", "=C33*P_FTE_PERSONA", "FTE", "SUPUESTO", "= personas × P_FTE_PERSONA", "Una persona no equivale a un FTE productivo completo."),
        ("− Formación", "=-K_FORM", "FTE", "", "07", ""),
        ("− Descanso compensatorio (libranzas)", "=-K_SAB", "FTE", "", "07", ""),
        ("Capacidad disponible para el parque antes de ausencias", "=SUM(C34:C36)", "FTE", "", "= FTE − formación − libranzas", ""),
        ("Capacidad efectiva con cobertura mínima (÷ (1 + P_COB_MIN))", "=C37/(1+P_COB_MIN)", "FTE", "ORIENTATIVO", "Vacaciones, absentismo, turnos…", "Estimación orientativa: el factor está sujeto a validación organizativa."),
        ("Capacidad efectiva con cobertura máxima (÷ (1 + P_COB_MAX))", "=C37/(1+P_COB_MAX)", "FTE", "ORIENTATIVO", "", ""),
        ("Necesidad del parque y carga planificada (A + B)", "=C7+C8", "FTE", "", "04 + 08", "Lo que exige el parque sin pérdidas de capacidad."),
        ("Capacidad efectiva (mín.–máx.) ÷ necesidad A + B", '=TEXT(C39/C40,"0%")&" – "&TEXT(C38/C40,"0%")', "%", "ORIENTATIVO", "", ""),
    ]
    for i, (lbl, f, u, est, fx, obs) in enumerate(cap):
        r = 33 + i
        put(ws, f"B{r}", lbl, F_BOLD if "Capacidad disponible" in lbl else F_BASE, align=WRAP)
        put(ws, f"C{r}", f, F_BOLD if "Capacidad" in lbl else F_BASE, fmt="0.00" if u != "personas" else "int")
        put(ws, f"D{r}", u)
        put(ws, f"H{r}", est)
        put(ws, f"I{r}", fx, align=WRAP)
        put(ws, f"J{r}", obs, align=WRAP)
    status_format(ws, "H33:H41")

    section(ws, 43, "4. ¿Qué parte de la necesidad base corresponde al parque actual y qué parte a nuevos productos?", 10)
    header_row(ws, 44, 1, ["", "Fuente del reparto", "Parque actual", "Nuevos productos (MEB ECO)", "Sin asignar", "Total", "", "Estado", "Cálculo", "Observación"])
    grp = c.kpi.get("grupos_tec", [])
    rows4 = [
        ("Diapositiva 7 (referencia presentada)", "=SUMIFS('05_FTE_POR_GRUPOS'!$B$7:$B$14,'05_FTE_POR_GRUPOS'!$A$7:$A$14,\"MQ*\")",
         "=SUM('05_FTE_POR_GRUPOS'!$B$9:$B$13)", "=K_FTE_BASE-C45-D45", "=SUM(C45:E45)", "PENDIENTE DE VALIDACIÓN",
         "MQ200+MQ281 / EBOX+MANGUETAS+KMM+CONJ.DIF.+ALUMINIO", "El PPT indica 'PENDIENTE REALIZAR'. 41,8 FTE sin asignar."),
        ("DIMENSIONAMIENTO 03_Secciones (fuente)", "=SUM('05_FTE_POR_GRUPOS'!$G$21:$G$22)", "=SUM('05_FTE_POR_GRUPOS'!$G$23:$G$27)",
         "='05_FTE_POR_GRUPOS'!$G$28", "=SUM(C46:E46)", "DIFERENCIA (ver 17)", "Filas 6–13 de la fuente", "Incluye posible doble conteo A + A_MEB."),
        ("Reconstrucción técnica activo a activo (03)", '=SUMIFS(S_FTE,S_PROGRAMA,"Convencional")+SUMIFS(S_FTE,S_PROGRAMA,"Común")',
         '=SUMIFS(S_FTE,S_PROGRAMA,"MEB ECO")', '=SUMIFS(S_FTE,S_PROGRAMA,"PENDIENTE")', "=SUM(C47:E47)", "PROPUESTA TÉCNICA – validar",
         "SUMIFS sobre 03 por programa", "Base técnica, no oficial (parque fila a fila < 2.043)."),
        ("Mínimo directo: 75 activos MEB críticos confirmados × clase A", "", "=75/P_RATIO_A", "", "=D48", "OK – dato de negocio",
         "= 75 ÷ P_RATIO_A", "MAESTRO_FABRICA 02_MEB_ECO_Oficial: 56 montaje + 19 mecanizado."),
    ]
    for i, (lbl, a, b, cc, d, est, calc, obs) in enumerate(rows4):
        r = 45 + i
        put(ws, f"B{r}", lbl, align=WRAP)
        put(ws, f"C{r}", a, fmt="0.00")
        put(ws, f"D{r}", b, F_BOLD, fmt="0.00")
        put(ws, f"E{r}", cc, fmt="0.00")
        put(ws, f"F{r}", d, fmt="0.00")
        put(ws, f"H{r}", est)
        put(ws, f"I{r}", calc, align=WRAP)
        put(ws, f"J{r}", obs, align=WRAP)
    status_format(ws, "H45:H48")
    note(ws, "B49", "Los nuevos productos son ampliación de carga sobre el parque existente (no sustitución). No se suma carga si el activo ya está contado en el parque.")

    section(ws, 51, "5. Turnos, especialidades y simultaneidad de averías", 10)
    header_row(ws, 52, 1, ["", "Especialidad", "Activos con tecnología relacionada (informativo)", "FTE por especialidad", "", "", "", "Estado", "Tecnologías de 03 asociadas", "Observación"])
    esp = [
        ("Mecánica", '=COUNTIF(S_TECNO,"Hornos / TT")+COUNTIF(S_TECNO,"Fundición")+COUNTIF(S_TECNO,"Lavado")+COUNTIF(S_TECNO,"Auxiliares")+COUNTIF(S_TECNO,"Manutención / transporte")', "Hornos/TT, Fundición, Lavado, Auxiliares, Manutención"),
        ("Electricidad", '=COUNTIF(S_TECNO,"Instalaciones / servicios")', "Instalaciones / servicios"),
        ("Automatización y PLC", '=COUNTIF(S_TECNO,"Automatización / PLC (estación)")', "Automatización / PLC (estación)"),
        ("Robótica", '=COUNTIF(S_TECNO,"Robótica")', "Robótica"),
        ("CNC", '=COUNTIF(S_TECNO,"Mecanizado CNC")', "Mecanizado CNC"),
        ("Matricería", '=COUNTIFS(S_GRUPO,"TALLER Y MANTENIMIENTO")', "Taller y mantenimiento (incl. moldes)"),
        ("Visión y medición", '=COUNTIF(S_TECNO,"Visión / medición")', "Visión / medición"),
        ("Instalaciones y servicios", '=COUNTIF(S_GRUPO,"INSTALACIONES Y SERVICIOS")', "Grupo instalaciones y servicios"),
    ]
    for i, (e, f, t) in enumerate(esp):
        r = 53 + i
        put(ws, f"B{r}", e, F_BOLD)
        put(ws, f"C{r}", f, fmt="int")
        put(ws, f"D{r}", "PENDIENTE", fill_="pendiente")
        put(ws, f"H{r}", "PENDIENTE DE ASIGNACIÓN POR ESPECIALIDAD")
        put(ws, f"I{r}", t, align=WRAP)
        put(ws, f"J{r}", "No se inventan porcentajes: falta plantilla por especialidad y turno." if i == 0 else "")
    status_format(ws, "H53:H60")
    av = c.M["av"]
    per = av[(av["FECHA"] >= "2025-01-01")]
    diario = per.groupby(per["FECHA"].dt.date).size()
    diario_paro = per[per["MIN"] > 0].groupby(per[per["MIN"] > 0]["FECHA"].dt.date).size()
    r = 62
    header_row(ws, r, 1, ["", "Indicador de simultaneidad (SAP averías 2025→)", "Valor", "Unidad", "", "", "", "Estado", "Cálculo", "Observación"])
    sim = [
        ("Turnos de mantenimiento", "DATO NO DISPONIBLE", "", "PENDIENTE", "", "Necesario para convertir FTE en personas por turno."),
        ("Averías por día natural con registro (media)", round(float(diario.mean()), 1), "avisos/día", "DATO FUENTE", "media diaria", "Fuera de montaje (el export excluye montaje y TPM)."),
        ("Averías por día – percentil 95", round(float(diario.quantile(0.95)), 1), "avisos/día", "DATO FUENTE", "P95 diario", "Referencia para dimensionar respuesta simultánea."),
        ("Averías con paro > 0 por día (media)", round(float(diario_paro.mean()), 1), "avisos/día", "DATO FUENTE", "media diaria", ""),
        ("Máximo de averías en un día", int(diario.max()), "avisos/día", "DATO FUENTE", "máximo", "No se convierte a FTE: se cubre en la cobertura organizativa."),
    ]
    for i, (lbl, v, u, est, calc, obs) in enumerate(sim):
        rr = r + 1 + i
        put(ws, f"B{rr}", lbl)
        put(ws, f"C{rr}", v, fill_="dato" if not isinstance(v, str) else "pendiente", fmt="0.0")
        put(ws, f"D{rr}", u)
        put(ws, f"H{rr}", est)
        put(ws, f"I{rr}", calc)
        put(ws, f"J{rr}", obs, align=WRAP)
    status_format(ws, f"H{r + 1}:H{r + 5}")
    ws.freeze_panes = "C7"


# ======================================================================================
# 09_MEB_ECO
# ======================================================================================
MEB_GRUPOS = [("KMM", "KMM (MEB)"), ("EBOX", "EBOX (MEB)"), ("Conjunto Diferencial", "CONJ. DIFERENCIAL (MEB)"),
              ("Manguetas", "MANGUETAS (MEB)"), ("Aluminio GGH/MGH", "ALUMINIO GGH/MGH (MEB)")]


def hoja_meb(c: Ctx):
    ws = c.ws["09_MEB_ECO"]
    title(ws, "09 · NUEVOS PRODUCTOS – MEB ECO",
          "Ampliación de carga sobre el parque existente (no sustitución). Se cruza con SAP para no sumar dos veces un activo ya contado.")
    widths(ws, {"A": 22, "B": 22, "C": 13, "D": 30, "E": 12, "F": 12, "G": 12, "H": 12, "I": 13, "J": 13, "K": 13, "L": 30, "M": 44})
    of = c.M["meb_of"]
    section(ws, 5, "1. Activos críticos MEB ECO confirmados por negocio (MAESTRO_FABRICA_FINAL · 02_MEB_ECO_Oficial)", 13)
    header_row(ws, 6, 1, ["Producto", "Proceso", "Criticidad", "Tipo de activo crítico", "Total crítico confirmado", "Identificados en ficheros",
                          "Pendientes inventario exacto", "Fila fuente", "", "", "", "Estado", "Trazabilidad"])
    r = 7
    for _, x in of.iterrows():
        put(ws, f"A{r}", x["Producto"], F_BOLD)
        put(ws, f"B{r}", x["Proceso"])
        put(ws, f"C{r}", "A (A1)", F_BOLD, "A")
        ws[f"C{r}"].font = F_HEAD
        put(ws, f"D{r}", x["Tipo activo crítico"], align=WRAP)
        put(ws, f"E{r}", int(x["Total crítico confirmado"]), fill_="dato", fmt="int")
        put(ws, f"F{r}", int(x["Identificados en ficheros"]), fill_="dato", fmt="int")
        put(ws, f"G{r}", int(x["Pendientes inventario exacto"]), fill_="dato", fmt="int")
        put(ws, f"H{r}", int(x["FILA_ORIGEN"]))
        put(ws, f"L{r}", "PENDIENTE DE VALIDACIÓN" if int(x["Pendientes inventario exacto"]) > 0 else "OK")
        put(ws, f"M{r}", x["Estado trazabilidad"], align=WRAP)
        r += 1
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "EFG":
        put(ws, f"{col}{r}", f"=SUM({col}7:{col}{r - 1})", F_BOLD, "total", "int")
    put(ws, f"L{r}", f'=IF(E{r}=75,"OK","REVISAR")', F_BOLD)
    put(ws, f"M{r}", "Total confirmado 75 = 56 estaciones de montaje + 19 máquinas de mecanizado (4 Reishauer CDIF + 15 Aluminio GGH/MGH).")
    rt = r
    r += 1
    put(ws, f"A{r}", "de ellos montaje", F_BOLD)
    put(ws, f"E{r}", f'=SUMIFS(E7:E{rt - 1},B7:B{rt - 1},"Montaje")', fmt="int")
    r += 1
    put(ws, f"A{r}", "de ellos mecanizado", F_BOLD)
    put(ws, f"E{r}", f'=SUMIFS(E7:E{rt - 1},B7:B{rt - 1},"Mecanizado")', fmt="int")
    define(c.wb, "K_MEB_TOTAL", "09_MEB_ECO", f"$E${rt}")
    status_format(ws, f"L7:L{rt}")
    r += 2

    section(ws, r, "2. Cruce con el inventario SAP y el modelo: ¿incluido, parcial, alta futura o doble conteo?", 13)
    header_row(ws, r + 1, 1, ["Producto", "Grupo en el modelo", "Críticos oficiales", "Estaciones / máquinas críticas identificadas en el modelo",
                              "Diferencia", "Registros SAP del grupo", "En parque principal", "Excluidos / componentes", "Activos A en el modelo",
                              "FTE técnico del grupo", "FTE presentado (diap. 7)", "Situación", "Lectura"])
    r += 2
    r0 = r
    for prod, grp in MEB_GRUPOS:
        put(ws, f"A{r}", prod, F_BOLD)
        put(ws, f"B{r}", grp)
        put(ws, f"C{r}", f'=SUMIFS($E$7:$E${rt - 1},$A$7:$A${rt - 1},A{r})', fill_="oficial", fmt="int")
        put(ws, f"D{r}", f'=COUNTIFS(S_GRUPO,B{r},S_MONT,"S")+COUNTIFS(S_GRUPO,B{r},S_MEB,"S")', F_BOLD, fmt="int")
        put(ws, f"E{r}", f"=D{r}-C{r}", fmt="int")
        put(ws, f"F{r}", f'=COUNTIF(M_GRUPO,B{r})', fmt="int")
        put(ws, f"G{r}", f'=COUNTIFS(M_GRUPO,B{r},M_PARQUE,"S")', fmt="int")
        put(ws, f"H{r}", f'=COUNTIFS(M_GRUPO,B{r},M_PARQUE,"N")', fmt="int")
        put(ws, f"I{r}", f'=COUNTIFS(S_GRUPO,B{r},S_CLASE,"A")', fmt="int")
        put(ws, f"J{r}", f'=SUMIFS(S_FTE,S_GRUPO,B{r})', fmt="fte")
        ppt_lbl = {"KMM": "KMM", "EBOX": "EBOX", "Conjunto Diferencial": "CONJ. DIF.", "Manguetas": "MANGUETAS", "Aluminio GGH/MGH": "ALUMINIO"}[prod]
        put(ws, f"K{r}", f"=INDEX('05_FTE_POR_GRUPOS'!$B$7:$B$14,MATCH(\"{ppt_lbl}\",'05_FTE_POR_GRUPOS'!$A$7:$A$14,0))", fill_="oficial", fmt="fte")
        put(ws, f"L{r}", f'=IF(D{r}=0,"PENDIENTE – sin inventario identificado",IF(E{r}=0,"INCLUIDO EN PARQUE (cuadra)",IF(E{r}<0,"PARCIAL – faltan "&-E{r},"INCLUIDO – "&E{r}&" estaciones más que el oficial")))', F_BOLD)
        lect = {
            "KMM": "Estaciones de la línea KMM ya están en SAP (TG00-1-KMM0). No es alta futura: no sumar de nuevo.",
            "EBOX": "SAP detalla más estaciones que el oficial (incluye OP auxiliares). Validar cuáles son las 14 críticas.",
            "Conjunto Diferencial": "Estaciones AFO activas + 4 rectificadoras Reishauer localizadas en SAP (TG00-1-CDIF-MONT / CDIF-CONJ). La fuente decía 6 identificadas.",
            "Manguetas": "Puestos 'S/R' son standby/reparación (no críticos). Validar las 17 críticas frente a 13 estaciones activas.",
            "Aluminio GGH/MGH": "Ninguna de las 15 identificada por inventario. Candidatos en SAP (sección 4). No se añaden filas ficticias.",
        }[prod]
        put(ws, f"M{r}", lect, align=WRAP)
        ws.row_dimensions[r].height = 42
        r += 1
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "CDEFGHIJK":
        put(ws, f"{col}{r}", f"=SUM({col}{r0}:{col}{r - 1})", F_BOLD, "total", "int" if col not in "JK" else "fte")
    status_format(ws, f"L{r0}:L{r}")
    define(c.wb, "K_MEB_IDENT", "09_MEB_ECO", f"$D${r}")
    r += 2

    section(ws, r, "3. Registros SAP de los grupos MEB por tipo de activo (crítico principal, estación, robot, periférico, manutención…)", 13)
    tipos = [t for t in c.inv[c.inv["PROGRAMA"] == "MEB ECO"]["TIPO_ACTIVO"].value_counts().index if t]
    header_row(ws, r + 1, 1, ["Tipo de activo"] + [g for _, g in MEB_GRUPOS] + ["TOTAL"], height=36)
    r += 2
    r1 = r
    for t in tipos:
        put(ws, f"A{r}", t, F_BOLD)
        for j, (_, g) in enumerate(MEB_GRUPOS):
            put(ws, f"{'BCDEF'[j]}{r}", f'=COUNTIFS(M_GRUPO,"{g}",M_TIPO,$A{r})', fmt="int")
        put(ws, f"G{r}", f"=SUM(B{r}:F{r})", F_BOLD, fmt="int")
        r += 1
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "BCDEFG":
        put(ws, f"{col}{r}", f"=SUM({col}{r1}:{col}{r - 1})", F_BOLD, "total", "int")
    r += 2

    section(ws, r, "4. Candidatos para los 15 críticos de Aluminio GGH/MGH y Reishauer CDIF (PENDIENTE DE VALIDACIÓN)", 13)
    header_row(ws, r + 1, 1, ["ID", "Descripción", "Grupo", "Ubicación", "Tipo SAP", "Cuello botella", "Clase modelo", "Averías", "Horas paro", "", "", "Situación", "Motivo de candidatura"])
    inv = c.inv
    cand = inv[(inv["EN_PARQUE"] == "S") & (
        ((inv["GRUPO_MANT"] == "ALUMINIO GGH/MGH (MEB)") & inv["TIPOLOGIA"].isin(["T03", "T05", "T07"])) |
        (inv["CB_DETALLE"].fillna("").str.contains("E-PRODUCTS-AL")) |
        (inv["MEB_CRIT"] == "S"))]
    r += 2
    for _, x in cand.sort_values(["GRUPO_MANT", "ID"]).iterrows():
        put(ws, f"A{r}", x["ID"])
        put(ws, f"B{r}", x["DESC"])
        put(ws, f"C{r}", x["GRUPO_MANT"])
        put(ws, f"D{r}", x["UBIC"])
        put(ws, f"E{r}", x["TP"])
        put(ws, f"F{r}", x["CB"], align=CENTER)
        put(ws, f"G{r}", f'=INDEX(S_ETIQ,MATCH(A{r},S_ID,0))')
        put(ws, f"H{r}", int(x["AV_N"]), fmt="int")
        put(ws, f"I{r}", float(x["AV_H"]), fmt="h")
        put(ws, f"L{r}", "PENDIENTE DE VALIDACIÓN")
        mot = []
        if x["MEB_CRIT"] == "S":
            mot.append("Rectificadora Reishauer en TG00-1-CDIF-CONJ")
        if "E-PRODUCTS-AL" in _s(x["CB_DETALLE"]):
            mot.append("Cuello de botella E-PRODUCTS-AL (" + _s(x["CB_DETALLE"])[:40] + ")")
        if x["GRUPO_MANT"] == "ALUMINIO GGH/MGH (MEB)":
            mot.append("Máquina principal en sección aluminio GGH/MGH")
        put(ws, f"M{r}", "; ".join(mot), align=WRAP)
        r += 1
    class_format(ws, f"G{r - len(cand)}:G{r}")
    status_format(ws, f"L{r - len(cand)}:L{r}")
    r += 1
    section(ws, r, "5. Detalle MEB del fichero de montaje (MAESTRO_FABRICA_FINAL · 03_MEB_ECO_Detalle) cruzado con el modelo", 13)
    header_row(ws, r + 1, 1, ["Inventario", "Producto", "Proceso", "Fabricante", "Estado fuente", "Fila fuente", "ID en el modelo", "", "", "", "", "Situación", "Comentario fuente"])
    r += 2
    det = c.M["meb_det"]
    for _, x in det.iterrows():
        put(ws, f"A{r}", _s(x["Inventario / Activo"]))
        put(ws, f"B{r}", _s(x["Producto"]))
        put(ws, f"C{r}", _s(x["Proceso"]))
        put(ws, f"D{r}", _s(x["Fabricante"]))
        put(ws, f"E{r}", _s(x["Estado"]))
        put(ws, f"F{r}", int(x["FILA_ORIGEN"]))
        put(ws, f"G{r}", _s(x["ID_MODELO"]))
        sit = ("INCLUIDO EN PARQUE (SAP)" if _s(x["ID_MODELO"]) else "PENDIENTE – no localizado en SAP")
        if "Standby" in _s(x["Estado"]):
            sit = "STANDBY / RESERVA – " + sit
        put(ws, f"L{r}", sit)
        put(ws, f"M{r}", _s(x["Comentario / carga"]), align=WRAP)
        r += 1
    status_format(ws, f"L{r - len(det)}:L{r}")
    ws.freeze_panes = "B5"


# ======================================================================================
# 18 / 19 / 20 – registros de cruce
# ======================================================================================
def hoja_reg_averias(c: Ctx):
    ws = c.ws["18_REG_AVERIAS"]
    title(ws, "18 · REGISTROS DE AVERÍA SAP (2025→actual) CON SU ACTIVO ASIGNADO",
          "Fuente: SAP.EXPORT.AVERIAS.LISTADO.xlsx · hoja AVERIAS. 'Tiempo de Paro' en minutos (validado: horas × 60 del histórico de paros).")
    cols = [("Fila origen", 8), ("Fecha inicio avería", 12), ("Orden", 11), ("P (paro)", 6), ("Objeto técnico", 40), ("Descripción avería", 46),
            ("Ubicación", 32), ("Tiempo de paro (min)", 10), ("Horas", 8), ("ID asignado", 22), ("Cruce", 26), ("Grupo mantenimiento", 24), ("Clase del activo", 12)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    av = c.M["av"].sort_values("FECHA").reset_index(drop=True)
    gm = dict(zip(c.inv["ID"], c.inv["GRUPO_MANT"]))
    parque = set(c.S["ID"])
    for i, x in av.iterrows():
        r = 6 + i
        ws.cell(r, 1, int(x["FILA_ORIGEN"]))
        ws.cell(r, 2, x["FECHA"].date() if pd.notna(x["FECHA"]) else None).number_format = "dd/mm/yyyy"
        ws.cell(r, 3, int(x["Orden"]) if pd.notna(x["Orden"]) else None)
        ws.cell(r, 4, _s(x["P"]))
        ws.cell(r, 5, _s(x["Denominación de objeto técnico"]))
        ws.cell(r, 6, _s(x["Descripción"]))
        ws.cell(r, 7, _s(x["Denominación ubicación"]))
        ws.cell(r, 8, float(x["MIN"]))
        ws.cell(r, 9, f"=H{r}/60").number_format = "0.00"
        ws.cell(r, 10, _s(x["ID_ASIGNADO"]))
        ws.cell(r, 11, _s(x["CRUCE"]))
        ws.cell(r, 12, gm.get(x["ID_ASIGNADO"], "SIN ACTIVO"))
        ws.cell(r, 13, f'=IFERROR(INDEX(S_ETIQ,MATCH(J{r},S_ID,0)),"—")' if x["ID_ASIGNADO"] in parque else "—")
    last = 5 + len(av)
    ws.freeze_panes = "E6"
    ws.auto_filter.ref = f"A5:M{last}"
    for name, col in (("AV_MIN", "H"), ("AV_ID", "J"), ("AV_CRUCE", "K"), ("AV_GRUPO", "L"), ("AV_FECHA", "B"), ("AV_P", "D")):
        define(c.wb, name, "18_REG_AVERIAS", f"${col}$6:${col}${last}")
    class_format(ws, f"M6:M{last}")


def hoja_cruce_montaje(c: Ctx):
    ws = c.ws["19_CRUCE_MONTAJE"]
    title(ws, "19 · FICHERO DE MONTAJE (264 FILAS) CRUZADO CON SAP",
          "Fuente: MAESTRO_FABRICA_FINAL.xlsx · 04_Montaje_Completo. Clave de cruce: (inventario de línea, nº de estación); 90410 → 7016420001 y 90510 → 7016420002.")
    cols = [("Fila fuente", 8), ("Inventario", 20), ("Inventario antiguo", 26), ("Programa", 12), ("Producto", 18), ("Fabricante", 16),
            ("Estado fuente", 22), ("Comentario", 44), ("Clave (línea, estación)", 22), ("En SAP", 8), ("ID SAP", 34), ("Situación en el modelo", 40)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    m = c.M["mont"].reset_index(drop=True)
    ids_extra = set(c.inv["ID"])
    for i, x in m.iterrows():
        r = 6 + i
        vals = [int(x["FILA_ORIGEN"]), _s(x["Inventario"]), _s(x["Inventario antiguo"]), _s(x["Programa"]), _s(x["Producto"]),
                _s(x["Fabricante"]), _s(x["Estado"]), _s(x["Comentario"]), str(x["CLAVE_EST"]) if x["CLAVE_EST"] else "", x["EN_SAP"], _s(x["ID_SAP"])]
        for j, v in enumerate(vals):
            ws.cell(r, j + 1, v)
        if x["EN_SAP"] == "S":
            sit = "Contado una vez en el registro SAP"
        elif any(i.startswith("MONT-" + _s(x["Inventario"])) for i in ids_extra if _s(x["Inventario"])):
            sit = "Añadido como registro de montaje no localizado en SAP"
        else:
            sit = "Registro pendiente (sin inventario o datos comunes)"
        ws.cell(r, 12, sit)
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:L{5 + len(m)}"


def hoja_cruce_cuellos(c: Ctx):
    ws = c.ws["20_CRUCE_CUELLOS"]
    title(ws, "20 · CUELLOS DE BOTELLA (129) CRUZADOS CON SAP",
          "Fuente: CUELLOS_DE_BOTELLA.xlsx + mapeo de MAESTRO_INTEGRAL 05_Cuellos_Botella. Todo cuello confirmado = A (C1 = 4, C3 = 2).")
    cols = [("Fila fuente", 8), ("Modelo", 18), ("Línea", 18), ("Cuello de botella", 26), ("OEE", 8), ("Mapeado (INTEGRAL)", 10),
            ("Activo destino", 40), ("Método cruce", 14), ("ID asignado en el modelo", 24), ("Estado del cruce", 36)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    cb = c.M["cb"].reset_index(drop=True)
    for i, x in cb.iterrows():
        r = 6 + i
        vals = [int(x["FILA_ORIGEN"]), _s(x["MODELO"]), _s(x["LINEA"]), _s(x["C.BOTELLA"]), float(x["OEE"]), _s(x["MAPEADO"]),
                _s(x["ACTIVO_DESTINO"]), _s(x["METODO_CRUCE"]), _s(x["ID_ASIGNADO"]) or ("CB-%03d" % int(x["FILA_ORIGEN"])), _s(x["ESTADO_CRUCE"])]
        for j, v in enumerate(vals):
            ws.cell(r, j + 1, v)
        ws.cell(r, 5).number_format = "0.0%"
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:J{5 + len(cb)}"
    status_format(ws, f"J6:J{5 + len(cb)}")


# ======================================================================================
# 10_AVERIAS
# ======================================================================================
def hoja_averias(c: Ctx):
    ws = c.ws["10_AVERIAS"]
    title(ws, "10 · AVERÍAS, RECURRENCIA Y PARO (periodo principal 2025→actual)",
          "La ausencia de averías recientes no convierte un activo en poco crítico: C4 pesa como máximo 1,5 de 10 puntos.")
    widths(ws, {"A": 44, "B": 14, "C": 14, "D": 14, "E": 14, "F": 16, "G": 40, "H": 30})
    meta = c.M["av_meta"]
    av = c.M["av"]
    section(ws, 5, "1. Fichero y unidad de paro", 8)
    header_row(ws, 6, 1, ["Concepto", "Valor", "Unidad", "", "", "Estado", "Fuente / explicación", ""])
    rows = [
        ("Registros con orden en el export", "=COUNT(AV_MIN)", "registros", "DATO FUENTE", "SAP.EXPORT.AVERIAS.LISTADO.xlsx · AVERIAS"),
        ("Primera fecha de avería en las filas", av["FECHA"].min().date(), "fecha", "DATO FUENTE", ""),
        ("Última fecha de avería en las filas", av["FECHA"].max().date(), "fecha", "DATO FUENTE", ""),
        ("Suma de 'Tiempo de Paro' de las filas", "=SUM(AV_MIN)", "min", "DATO FUENTE", ""),
        ("Horas de paro de las filas", "=B10/60", "h", "DATO FUENTE", "min ÷ 60"),
        ("Fila 'Total' del export", meta["total_fila_min"], "min", "REVISAR",
         f"La fila Total (desde {meta['total_fila_desde']}) corresponde a la vista completa 2022→: 33.758,6 h. Las filas exportadas solo cubren 2025→."),
        ("Filtro aplicado en el origen (Power BI)", (meta["filtro"] or "").replace("\n", " · "), "", "OK", "Excluye TPM (G9RTPM) y las zonas MONTAJE_MQ200 / MONTAJE_MQ281."),
        ("Unidad de 'Tiempo de Paro'", "minutos", "", "OK", "Comprobado: 05_Paros_Historico 'Minutos paro' = 'Horas paro' × 60 (41_Comprobaciones fila 8)."),
        ("Unidad de 'CP' en data.1 / data.2 (montaje)", "minutos (supuesto)", "", "PENDIENTE DE VALIDACIÓN", "Mediana 1–4 por registro; existen valores atípicos (6.145, 7.185)."),
    ]
    for i, (a, b, u, e, f) in enumerate(rows):
        r = 7 + i
        put(ws, f"A{r}", a)
        put(ws, f"B{r}", b, fill_="dato", fmt="dd/mm/yyyy" if u == "fecha" else "#,##0.0")
        put(ws, f"C{r}", u)
        put(ws, f"F{r}", e)
        put(ws, f"G{r}", f, align=WRAP)
        ws.merge_cells(f"G{r}:H{r}")
    status_format(ws, "F7:F15")
    define(c.wb, "K_AV_H", "10_AVERIAS", "$B$11")

    section(ws, 17, "2. Discrepancia: 'Horas avería 2025→actual = 29.627 h' (DIMENSIONAMIENTO 00_Resumen B9 · 05_Datos_2025_actual B7)", 8)
    header_row(ws, 18, 1, ["Concepto", "Presentado / fuente", "Recalculado", "Diferencia", "", "Estado", "Explicación", ""])
    put(ws, "A19", "Horas de paro 2025→actual")
    put(ws, "B19", 29627.2, fill_="dato", fmt="#,##0.0")
    put(ws, "C19", "=K_AV_H", F_BOLD, fmt="#,##0.0")
    put(ws, "D19", "=C19-B19", fmt="#,##0.0")
    put(ws, "F19", "DIFERENCIA", F_BOLD)
    put(ws, "G19", "El filtro 'campo Años contiene 2025/2026' suma el histórico COMPLETO (2022–2026) de las máquinas con actividad en 2025/26. "
        "No se usó en el PPT final; solo afecta al criterio C4.", align=WRAP)
    ws.merge_cells("G19:H19")
    ws.row_dimensions[19].height = 45
    status_format(ws, "F19")

    section(ws, 21, "3. Evolución mensual", 8)
    header_row(ws, 22, 1, ["Mes", "Registros", "Con paro (P = X)", "Horas de paro", "Horas por registro", "", "", ""])
    av_m = av.copy()
    av_m["MES"] = av_m["FECHA"].dt.to_period("M")
    g = av_m.groupby("MES").agg(N=("MIN", "size"), PX=("P", lambda s: int((s.astype(str) == "X").sum())), H=("MIN", lambda s: s.sum() / 60))
    r = 23
    for mes, x in g.iterrows():
        put(ws, f"A{r}", str(mes))
        put(ws, f"B{r}", int(x["N"]), fill_="dato", fmt="int")
        put(ws, f"C{r}", int(x["PX"]), fill_="dato", fmt="int")
        put(ws, f"D{r}", round(float(x["H"]), 1), fill_="dato", fmt="#,##0.0")
        put(ws, f"E{r}", f"=D{r}/B{r}", fmt="0.00")
        r += 1
    put(ws, f"A{r}", "TOTAL", F_BOLD, "total")
    for col in "BCD":
        put(ws, f"{col}{r}", f"=SUM({col}23:{col}{r - 1})", F_BOLD, "total", "#,##0.0" if col == "D" else "int")
    r += 2

    section(ws, r, "4. Cruce de averías con el inventario y reparto por grupo de mantenimiento (fórmulas sobre 18_REG_AVERIAS)", 8)
    header_row(ws, r + 1, 1, ["Grupo / tipo de cruce", "Registros", "Horas de paro", "% horas", "", "", "", ""])
    r += 2
    r0 = r
    for k in ["Exacto por descripción", "Ambiguo", "SIN ACTIVO SAP"]:
        put(ws, f"A{r}", "Cruce: " + k, F_BOLD)
        put(ws, f"B{r}", f'=COUNTIF(AV_CRUCE,"{k}*")', fmt="int")
        put(ws, f"C{r}", f'=SUMIF(AV_CRUCE,"{k}*",AV_MIN)/60', fmt="#,##0.0")
        put(ws, f"D{r}", f"=C{r}/K_AV_H", fmt="pct")
        r += 1
    r += 1
    grupos = list(av.assign(G=av["ID_ASIGNADO"].map(dict(zip(c.inv["ID"], c.inv["GRUPO_MANT"]))).fillna("SIN ACTIVO"))
                  .groupby("G")["MIN"].sum().sort_values(ascending=False).index)
    for gname in grupos:
        put(ws, f"A{r}", gname)
        put(ws, f"B{r}", f'=COUNTIF(AV_GRUPO,A{r})', fmt="int")
        put(ws, f"C{r}", f'=SUMIF(AV_GRUPO,A{r},AV_MIN)/60', fmt="#,##0.0")
        put(ws, f"D{r}", f"=C{r}/K_AV_H", fmt="pct")
        r += 1
    r += 1

    section(ws, r, "5. Top 30 activos por horas de paro 2025→actual", 8)
    header_row(ws, r + 1, 1, ["Activo", "ID", "Averías", "Horas de paro", "Parada máx. (min)", "Clase en el modelo", "Grupo", "Cuello de botella"])
    r += 2
    top = c.S.sort_values("AV_H", ascending=False).head(30)
    r1 = r
    for _, x in top.iterrows():
        put(ws, f"A{r}", x["DESC"])
        put(ws, f"B{r}", x["ID"])
        put(ws, f"C{r}", int(x["AV_N"]), fmt="int")
        put(ws, f"D{r}", float(x["AV_H"]), fmt="#,##0.0")
        put(ws, f"E{r}", float(x["AV_MAX"]), fmt="int")
        put(ws, f"F{r}", f"=INDEX(S_ETIQ,MATCH(B{r},S_ID,0))")
        put(ws, f"G{r}", x["GRUPO_MANT"])
        put(ws, f"H{r}", x["CB"], align=CENTER)
        r += 1
    class_format(ws, f"F{r1}:F{r}")
    r += 1

    section(ws, r, "6. Paros de las líneas de montaje 2025 (data.1 = MQ200 Línea 1 · data.2 = MQ281 Línea 3)", 8)
    header_row(ws, r + 1, 1, ["Estación (EST)", "Línea", "Registros", "CP total (min)", "ID asignado", "", "Archivo", ""])
    logs = c.M["logs"]
    gl = logs.groupby(["EST", "LINEA", "ARCHIVO"]).agg(N=("CP", "count"), CP=("CP", "sum"), ID=("ID_ASIGNADO", "first")).reset_index().sort_values("CP", ascending=False).head(30)
    r += 2
    for _, x in gl.iterrows():
        put(ws, f"A{r}", x["EST"])
        put(ws, f"B{r}", x["LINEA"])
        put(ws, f"C{r}", int(x["N"]), fmt="int")
        put(ws, f"D{r}", float(x["CP"]), fmt="#,##0")
        put(ws, f"E{r}", x["ID"] or "sin estación SAP")
        put(ws, f"G{r}", x["ARCHIVO"])
        r += 1
    note(ws, f"A{r + 1}", "Los paros de montaje solo se usan para C4 de estaciones que ya son A por la regla de montaje principal; no cambian su clase.")
    ws.freeze_panes = "B5"


# ======================================================================================
# 11_TRAZABILIDAD
# ======================================================================================
def hoja_trazabilidad(c: Ctx):
    ws = c.ws["11_TRAZABILIDAD"]
    title(ws, "11 · TRAZABILIDAD: POWERPOINT → RESULTADO → CÁLCULO → GRUPO → ACTIVO → CRITERIO → DATO → ARCHIVO",
          "Cada KPI del PPT con su valor oficial, la cifra que reproduce el modelo, la fórmula, el supuesto y la fuente. Cuando no se puede completar: TRAZABILIDAD INCOMPLETA.")
    cols = [("#", 5), ("KPI", 34), ("Valor oficial PPT", 13), ("Diap.", 7), ("Valor modelo", 13), ("Diferencia", 11), ("Archivo de origen", 34),
            ("Hoja", 24), ("Registro / celda", 22), ("Fórmula", 40), ("Supuesto utilizado", 30), ("Estado", 24), ("Observaciones", 60), ("Ir a", 14)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    DIM = "DIMENSIONAMIENTO_MANTENIMIENTO_DATOS_FINAL.xlsx"
    FAB = "MAESTRO_FABRICA_FINAL.xlsx"
    of = O.OF
    G = "'05_FTE_POR_GRUPOS'!"
    K = [
        ("Parque total de activos", of["PARQUE"][0], "4 · 8", "=K_PARQUE_CLASES", FAB + " · " + DIM, "07_Contexto_Global · 02_Parque_Criticidad", "B4:B12 · B6:B15",
         "= 430+92+264+58+290+74+355+386+94", "P_PARQUE_OF", "auto", "Suma de familias globales, no inventario fila a fila (SAP depurado 2.078; parque del modelo en 04).", "04_FTE_CRITICIDAD"),
        ("Activos clase A", of["N_A"][0], "4 · 8", "=P_N_A", FAB, "01_Maestro_Fabrica", "Bloques CB, A2, Montaje, MEB",
         "= 129 + 55 + 185 + 43 + 6 filas", "P_N_A", "RECONSTRUIDO con observaciones", "Recuento de filas con granularidad mixta (ver 04 §3 y 17).", "04_FTE_CRITICIDAD"),
        ("Activos clase B", of["N_B"][0], "4", "=P_N_B", FAB, "01_Maestro_Fabrica", "Histórico paros 'B - Alta carga' = 119",
         "Declarado", "P_N_B", "TRAZABILIDAD INCOMPLETA", "+156 sin registro de origen.", "17_DIFERENCIAS"),
        ("Activos clase C", of["N_C"][0], "4", "=P_N_C", FAB, "01_Maestro_Fabrica", "C - Media 153 + C/Pendiente 36",
         "Declarado", "P_N_C", "TRAZABILIDAD INCOMPLETA", "Valor redondo declarado.", "17_DIFERENCIAS"),
        ("Activos clase D", of["N_D"][0], "4", "=P_N_D", FAB, "01_Maestro_Fabrica", "D - Baja/ruta 920 + 397 no clasificados",
         "Declarado", "P_N_D", "TRAZABILIDAD INCOMPLETA", "Reasignación declarada coherente con 'D solo con bajo impacto demostrado'.", "17_DIFERENCIAS"),
        ("Ratios A/B/C/D (activos por FTE)", "12 / 20 / 35 / 60", "3 · 4", '=P_RATIO_A&" / "&P_RATIO_B&" / "&P_RATIO_C&" / "&P_RATIO_D', O.PPT + " · " + DIM, "02_Parque_Criticidad", "C19:C22",
         "", "P_RATIO_A…D", "OK", "Criterio interno de cartera; extremo superior de los rangos de la diapositiva 3.", "12_SUPUESTOS"),
        ("FTE base clase A", of["FTE_A"][0], "4", "='04_FTE_CRITICIDAD'!D7", DIM, "02_Parque_Criticidad", "D19", "= 418 ÷ 12", "P_N_A, P_RATIO_A", "auto", "", "04_FTE_CRITICIDAD"),
        ("FTE base clase B", of["FTE_B"][0], "4", "='04_FTE_CRITICIDAD'!D8", DIM, "02_Parque_Criticidad", "D20", "= 275 ÷ 20", "P_N_B, P_RATIO_B", "auto", "", "04_FTE_CRITICIDAD"),
        ("FTE base clase C", of["FTE_C"][0], "4", "='04_FTE_CRITICIDAD'!D9", DIM, "02_Parque_Criticidad", "D21", "= 1.200 ÷ 35", "P_N_C, P_RATIO_C", "auto", "", "04_FTE_CRITICIDAD"),
        ("FTE base clase D", of["FTE_D"][0], "4", "='04_FTE_CRITICIDAD'!D10", DIM, "02_Parque_Criticidad", "D22", "= 150 ÷ 60", "P_N_D, P_RATIO_D", "auto", "", "04_FTE_CRITICIDAD"),
        ("FTE base por criticidad", of["FTE_BASE"][0], "4 · 5", "=K_FTE_BASE", DIM, "02_Parque_Criticidad", "D23", "= Σ FTE clases", "Ratios oficiales", "auto", "Diap. 5 muestra 85 (redondeo).", "04_FTE_CRITICIDAD"),
        ("Preventivos (PM + avisos)", of["W_PM"][0], "5", "=K_PM", DIM, "04_Calculo_FTE · 01_Supuestos_Fuentes", "B7 · B12:B13", "= 60.000 ÷ 6.300", "P_PRODUCT (hipótesis)", "auto",
         "Productividad sin horas reales (~13,8 min/unidad). TPM no sumado.", "08_PM_AVISOS"),
        ("Formación", of["W_FORM"][0], "5", "=K_FORM", DIM, "04_Calculo_FTE", "B8", "= 80 × 55 ÷ 1.450", "P_FORM_H, P_HORAS_PROD", "auto", "Exacto 3,03.", "07_SABADOS_FORMACION"),
        ("Sábados / libranzas", of["W_SAB"][0], "5", "=K_SAB", DIM, "04_Calculo_FTE", "B9", "= 10 × 45 ÷ 220", "P_PERS_SAB, P_SABADOS, P_JORNADAS", "auto", "Exacto 2,05. Descanso compensatorio.", "07_SABADOS_FORMACION"),
        ("Cobertura organizativa mínima", of["W_COB_MIN"][0], "5", "=K_COB_MIN", DIM, "04_Calculo_FTE", "B10", "= base × 15 %", "P_COB_BASE, P_COB_MIN", "auto",
         "+13 = 15 % solo sobre la base (selector 3). Con el método de la diap. 2 (subtotal) = 15,0.", "06_CAPACIDAD_PERSONAL"),
        ("Cobertura organizativa máxima", of["W_COB_MAX"][0], "5", "=K_COB_MAX", DIM, "04_Calculo_FTE", "B11", "= base × 25 %", "P_COB_BASE, P_COB_MAX", "auto",
         "+21 = 25 % solo sobre la base (selector 3).", "06_CAPACIDAD_PERSONAL"),
        ("Necesidad total mínima (diap. 7)", of["S7_REF_MIN"][0], "7 · 5", "=K_NEC_MIN", DIM, "04_Calculo_FTE", "B12", "= subtotal + cobertura mín.", "P_COB_BASE", "auto",
         "Diap. 5 muestra también '113–121' (otra base de cobertura).", "06_CAPACIDAD_PERSONAL"),
        ("Necesidad total máxima (diap. 7)", of["S7_REF_MAX"][0], "7 · 5", "=K_NEC_MAX", DIM, "04_Calculo_FTE", "B13", "= subtotal + cobertura máx.", "P_COB_BASE", "auto", "", "06_CAPACIDAD_PERSONAL"),
        ("Necesidad total mínima (diap. 5 waterfall)", of["W_NEC_MIN"][0], "5", "='06_CAPACIDAD_PERSONAL'!F28", O.PPT, "Diap. 5", "113–121", "= subtotal + 15 % × base", "Selector 3",
         "auto", "Reproducida con la cobertura aplicada solo a la base.", "06_CAPACIDAD_PERSONAL"),
        ("Plantilla actual", of["PLANTILLA"][0], "5 · 7", "=P_PLANTILLA", DIM, "01_Supuestos_Fuentes", "B6", "", "P_PLANTILLA", "PENDIENTE DE VALIDACIÓN",
         "'~80 personas' facilitado por el usuario; sin fichero de organización.", "12_SUPUESTOS"),
        ("GAP (diap. 7)", of["S7_GAP"][0], "7", "=K_GAP_MIN", DIM, "04_Calculo_FTE", "B15 = 34,21", "= necesidad mín. − plantilla", "P_COB_BASE", "auto",
         "+34 coincide con el Excel de dimensionamiento (base + PM); con el método subtotal = 35,0.", "06_CAPACIDAD_PERSONAL"),
        ("GAP (diap. 5)", of["S5_GAP"][0], "5", "='06_CAPACIDAD_PERSONAL'!F28-P_PLANTILLA", O.PPT, "Diap. 5", "", "= 113 − 80", "Selector 3", "auto", "", "06_CAPACIDAD_PERSONAL"),
        ("Cobertura actual estimada", of["COB_ACTUAL"][0], "5 · 7", "=K_COB_ACTUAL", O.PPT, "Diap. 5 y 7", "", "= plantilla ÷ subtotal A+B+C+D", "", "auto", "80 ÷ 99,97 = 80,0 %.", "06_CAPACIDAD_PERSONAL"),
    ]
    filas_g = {"MQ200": 21, "MQ281": 22, "EBOX": 23, "MANGUETAS": 24, "KMM": 25, "CONJ. DIF.": 26, "ALUMINIO": 27}
    for g, v in O.GRUPOS_S7:
        if g in filas_g:
            K.append((f"FTE grupo {g}", v, "7", f"={G}G{filas_g[g]}", DIM, "03_Secciones", f"fila {filas_g[g] - 15}", "= A/12 + B/20 + C/35 + D/60 + A_MEB/12",
                      "Ratios oficiales", "auto", "Reparto marcado 'PENDIENTE REALIZAR' en el PPT.", "05_FTE_POR_GRUPOS"))
        else:
            K.append((f"FTE grupo {g}", v, "7", "", O.PPT, "Tabla incrustada diap. 7", "fórmula D7-D6 sobre celdas vacías", "", "", "TRAZABILIDAD INCOMPLETA",
                      "41,8 FTE de la base no están repartidos en el gráfico.", "05_FTE_POR_GRUPOS"))
    K += [
        ("Puntos de preventivo 'yoY'", of["S8_PUNTOS_PM"][0], "8", "=P_PUNTOS_PM", O.PPT_CTX, "Diap. 1", "'50.000 puntos'", "", "P_PUNTOS_PM", "auto",
         "El cálculo de preventivos usa 50.000 puntos; el 15.000 de la diap. 8 no tiene fuente.", "08_PM_AVISOS"),
        ("Activos críticos MEB ECO (fuente)", 75, "—", "=K_MEB_TOTAL", FAB, "02_MEB_ECO_Oficial", "E4:E9", "= 15+14+10+4+17+15", "", "auto",
         "No aparece en el PPT V2 como cifra; base de los grupos MEB.", "09_MEB_ECO"),
    ]
    for i, (kpi, val, diap, modelo, arch, hoja, reg, fx, sup, est, obs, ira) in enumerate(K):
        r = 6 + i
        put(ws, f"A{r}", i + 1, align=CENTER)
        put(ws, f"B{r}", kpi, F_BOLD, align=WRAP)
        num = isinstance(val, (int, float))
        put(ws, f"C{r}", val, F_BOLD, "oficial", ("pct" if (num and val < 1 and val > 0) else ("0.00" if num and not float(val).is_integer() else "#,##0")) if num else None)
        put(ws, f"D{r}", diap, align=CENTER)
        put(ws, f"E{r}", modelo or "—", F_BOLD, None, ("pct" if (num and 0 < val < 1) else "0.00") if num else None)
        if num and modelo:
            put(ws, f"F{r}", f"=E{r}-C{r}", fmt="0.0%" if 0 < val < 1 else "0.00")
        else:
            put(ws, f"F{r}", "")
        put(ws, f"G{r}", arch, align=WRAP)
        put(ws, f"H{r}", hoja, align=WRAP)
        put(ws, f"I{r}", reg, align=WRAP)
        put(ws, f"J{r}", fx, align=WRAP)
        put(ws, f"K{r}", sup, align=WRAP)
        if est == "auto":
            tol = "0.01" if (num and 0 < val < 1) else "P_TOL_RED"
            put(ws, f"L{r}", f'=IF(ABS(F{r})<={tol},"OK","DIFERENCIA – ver 17")', F_BOLD)
        else:
            put(ws, f"L{r}", est, F_BOLD)
        put(ws, f"M{r}", obs, align=WRAP)
        put(ws, f"N{r}", hl(f"#'{ira}'!A1", "→ " + ira), F_LINK)
        ws.row_dimensions[r].height = 30
    last = 5 + len(K)
    status_format(ws, f"L6:L{last}")
    define(c.wb, "T_ESTADOS", "11_TRAZABILIDAD", f"$L$6:$L${last}")
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:N{last}"

    # Recorrido de ejemplo para un activo
    r = last + 3
    section(ws, r, "Recorrido completo de ejemplo (un cuello de botella con averías recurrentes)", 14)
    S = c.S_sorted
    cand = S[(S["CB"] == "S") & (S["AV_N"] > 0)].sort_values("AV_H", ascending=False)
    if len(cand):
        x = cand.iloc[0]
        srow = 6 + int(S.index[S["ID"] == x["ID"]][0])
        mrow = c.maestro_row.get(x["ID"])
        pasos = [
            ("POWERPOINT", f"Diapositiva 4: A = 418 activos → 418 ÷ 12 = 34,83 FTE", "→ 04_FTE_CRITICIDAD", "#'04_FTE_CRITICIDAD'!A6"),
            ("RESULTADO OFICIAL", "FTE base 85,37 (04 D11) → necesidad 115–125 (06)", "→ 06_CAPACIDAD_PERSONAL", "#'06_CAPACIDAD_PERSONAL'!A6"),
            ("CÁLCULO", "Clase A × ratio 12 = 1/12 FTE por activo (03 col. AK)", "→ 03 fila " + str(srow), f"#'03_SCORE_CRITICIDAD'!A{srow}"),
            ("GRUPO", f"{x['GRUPO_MANT']} · {x['LINEA']} (05 §3)", "→ 05_FTE_POR_GRUPOS", "#'05_FTE_POR_GRUPOS'!A1"),
            ("ACTIVO", f"{x['ID']} · {x['DESC']}", "→ 02 fila " + str(mrow), f"#'02_MAESTRO_CRITICIDAD'!A{mrow}"),
            ("CRITERIOS", "C1–C5 en 03 (columnas V–Z) · regla forzada: " + x["REGLA_FORZADA"], "→ 03 fila " + str(srow), f"#'03_SCORE_CRITICIDAD'!V{srow}"),
            ("DATO ORIGINAL", f"Cuello: {x['CB_DETALLE'][:70]} · {int(x['AV_N'])} averías / {x['AV_H']:.1f} h", "→ 20 / 18", "#'20_CRUCE_CUELLOS'!A1"),
            ("ARCHIVO FUENTE", f"{x['FUENTE']} fila {x['FILA']} · CUELLOS_DE_BOTELLA.xlsx · SAP.EXPORT.AVERIAS.LISTADO.xlsx", "→ 14_FUENTES", "#'14_FUENTES'!A1"),
        ]
        for i, (paso, txt, lk, href) in enumerate(pasos):
            rr = r + 1 + i
            put(ws, f"B{rr}", ("↓ " if i else "") + paso, F_BOLD, "section")
            put(ws, f"C{rr}", txt)
            ws.merge_cells(f"C{rr}:M{rr}")
            put(ws, f"N{rr}", hl(href, lk), F_LINK)


# ======================================================================================
# 17_DIFERENCIAS
# ======================================================================================
def hoja_diferencias(c: Ctx):
    ws = c.ws["17_DIFERENCIAS"]
    title(ws, "17 · DIFERENCIA DETECTADA ENTRE RESULTADO PRESENTADO Y FUENTE",
          "Ninguna diferencia se corrige en silencio: el resultado oficial se mantiene y aquí se documenta la fuente, la explicación, el impacto y la acción.")
    cols = [("ID", 6), ("Concepto", 28), ("Resultado presentado", 20), ("Diap.", 7), ("Resultado en la fuente / modelo", 26), ("Archivo de origen", 30),
            ("Hoja / registro", 26), ("Posible explicación", 50), ("Impacto en el estudio", 34), ("Acción recomendada para validarla", 44), ("Responsable recomendado", 22), ("Estado", 16)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    st = c.fab
    nparque = len(c.S)
    DIM = "DIMENSIONAMIENTO_MANTENIMIENTO_DATOS_FINAL.xlsx"
    FAB = "MAESTRO_FABRICA_FINAL.xlsx"
    filas = [
        ("Distribución clase B", "275", "4", "=\"119 (fuente) · \"&COUNTIF(S_CLASE,\"B\")&\" (score)\"", FAB, "01_Maestro_Fabrica · Histórico 'B - Alta carga'",
         "Reasignación declarada de clases al cerrar el estudio ('los 2.043 activos se consideran clasificados').", "Alto: B aporta 13,75 FTE.", "Documentar el criterio de la reasignación o listar los 275 activos B.", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Distribución clase C", "1.200", "4", "=\"189 (fuente) · \"&COUNTIF(S_CLASE,\"C\")&\" (score)\"", FAB, "01_Maestro_Fabrica · 'C - Media' + 'C/Pendiente'",
         "Valor redondo: absorbe no clasificados y D sin impacto demostrado.", "Alto: C aporta 34,29 FTE (40 % de la base).", "Sustituir por el recuento activo a activo validado (03).", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Distribución clase D", "150", "4", "=\"1.317 (fuente) · \"&COUNTIF(S_CLASE,\"D\")&\" (score)\"", FAB + " · " + DIM, "Histórico 'D - Baja/ruta' 920 + 397 no clasificados",
         "La regla 'D solo con bajo impacto demostrado' mueve la mayoría a C.", "Medio: cada 35 activos D→C añaden ~0,4 FTE.", "Validar la lista de activos D con Producción.", "Producción + Mantenimiento", "ABIERTA"),
        ("Clase A = 418", "418", "4 · 8", f"{st['CB'][0]} CB + {st['A2'][0]} A2 + {st['MONT_A1'][0]} montaje + {st['MEB_MONT'][0]} MEB montaje + {st['MEB_OF'][0]} MEB oficial", FAB, "01_Maestro_Fabrica",
         f"Recuento de filas: 6 filas MEB representan 75 activos; los 129 cuellos son {st['CB_UNICOS']} máquinas SAP + {st['CB_SIN']} entradas sin equivalencia (líneas/islas); {st['SOLAPE_CB_A2']} máquinas A2 también son cuello.",
         "Medio: posible sobre/infra-conteo de A (±10–60 activos).", "Recontar A sobre activos únicos (03 da el contraste).", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Parque total 2.043", "2.043", "4 · 8", f"SAP depurado 2.078 · parque fila a fila del modelo {nparque}", FAB + " · LISTADO MAQUINAS SAP.xlsx", "07_Contexto_Global · 01_MAESTRO_ORDENADO",
         "2.043 = suma de familias globales (incluye 386 medios de transporte y 355 instalaciones de transporte que la depuración SAP excluye como periféricos).",
         "Alto sobre el denominador del dimensionamiento.", "Conciliar familia a familia (13 §1) y fijar el inventario oficial.", "Ingeniería + Mantenimiento", "ABIERTA"),
        ("Necesidad total", "113–121 (diap. 5) · 115–125 (diap. 7)", "5 · 7", '=TEXT(K_NEC_MIN,"0.0")&"–"&TEXT(K_NEC_MAX,"0.0")&" (selector "&P_COB_BASE&")"', O.PPT + " · " + DIM, "Diap. 2, 5, 7 · 04_Calculo_FTE",
         "La cobertura 15–25 % se aplica sobre bases distintas: base (diap. 5), base + PM (Excel), subtotal (método diap. 2).", "±2 FTE.", "Fijar la base de cobertura (P_COB_BASE) y unificar diapositivas.", "Dirección técnica", "ABIERTA"),
        ("GAP", "+33 (diap. 5) · +34 (diap. 7)", "5 · 7", '=TEXT(K_GAP_MIN,"0.0")&" / "&TEXT(K_GAP_MAX,"0.0")', O.PPT + " · " + DIM, "04_Calculo_FTE B15",
         "+33 = 113 − 80; +34 = 114,2 − 80 (Excel); con método subtotal = 35,0. Máximo no presentado (44–45).", "±2 FTE.", "Presentar GAP mín.–máx. con una única base.", "Dirección técnica", "ABIERTA"),
        ("Resumen del Excel de dimensionamiento", "—", "—", "68,1 / 11,6 / 19,4 / 114 / 124 / '+14 a +22'", DIM, "00_Resumen B10:B19 y A23",
         "Valores escritos a mano de una versión anterior (A 418 con B/C/D de la fuente = 68,1 FTE). La conclusión '+14 a +22' contradice sus propias filas 34/44.", "Bajo (no usado en el PPT).", "No usar 00_Resumen; usar 04_Calculo_FTE o este maestro.", "Autor del estudio", "DOCUMENTADA"),
        ("FTE por grupo MEB", "EBOX 4,66 · MANGUETAS 5,21 · KMM 4,22 · ALUMINIO 6,15", "7", f"=TEXT({chr(39)}05_FTE_POR_GRUPOS{chr(39)}!G23,\"0.00\")&\" · \"&TEXT({chr(39)}05_FTE_POR_GRUPOS{chr(39)}!G24,\"0.00\")&\" · \"&TEXT({chr(39)}05_FTE_POR_GRUPOS{chr(39)}!G25,\"0.00\")&\" · \"&TEXT({chr(39)}05_FTE_POR_GRUPOS{chr(39)}!G27,\"0.00\")",
         DIM, "03_Secciones filas 8–12", "La tabla incrustada del gráfico está escrita a mano; no coincide con la fuente (MQ200, MQ281 y CONJ. DIF. sí coinciden).", "Medio: 10 FTE de diferencia en nuevos productos.",
         "Recalcular el reparto desde activos validados (05 §3).", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Doble conteo MEB en 03_Secciones", "—", "7", "A = A_MEB (14, 17, 15, 14, 15) y se suman ambos", DIM, "03_Secciones B8:F12",
         "El activo crítico MEB se cuenta en A y otra vez en A_MEB.", "Medio: +6,25 FTE en grupos MEB.", "Eliminar A_MEB o A en el reparto definitivo.", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Cierre 'RESTO PLANTA'", "—", "7", "51,69 FTE", DIM, "03_Secciones G13", "Se completa el reparto por diferencia 'para cuadrar'. El PPT muestra 0 en SERV. COMUNES.", "Alto: 60 % de la base sin grupo.",
         "No cerrar por diferencia: asignar desde activos.", "Ingeniería de Mantenimiento", "ABIERTA"),
        ("Puntos de preventivo", "15.000 'yoY'", "8", "50.000 puntos (usados en el cálculo)", O.PPT_CTX, "Diap. 1", "Cifra sin fuente en el repositorio; el cálculo de preventivos usa 50.000.",
         "Alto si fueran 15.000: preventivos pasarían de 9,5 a ~4 FTE.", "Confirmar puntos anuales reales en SAP PM.", "Planificación de Mantenimiento", "ABIERTA"),
        ("Horas de avería 2025→actual", "29.627 h (Excel)", "—", "=TEXT(K_AV_H,\"#,##0\")&\" h\"", DIM + " · SAP.EXPORT.AVERIAS.LISTADO.xlsx", "00_Resumen B9 · 05_Datos_2025_actual B7",
         "El filtro 'Años contiene 2025/2026' suma el histórico completo 2022–2026 de esas máquinas.", "Bajo (no está en el PPT final; solo C4).", "Usar la suma por fecha de avería (10).", "Ingeniería de Mantenimiento", "DOCUMENTADA"),
        ("Estaciones de montaje", "264 (filas)", "—", "183 estaciones (presentación de contexto)", FAB + " · " + O.PPT_CTX, "07_Contexto_Global fila 6 · diap. 1",
         "El fichero de montaje incluye standby, datos comunes, palets, armarios e hidráulicos.", "Medio sobre A (montaje = A).", "Contar solo estaciones activas principales (19 y 03).", "Mantenimiento Montaje", "ABIERTA"),
        ("Identificación MEB", "KMM 15 · EBOX 14 · CDIF 10+4 · MANG 17 · ALU 15", "—", "=\"Modelo: \"&K_MEB_IDENT&\" identificados\"", FAB + " · LISTADO_EQUIPOS_SAP.xlsx", "02_MEB_ECO_Oficial · 09 §2",
         "CDIF: estaciones AFO + 4 Reishauer localizadas en SAP (la fuente decía 6). EBOX y Manguetas no coinciden 1:1. Aluminio: 0 identificados.", "Medio.", "Validar la lista nominal de los 75 con Ingeniería MEB.", "Ingeniería MEB ECO", "ABIERTA"),
        ("Export de averías", "—", "—", "Fila 'Total' 2.025.517 min vs filas 434.029 min", "SAP.EXPORT.AVERIAS.LISTADO.xlsx", "AVERIAS fila 2",
         "La fila Total es de la vista completa 2022→; el export solo contiene 2025→.", "Ninguno en el cálculo.", "Ninguna (documentado).", "—", "DOCUMENTADA"),
        ("Ubicación de los centros MAG L5/L6", "—", "—", "Gr. 66511/06 y /07 → MECANIZADO ALUMINIO MQ200/MQ281", "LISTADO MAQUINAS SAP.xlsx · CUELLOS_DE_BOTELLA.xlsx", "01_MAESTRO_ORDENADO · Export",
         "Los cuellos E-PRODUCTS-AL-GGH/MGH apuntan a centros MAG de las líneas 5/6, clasificadas por ubicación como convencional.", "Medio sobre el reparto MEB.", "Confirmar si L5/L6 son Aluminio GGH/MGH (MEB).", "Ingeniería MEB ECO", "ABIERTA"),
    ]
    for i, row in enumerate(filas):
        r = 6 + i
        put(ws, f"A{r}", f"D{i + 1:02d}", F_BOLD, align=CENTER)
        for j, v in enumerate(row):
            col = "BCDEFGHIJKL"[j]
            put(ws, f"{col}{r}", v, F_BOLD if col in "BC" else F_BASE, "oficial" if col == "C" else None, align=WRAP)
        ws.row_dimensions[r].height = 64
    last = 5 + len(filas)
    status_format(ws, f"L6:L{last}")
    define(c.wb, "T_DIF", "17_DIFERENCIAS", f"$L$6:$L${last}")
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:L{last}"


# ======================================================================================
# 13_METODOLOGIA
# ======================================================================================
def _parrafos(ws, r, titulo, lineas, ncols=8):
    section(ws, r, titulo, ncols)
    r += 1
    for ln in lineas:
        c = ws.cell(r, 1, ln)
        c.font = F_BASE
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.row_dimensions[r].height = max(15, 15 * (1 + len(ln) // 150))
        r += 1
    return r + 1


def hoja_metodologia(c: Ctx):
    ws = c.ws["13_METODOLOGIA"]
    title(ws, "13 · METODOLOGÍA", "Cómo se construye cada cifra. Orden: Inventario → Criticidad → Carga → Disponibilidad real → Necesidad → GAP.")
    widths(ws, {"A": 22, "B": 30, "C": 26, "D": 18, "E": 30, "F": 22, "G": 22, "H": 46})
    r = 5
    r = _parrafos(ws, r, "0. Principio de trabajo", [
        "POWERPOINT FINAL (PK-16.DATOS.MANTENIMIENTO.V2.pptx) = RESULTADO OFICIAL PRESENTADO. EXCEL MAESTRO = DATOS + CÁLCULOS + CRITERIOS + METODOLOGÍA + TRAZABILIDAD.",
        "El Excel reproduce el resultado oficial con fórmulas que leen 12_SUPUESTOS. No genera una versión alternativa: cuando una fuente da otra cifra, se registra en 17_DIFERENCIAS.",
        "El contraste activo a activo (03) se muestra para justificar, priorizar y detectar huecos. No sustituye al resultado oficial mientras no se valide.",
        "No se ajustan datos para justificar una plantilla: la conclusión sale de los datos y de supuestos visibles.",
    ])
    r = _parrafos(ws, r, "1. Construcción del inventario (02_MAESTRO_CRITICIDAD)", [
        "Base: exportación SAP completa LISTADO_EQUIPOS_SAP.xlsx (2.857 equipos). Ninguna fila se borra: cada registro conserva su estado y motivo.",
        "Depuración previa (LISTADO MAQUINAS SAP.xlsx): 779 periféricos excluidos (carros, potros, puestos de control, transferizados, satélites EWAB, Hoffmann, microscopios, filtros, calibres…). Se mantienen como 'Excluido – auxiliar / periférico' con su grupo heredado.",
        "Depuración ampliada del modelo (PENDIENTE DE VALIDACIÓN): palets, mesas, utillajes, maletas de programación y carros que SAP registra como estaciones.",
        "Componentes integrados: un equipo cuya ubicación cuelga de la ubicación de otra estación (p. ej. ST10-1 bajo el puesto 10) se cuenta en el activo principal. Los robots se mantienen como activos propios.",
        "Agrupadores de línea ('GRUPO SECCIÓN…', 'LÍNEA MONTAJE DIFERENCIAL 7016426001', 'GRUPO A/B/C/D MEDIOS') no se cuentan para no duplicar sus estaciones.",
        "Montaje convencional: manda el fichero de montaje actualizado (264 filas). Las estaciones SAP que no aparecen en él quedan como 'posible baja' fuera del recuento; las del fichero que no están en SAP se añaden como 'no localizado en SAP'. Cruce por (línea, estación) con 90410 → 7016420001 y 90510 → 7016420002.",
        "Montaje MEB ECO (líneas nuevas, fichero incompleto): manda SAP a nivel estación. Standby / reserva / reparación ('S/R', 'STANDBY') se conservan con ese estado.",
        "Cuellos de botella sin equivalencia SAP (líneas o islas completas) se registran como 'pendiente de confirmar' sin sumarlos al parque, para no duplicar máquinas ya contadas.",
    ])
    r = _parrafos(ws, r, "2. Clasificación por ubicación técnica (regla jerárquica)", [
        "Primero se aplica la ubicación más específica; después la ubicación superior. Cada equipo hereda producto, línea y proceso de la raíz más larga que contiene su ubicación (coincidencia exacta o por descendencia, nunca por texto).",
        "Control: esta regla reproduce el 100 % de la asignación de la depuración previa (2.078 registros). Las raíces marcadas 'Propuesta del modelo' clasifican los 516 registros que estaban sin coincidencia y quedan PENDIENTES DE VALIDACIÓN.",
    ])
    header_row(ws, r, 1, ["Ubicación raíz", "Grupo SAP (depuración)", "Línea / sección", "Programa", "Producto", "Proceso", "Grupo mantenimiento", "Origen de la regla"])
    r += 1
    for _, x in D.ROOTS_DF.sort_values("RAIZ").iterrows():
        for j, f in enumerate(["RAIZ", "GRUPO_SAP", "LINEA", "PROGRAMA", "PRODUCTO", "PROCESO", "GRUPO_MANT", "ORIGEN_REGLA"]):
            put(ws, f"{'ABCDEFGH'[j]}{r}", x[f], fill_="pendiente" if "PENDIENTE" in x["ORIGEN_REGLA"] and j == 7 else None, align=WRAP)
        r += 1
    r += 1
    section(ws, r, "3. Criterios de criticidad (diapositiva 4) y cómo se valoran con los datos disponibles", 8)
    r += 1
    header_row(ws, r, 1, ["Criterio", "Escala", "", "Máximo", "Qué valora", "", "Evidencia usada por el modelo", "Limitación"])
    crit = [
        ("C1 Impacto en producción", "0 sin impacto · 1 pérdida menor · 2 parada de estación con alternativa · 3 pérdida importante de proceso · 4 parada de línea/producto/proceso crítico", "4",
         "Qué ocurre cuando falla.", "Tipología del activo (12 §12); 4 si cuello de botella, montaje principal o MEB crítico; mínimo 3 si SAP ABC = A.", "No hay matriz de impacto por activo validada con Producción."),
        ("C2 Seguridad, calidad, cliente y trazabilidad", "0 · 1 · 2", "2", "Riesgo personas, conformidad, defectos, escape a cliente, trazabilidad.",
         "Tipología + indicio en la descripción (test, fugas, estanqueidad, control, medición, atornillado, prensado, soldadura, láser/grabado, hornos/TT, visión).", "Sin AMFE ni registro de incidencias de calidad por activo."),
        ("C3 Unicidad, redundancia y cuello de botella", "0 redundante · 0,5 existe equivalente · 1 redundancia no confirmada · 2 único / cuello", "2",
         "Alternativa, transferencia, recuperación.", "Cuello de botella confirmado = 2 (regla). Equivalente = mismo tipo SAP, fabricante y ubicación.", "Transferibilidad real no documentada."),
        ("C4 Averías, recurrencia y paro (2025→actual)", "0 · 0,5 · 1 · 1,5", "1,5", "Número, horas, parada máxima, recurrencia.",
         "SAP.EXPORT.AVERIAS (2025-01-03 → último registro) por activo; montaje con data.1/data.2. Umbrales en 12 §10.", "El export excluye montaje y TPM; no hay MTTR ni complejidad."),
        ("C5 Obsolescencia, repuestos y soporte", "0 · 0,25 · 0,5", "0,5", "Antigüedad, PLC/CNC, repuestos, soporte.",
         "Año de puesta en servicio SAP (1900 = desconocido → valor medio).", "Sin inventario de PLC/CNC, descatalogados ni plazos de repuesto."),
    ]
    r += 1
    for row in crit:
        for j, v in enumerate(row):
            put(ws, f"{'ABDEGH'[j]}{r}", v, F_BOLD if j == 0 else F_BASE, align=WRAP)
        ws.merge_cells(f"B{r}:C{r}")
        ws.merge_cells(f"E{r}:F{r}")
        ws.row_dimensions[r].height = 62
        r += 1
    r += 1
    r = _parrafos(ws, r, "4. Score, clases y reglas especiales", [
        "SCORE TOTAL = C1 + C2 + C3 + C4 + C5 (máximo 10). Clases oficiales: A ≥ 8 · B 6–7,9 · C 3,5–5,9 · D < 3,5.",
        "Reglas que prevalecen: cuello de botella confirmado = A (máxima valoración en unicidad) · activo capaz de parar una línea completa = máxima valoración productiva · montaje principal = A (regla acordada) · activos críticos de nuevos productos confirmados conservan su prioridad (A).",
        "La clase D se reserva a bajo impacto demostrado. La falta de información nunca envía un activo a D: con datos insuficientes (tipología no tipificada o grupo propuesto) la clase queda como 'C provisional' y PENDIENTE DE VALIDACIÓN.",
        "Toda discrepancia entre la clase calculada y la clase presentada se registra (04 §4 y 17); la clase oficial del PPT no se sustituye.",
        "Motivo de criticidad (03 col. AH): frase generada por fórmula a partir de los criterios (impacto, redundancia, cuello, seguridad/calidad, averías y horas, obsolescencia, información pendiente).",
    ])
    r = _parrafos(ws, r, "5. Puntos equivalentes y dimensionamiento por criticidad", [
        "Puntos equivalentes por clase (A 4 · B 2,5 · C 1,5 · D 1; MAESTRO_CRITICIDAD_MAQUINAS · Criterios): sirven para visualizar la cartera, comparar grupos y priorizar. No son horas y no sustituyen al FTE.",
        "FTE de la clase = Número de activos de la clase ÷ Ratio aplicado. FTE base = FTE A + FTE B + FTE C + FTE D.",
        "Los ratios A 12 · B 20 · C 35 · D 60 activos por mantenedor son un criterio interno de dimensionamiento de cartera, apoyado en metodología de carga de trabajo, criticidad, benchmarking y contexto operativo. No son una norma universal publicada literalmente por SMRP ni por NASA.",
    ])
    r = _parrafos(ws, r, "6. Capas de capacidad (diapositivas 2, 3 y 5)", [
        "01 Base por criticidad: Σ activos por clase ÷ ratio.",
        "02 Carga planificada: (puntos PM + avisos) ÷ productividad por FTE. Las unidades no consumen el mismo tiempo: se marca como HIPÓTESIS hasta tener horas reales (frecuencia × duración de los planes SAP).",
        "TPM informado como actividad operativa. No sumado como FTE central adicional cuando ya está incluido en la actividad del grupo.",
        "03 Formación = Plantilla nominal × horas de formación por persona ÷ horas productivas anuales por FTE (resultado exacto y valor redondeado presentado).",
        "04 Fines de semana y festivos: descanso compensatorio (libranza) = personas por jornada × jornadas anuales ÷ jornadas efectivas por FTE; con horas reales, horas ÷ horas productivas. No son vacaciones.",
        "05 Cobertura organizativa = base × factor 15–25 % (vacaciones, absentismo, turnos, especialidades mínimas, simultaneidad, reuniones, indisponibilidades). No reincluye formación ni libranzas. La base es un selector (P_COB_BASE) porque el PPT usa tres bases distintas.",
        "06 Necesidad total = base + carga + pérdidas + cobertura. GAP = necesidad − plantilla nominal. El GAP es ORIENTATIVO mientras plantilla, turnos, especialidades y pérdidas de capacidad no estén validados. La plantilla nominal no equivale a capacidad disponible (06 §3).",
        "Especialidades: sin plantilla por especialidad no se reparten FTE (PENDIENTE DE ASIGNACIÓN POR ESPECIALIDAD); se informa el nº de activos por tecnología como orientación.",
    ])
    section(ws, r, "7. Referencias metodológicas externas", 8)
    r += 1
    header_row(ws, r, 1, ["Referencia", "", "", "URL", "", "Qué respalda", "", ""])
    r += 1
    for ref, url, txt in O.REFERENCIAS_EXTERNAS:
        put(ws, f"A{r}", ref, F_BOLD, align=WRAP)
        ws.merge_cells(f"A{r}:C{r}")
        cl = put(ws, f"D{r}", url, F_LINK, align=WRAP)
        cl.hyperlink = url
        ws.merge_cells(f"D{r}:E{r}")
        put(ws, f"F{r}", txt, align=WRAP)
        ws.merge_cells(f"F{r}:H{r}")
        ws.row_dimensions[r].height = 45
        r += 1
    r = _parrafos(ws, r + 1, "Lectura de las referencias", [
        "Estas referencias respaldan una metodología basada en carga de trabajo, mantenimiento preventivo y correctivo, criticidad, recursos, planificación, benchmarking y contexto operativo.",
        "Los ratios concretos A/B/C/D utilizados en el estudio son un criterio interno de cartera y no un ratio universal publicado literalmente por SMRP o NASA.",
    ])


# ======================================================================================
# 14_FUENTES
# ======================================================================================
def _fecha_mod(path):
    try:
        import openpyxl
        if str(path).endswith(".xlsx"):
            wb = openpyxl.load_workbook(path, read_only=True)
            m = wb.properties.modified
            wb.close()
            return m.date() if m else None
        if str(path).endswith(".pptx"):
            from pptx import Presentation
            m = Presentation(path).core_properties.modified
            return m.date() if m else None
    except Exception:
        return None
    return None


def hoja_fuentes(c: Ctx):
    ws = c.ws["14_FUENTES"]
    title(ws, "14 · INVENTARIO DE FUENTES", "Repositorio ROCOSA00/seat-componentes (raíz) y release 'ARCHIVOS' (ficheros pesados). Confianza: Alta / Media / Baja.")
    cols = [("Tipo", 18), ("Archivo", 40), ("Ubicación", 16), ("Hojas usadas", 30), ("Registros", 10), ("Información aportada", 46), ("Confianza", 10),
            ("Actualización", 12), ("Uso en la presentación", 30), ("Disponibilidad", 16), ("Observaciones", 56)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=36)
    F = D.F
    av = c.M["av"]
    rows = [
        ("PowerPoint final", F["PPT"], "Release", "8 diapositivas", 8, "Resultado oficial: distribución A/B/C/D, ratios, FTE base, capas, necesidad, GAP, reparto por grupos.", "Oficial",
         "Diapositivas 1–8", "DISPONIBLE", "Diap. 6 vacía; diap. 8 con marcadores 'XX'. Gráficos con tablas incrustadas escritas a mano."),
        ("Presentación de contexto", F["PPT_CONTEXTO"], "Release", "2 diapositivas", 2, "Familias del parque (430 mecanizado, 92 robots…), 9.000 planes, 50.000 puntos PM, 10.000 avisos, 3.000 TPM, instalaciones de planta.",
         "Media", "Contexto / carga global", "DISPONIBLE", "Cifras globales ('+ de'), sin detalle por activo."),
        ("Inventario SAP", F["SAP_RAW"], "Release", "Hoja1", 2857, "Exportación completa de equipos: ubicación técnica, tipo, fabricante, año, indicador ABC.", "Alta",
         "Base del inventario fila a fila", "DISPONIBLE", "Sin 'Equipo superior' (jerarquía solo por ubicación). Columna 'Autor' no se reproduce (dato personal)."),
        ("Inventario SAP depurado", F["SAP_DEP"], "Repositorio", "01_MAESTRO_ORDENADO · 13_PENDIENTES · 14_ELIMINADOS · 15_FILTROS", 2078,
         "Clasificación por ubicación raíz (exacta/heredada), 516 pendientes, 779 eliminados con motivo.", "Alta", "Depuración del parque", "DISPONIBLE", "La regla jerárquica del modelo la reproduce al 100 %."),
        ("Maestro de fábrica", F["FAB"], "Repositorio", "01_Maestro_Fabrica · 02/03 MEB · 04_Montaje · 05_Paros · 06_Cuellos · 07_Contexto", 1660,
         "Origen de A = 418, MEB 75 críticos, montaje 264, contexto 2.043.", "Media", "Diap. 4, 8", "DISPONIBLE", "Mezcla filas agregadas e individuales."),
        ("Dimensionamiento (cálculo)", F["DIM"], "Repositorio", "01_Supuestos · 02_Parque · 03_Secciones · 04_Calculo_FTE", 0,
         "Supuestos (80, 55 h, 1.450 h, 10×45/220, 60.000/6.300, 15–25 %), cálculo FTE y reparto por secciones.", "Media", "Diap. 4, 5, 7", "DISPONIBLE", "00_Resumen desactualizado (ver 17)."),
        ("Libro consolidado previo", F["PROF"], "Repositorio", "30_Dimensionamiento · 31_Analisis · 41_Comprobaciones", 0, "Consolidación previa de fábrica + dimensionamiento.", "Media",
         "Versión anterior", "DISPONIBLE", "Duplica MAESTRO_FABRICA y DIMENSIONAMIENTO (ORG_*)."),
        ("Maestro de criticidad", F["CRIT"], "Release", "Maestro Criticidad · Criterios", 1247, "Criticidad previa por paros + cuellos; puntos equivalentes A 4/B 2,5/C 1,5/D 1; rangos de ratio.",
         "Media", "Ratios y puntos", "DISPONIBLE", "Distribución A31/B84/C185/D947 (criterio anterior)."),
        ("Maestro integral", F["INTEGRAL"], "Release", "01_Maestro_Activos · 05_Cuellos_Botella · 07_Preventivo_TPM_FTE", 1519,
         "Mapeo de cuellos a activo; escenarios PM/TPM; criticidad A1/A2 previa.", "Media", "Escenarios de contraste", "DISPONIBLE", ""),
        ("Averías", F["AVERIAS"], "Release", "AVERIAS", len(av), "Órdenes de avería con fecha, objeto técnico, ubicación y tiempo de paro (min) 2025→2026.",
         "Alta", "Criterio 'Averías y paro 2025→actual'", "DISPONIBLE", "Excluye montaje y TPM (filtro Power BI)."),
        ("Cuellos de botella", F["CB"], "Release", "Export", 129, "129 cuellos con modelo, línea y OEE.", "Alta", "Criterio cuello de botella", "DISPONIBLE", "Todos A por definición."),
        ("Paros montaje línea 1", F["MONT_L1"], "Release", "Export", int((c.M["logs"]["ARCHIVO"] == "data.1.xlsx").sum()), "Paros por estación MQ200 Línea 1 (2025).", "Media", "C4 montaje", "DISPONIBLE", "Unidad 'CP' a validar."),
        ("Paros montaje línea 3", F["MONT_L3"], "Release", "Export", int((c.M["logs"]["ARCHIVO"] == "data.2.xlsx").sum()), "Paros por estación MQ281 Línea 3 (2025).", "Media", "C4 montaje", "DISPONIBLE", "Unidad 'CP' a validar."),
        ("Análisis Power BI", F["TOP"], "Release", "Export", 0, "Vacío (exportación sin datos).", "—", "—", "SIN DATOS", "Contiene solo el aviso de Power BI."),
    ]
    falta = [
        ("Maestro oficial de activos", "MAQRAIZ", "—", "—", None, "Máquina raíz y jerarquía padre/hijo oficial.", "—", "—", "NO DISPONIBLE", "Sin él, la jerarquía se deduce por ubicación técnica (02)."),
        ("Histórico de paros original", "ANALISIS_PAROS_MAQUINAS_V2 3.xlsx", "—", "—", None, "Origen del histórico agregado 2022–2026 (1.247 máquinas).", "—", "—", "NO DISPONIBLE", "Referenciado en MAESTRO_INTEGRAL; su contenido agregado está en 05_Paros_Historico."),
        ("Líneas de montaje", "LINEAS DE MONTAJE GENERALES DE TODA LA FABRICA.xlsx", "—", "—", None, "Fichero de montaje original.", "—", "—", "NO DISPONIBLE (contenido en 04_Montaje_Completo)", ""),
        ("Preventivos", "Planes SAP PM (frecuencia, duración, horas)", "—", "—", None, "Horas reales de preventivo por máquina.", "—", "—", "NO DISPONIBLE", "Por eso PM se convierte con productividad hipotética."),
        ("Organización", "Plantilla de mantenimiento por especialidad y turno", "—", "—", None, "Personas, FTE, turnos, especialidades, vacantes.", "—", "—", "NO DISPONIBLE", "Plantilla '~80' facilitada verbalmente."),
        ("Turnos / fines de semana", "Calendario de sábados, festivos y cobertura", "—", "—", None, "Jornadas especiales reales y horas.", "—", "—", "NO DISPONIBLE", ""),
        ("Formación", "Plan de formación anual", "—", "—", None, "Horas reales por persona.", "—", "—", "NO DISPONIBLE", ""),
        ("Instalaciones de planta", "Inventario de utilities (compresores, AT/BT, EDAR, PCI…)", "—", "—", None, "Instalaciones listadas en la presentación de contexto (diap. 2).", "—", "—", "NO DISPONIBLE", "No forman parte de los 2.043."),
    ]
    r = 6
    for row in rows:
        tipo, path, ubic, hojas, nreg, info, conf, uso, disp, obs = row
        vals = [tipo, path.name, ubic, hojas, nreg, info, conf, _fecha_mod(path), uso, disp, obs]
        for j, v in enumerate(vals):
            put(ws, f"{get_column_letter(j + 1)}{r}", v, F_BOLD if j == 1 else F_BASE, "dato" if j == 1 else None,
                "dd/mm/yyyy" if j == 7 else ("#,##0" if j == 4 else None), align=WRAP)
        r += 1
    for tipo, arch, ubic, hojas, nreg, info, conf, uso, disp, obs in falta:
        vals = [tipo, arch, ubic, hojas, nreg, info, conf, None, uso, disp, obs]
        for j, v in enumerate(vals):
            put(ws, f"{get_column_letter(j + 1)}{r}", v, F_BOLD if j == 1 else F_BASE, "pendiente", align=WRAP)
        r += 1
    r0 = r
    for ref, url, txt in O.REFERENCIAS_EXTERNAS:
        vals = ["Referencia metodológica externa", ref, "Web", "—", None, txt, "Alta", None,
                "Diap. 3 'Fuentes externas – metodologías'", "DISPONIBLE (pública)", "Respalda la metodología por carga de trabajo; no es la fuente literal de los ratios internos."]
        for j, v in enumerate(vals):
            put(ws, f"{get_column_letter(j + 1)}{r}", v, F_BOLD if j == 1 else F_BASE, align=WRAP)
        ws[f"B{r}"].hyperlink = url
        r += 1
    status_format(ws, f"J6:J{r}")
    define(c.wb, "T_DISP", "14_FUENTES", f"$J$6:$J${r - 1}")
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:K{r - 1}"


# ======================================================================================
# 15_PENDIENTES_VALIDACION
# ======================================================================================
def hoja_pendientes(c: Ctx):
    ws = c.ws["15_PENDIENTES_VALIDACION"]
    title(ws, "15 · PENDIENTES DE VALIDACIÓN", "Incertidumbres abiertas. Prioridad por impacto en las cifras oficiales.")
    cols = [("ID", 6), ("Tema", 26), ("Descripción", 60), ("Nº registros", 11), ("Impacto", 10), ("KPI afectado", 22), ("Responsable recomendado", 24),
            ("Acción recomendada", 52), ("Estado", 16), ("Ir a", 14)]
    header_row(ws, 5, 1, [h for h, _ in cols], [w for _, w in cols], height=30)
    inv = c.inv
    n = lambda cond: int(cond.sum())
    P = [
        ("Identificación exacta de activos", "No existe MAQRAIZ ni maestro oficial en el repositorio; el inventario se reconstruye desde SAP + montaje.", None, "Alta", "Parque 2.043", "Ingeniería de Mantenimiento", "Aportar MAQRAIZ o validar 02_MAESTRO_CRITICIDAD.", "02_MAESTRO_CRITICIDAD"),
        ("Parque oficial vs inventario", "2.043 es suma de familias globales; el parque fila a fila del modelo y el SAP depurado no coinciden con ese total.", len(c.S), "Alta", "Parque, FTE base", "Ingeniería + Mantenimiento", "Conciliar familia a familia y fijar el inventario oficial.", "04_FTE_CRITICIDAD"),
        ("Distribución B/C/D", "275 / 1.200 / 150 no tienen registro activo a activo (fuente: 119 / 189 / 1.317).", None, "Alta", "FTE base", "Ingeniería de Mantenimiento", "Validar la clase de cada activo en 03 y recalcular.", "04_FTE_CRITICIDAD"),
        ("Máquina raíz y componentes", "Componentes deducidos por ubicación hija de una estación.", n(inv["ESTADO"].str.startswith("Componente")), "Media", "Parque", "Mantenimiento", "Confirmar qué sub-estaciones son activos independientes.", "02_MAESTRO_CRITICIDAD"),
        ("Bajas", "Estaciones SAP de montaje convencional no presentes en el fichero de montaje actualizado.", n(inv["ESTADO"].str.contains("posible baja")), "Media", "Parque, clase A", "Mantenimiento Montaje", "Confirmar baja, renombrado o reserva.", "02_MAESTRO_CRITICIDAD"),
        ("Duplicados", "Descripciones repetidas en SAP (misma descripción en varios equipos).", int(inv[inv["EQUIPO"] != ""]["DESC"].str.upper().duplicated(keep=False).sum()), "Media", "Averías por activo", "Mantenimiento", "Revisar equipos con descripción repetida; afecta al cruce de averías.", "18_REG_AVERIAS"),
        ("Reservas y standby", "Puestos standby / reserva / reparación: se mantienen en parque con baja criticidad.", n(inv["ESTADO"].str.startswith("Standby")), "Baja", "Clase", "Mantenimiento Montaje", "Confirmar qué puestos están en servicio.", "02_MAESTRO_CRITICIDAD"),
        ("Depuración ampliada", "Palets, mesas, utillajes, maletas y carros excluidos por el modelo.", n(inv["ESTADO"].str.contains("ampliada")), "Media", "Parque", "Mantenimiento", "Validar la exclusión.", "02_MAESTRO_CRITICIDAD"),
        ("Grupo propuesto", "Registros sin coincidencia en la depuración previa clasificados por la descripción de su ubicación.", n(inv["ESTADO"].str.contains("propuesta")), "Media", "Reparto por grupos", "Mantenimiento", "Validar la tabla de ubicaciones raíz 'Propuesta' (13 §2).", "13_METODOLOGIA"),
        ("Redundancia", "C3 = 1 por defecto (redundancia no confirmada) salvo cuello, montaje o equivalente detectado.", None, "Alta", "Clase", "Producción", "Indicar alternativa y transferencia por activo crítico.", "03_SCORE_CRITICIDAD"),
        ("Impacto productivo", "C1 por tipología, no por análisis de impacto activo a activo.", None, "Alta", "Clase", "Producción", "Validar C1 de los activos A/B y de los provisionales.", "03_SCORE_CRITICIDAD"),
        ("Seguridad y calidad", "C2 por tipología y palabras clave.", None, "Media", "Clase", "Calidad + Prevención", "Validar con AMFE / evaluación de riesgos.", "03_SCORE_CRITICIDAD"),
        ("Obsolescencia y repuestos", "C5 por año de puesta en servicio; muchos activos sin año.", n((inv["EN_PARQUE"] == "S") & inv["ANIO"].isna()), "Media", "Clase", "Mantenimiento + Compras", "Inventariar PLC/CNC, descatalogados y plazos.", "03_SCORE_CRITICIDAD"),
        ("Activos provisionales", "Activos con información insuficiente o grupo propuesto: nunca D, quedan C provisional.", int((c.S["PROVISIONAL"] == "S").sum()), "Media", "Clase C/D", "Mantenimiento", "Completar la información y revalorar.", "03_SCORE_CRITICIDAD"),
        ("Horas preventivas", "Solo existen 50.000 puntos globales; no hay horas por plan.", None, "Alta", "Preventivos +10", "Planificación de Mantenimiento", "Exportar planes SAP PM con frecuencia y duración.", "08_PM_AVISOS"),
        ("Puntos de preventivo", "Diapositiva 8: 15.000 'yoY' frente a 50.000 usados en el cálculo.", None, "Alta", "Preventivos", "Planificación de Mantenimiento", "Confirmar la cifra correcta.", "17_DIFERENCIAS"),
        ("Unidad de paro montaje", "'CP' de data.1/data.2 se interpreta en minutos.", None, "Baja", "C4 montaje", "Mantenimiento Montaje", "Confirmar unidad.", "10_AVERIAS"),
        ("Turnos", "Sin calendario de turnos ni cobertura por turno.", None, "Alta", "Cobertura, GAP", "Dirección de Mantenimiento + RRHH", "Aportar organización por turno.", "06_CAPACIDAD_PERSONAL"),
        ("Especialidades", "Sin plantilla por especialidad: no se reparten FTE.", None, "Alta", "Necesidad por especialidad", "Dirección de Mantenimiento + RRHH", "Aportar plantilla por especialidad.", "06_CAPACIDAD_PERSONAL"),
        ("Plantilla de referencia", "'~80 personas' sin fichero de organización (personas ≠ FTE).", None, "Alta", "GAP", "RRHH", "Confirmar plantilla nominal, vacantes y externos.", "06_CAPACIDAD_PERSONAL"),
        ("Base de cobertura", "El PPT aplica el 15–25 % sobre tres bases distintas.", None, "Media", "Necesidad, GAP", "Dirección técnica", "Fijar P_COB_BASE.", "06_CAPACIDAD_PERSONAL"),
        ("Solape nuevos productos", "MEB 75 críticos frente a estaciones SAP; Aluminio GGH/MGH sin inventario; A + A_MEB duplicados en 03_Secciones.", None, "Media", "Nuevos productos", "Ingeniería MEB ECO", "Lista nominal de los 75 y validación de L5/L6 MAG.", "09_MEB_ECO"),
        ("Reparto definitivo por grupos", "Diap. 7 'PENDIENTE REALIZAR': 41,8 FTE sin asignar y valores MEB distintos de la fuente.", None, "Alta", "FTE por grupo", "Ingeniería de Mantenimiento", "Adoptar 05 §3 tras validar clases.", "05_FTE_POR_GRUPOS"),
        ("Cruce de averías", "Objetos técnicos de avería sin activo SAP o con descripción repetida.", int((c.M["av"]["CRUCE"] != "Exacto por descripción").sum()), "Baja", "C4", "Mantenimiento", "Revisar nombres en SAP.", "10_AVERIAS"),
        ("Cruce de cuellos de botella", "Cuellos sin equivalencia SAP (líneas o islas completas).", int((c.M["cb"]["ID_ASIGNADO"] == "").sum()), "Media", "Clase A", "Producción", "Indicar la máquina limitante de cada línea/isla.", "20_CRUCE_CUELLOS"),
        ("Festivos y sábados", "Calendario real de jornadas especiales no disponible (festivos = 0).", None, "Baja", "Libranzas", "Dirección de Mantenimiento", "Aportar calendario.", "07_SABADOS_FORMACION"),
        ("Instalaciones de planta", "Utilities (compresores, AT/BT, EDAR, PCI…) sin inventario fila a fila ni en los 2.043.", None, "Media", "Parque", "Mantenimiento Instalaciones", "Inventariar y decidir si entran en el dimensionamiento.", "14_FUENTES"),
    ]
    for i, (tema, desc, nreg, imp, kpi, resp, acc, ira) in enumerate(P):
        r = 6 + i
        put(ws, f"A{r}", f"P{i + 1:02d}", F_BOLD, align=CENTER)
        put(ws, f"B{r}", tema, F_BOLD, align=WRAP)
        put(ws, f"C{r}", desc, align=WRAP)
        put(ws, f"D{r}", nreg, fmt="int")
        put(ws, f"E{r}", imp, F_BOLD, "ERROR" if imp == "Alta" else ("REVISAR" if imp == "Media" else "OK"), align=CENTER)
        put(ws, f"F{r}", kpi, align=WRAP)
        put(ws, f"G{r}", resp, align=WRAP)
        put(ws, f"H{r}", acc, align=WRAP)
        put(ws, f"I{r}", "PENDIENTE")
        put(ws, f"J{r}", hl(f"#'{ira}'!A1", "→ " + ira[:12]), F_LINK)
        ws.row_dimensions[r].height = 32
    last = 5 + len(P)
    status_format(ws, f"I6:I{last}")
    define(c.wb, "T_PEND", "15_PENDIENTES_VALIDACION", f"$I$6:$I${last}")
    define(c.wb, "T_PEND_IMP", "15_PENDIENTES_VALIDACION", f"$E$6:$E${last}")
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:J{last}"


# ======================================================================================
# 16_CONTROLES
# ======================================================================================
def hoja_controles(c: Ctx):
    ws = c.ws["16_CONTROLES"]
    title(ws, "16 · CONTROLES DE COHERENCIA", "Estados: OK · REVISAR · ERROR. Todos los controles son fórmulas sobre el propio libro.")
    widths(ws, {"A": 6, "B": 58, "C": 16, "D": 16, "E": 16, "F": 54, "G": 22})
    put(ws, "B5", "Resumen", F_BOLD)
    for i, (lbl, k) in enumerate((("OK", "OK"), ("REVISAR", "REVISAR"), ("ERROR", "ERROR"))):
        put(ws, f"C{5 + i}", lbl, F_BOLD, k, align=CENTER)
        put(ws, f"D{5 + i}", f'=COUNTIF($E$11:$E$60,"{lbl}*")', F_BOLD, fmt="int")
    define(c.wb, "K_CTRL_OK", "16_CONTROLES", "$D$5")
    define(c.wb, "K_CTRL_REV", "16_CONTROLES", "$D$6")
    define(c.wb, "K_CTRL_ERR", "16_CONTROLES", "$D$7")
    header_row(ws, 10, 1, ["#", "Control", "Resultado", "Esperado", "Estado", "Criterio / lectura", "Ir a"])
    C = [
        ("A + B + C + D coincide con el parque oficial", "=K_PARQUE_CLASES", "=P_PARQUE_OF", 'IF(C=D,"OK","ERROR")', "Suma de clases = 2.043.", "04_FTE_CRITICIDAD"),
        ("FTE base del modelo = FTE base del PPT (85,37)", "=K_FTE_BASE", O.OF["FTE_BASE"][0], 'IF(ABS(C-D)<=P_TOL_EXACT,"OK","ERROR")', "Reproducción exacta de la diap. 4.", "04_FTE_CRITICIDAD"),
        ("Ratios usados = ratios oficiales (12/20/35/60)", '=AND(P_RATIO_A=12,P_RATIO_B=20,P_RATIO_C=35,P_RATIO_D=60)', "VERDADERO", 'IF(C=TRUE,"OK","REVISAR")', "Cambiar ratios crea un escenario, no el oficial.", "12_SUPUESTOS"),
        ("Formación coincide con la presentación (+3)", "=K_FORM", O.OF["W_FORM"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "Exacto 3,03 → 3.", "07_SABADOS_FORMACION"),
        ("Libranzas coinciden con la presentación (+2)", "=K_SAB", O.OF["W_SAB"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "Exacto 2,05 → 2.", "07_SABADOS_FORMACION"),
        ("Preventivos coinciden con la presentación (+10)", "=K_PM", O.OF["W_PM"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "Exacto 9,52 → 10.", "08_PM_AVISOS"),
        ("TPM no duplicado (no sumado como FTE)", "=P_TPM_SUMAR", 0, 'IF(C=0,"OK","REVISAR")', "TPM informado como actividad operativa.", "08_PM_AVISOS"),
        ("Necesidad mínima = PPT diap. 7 (115)", "=K_NEC_MIN", O.OF["S7_REF_MIN"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "Depende de P_COB_BASE (ver 06 §2).", "06_CAPACIDAD_PERSONAL"),
        ("Necesidad máxima = PPT diap. 7 (125)", "=K_NEC_MAX", O.OF["S7_REF_MAX"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "", "06_CAPACIDAD_PERSONAL"),
        ("GAP mínimo = PPT diap. 7 (+34)", "=K_GAP_MIN", O.OF["S7_GAP"][0], 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "+34 viene de otra base de cobertura (17 D07).", "06_CAPACIDAD_PERSONAL"),
        ("Cobertura actual = 80 %", "=K_COB_ACTUAL", O.OF["COB_ACTUAL"][0], 'IF(ABS(C-D)<=0.01,"OK","REVISAR")', "Plantilla ÷ subtotal sin cobertura.", "06_CAPACIDAD_PERSONAL"),
        ("Nuevos productos cuadran con la referencia aprobada (75)", "=K_MEB_TOTAL", 75, 'IF(C=D,"OK","ERROR")', "02_MEB_ECO_Oficial.", "09_MEB_ECO"),
        ("Críticos MEB identificados en el inventario = 75", "=K_MEB_IDENT", 75, 'IF(C=D,"OK","REVISAR")', "Aluminio GGH/MGH sin inventario; ver 09 §2.", "09_MEB_ECO"),
        ("Suma por grupos (técnico) = total técnico", "=SUMIFS(S_FTE,S_GRUPO,\"<>\")", "=K_TEC_FTE", 'IF(ABS(C-D)<=P_TOL_EXACT,"OK","ERROR")', "Todo activo del parque tiene grupo.", "05_FTE_POR_GRUPOS"),
        ("FTE por grupos presentado (diap. 7) = FTE base", "=SUM('05_FTE_POR_GRUPOS'!B7:B14)", "=K_FTE_BASE", 'IF(ABS(C-D)<=P_TOL_RED,"OK","REVISAR")', "Reparto presentado incompleto: no se fuerza el cuadre.", "05_FTE_POR_GRUPOS"),
        ("No hay IDs duplicados en el maestro", "=ROWS(M_ID)-SUMPRODUCT(1/COUNTIF(M_ID,M_ID))", 0, 'IF(ROUND(C,6)=0,"OK","ERROR")', "Cada registro una sola vez.", "02_MAESTRO_CRITICIDAD"),
        ("No hay equipos SAP duplicados", '=SUMPRODUCT((M_EQUIPO<>"")*(COUNTIF(M_EQUIPO,M_EQUIPO)>1))', 0, 'IF(C=0,"OK","ERROR")', "", "02_MAESTRO_CRITICIDAD"),
        ("Bajas, excluidos y pendientes no suman FTE", '=SUMIFS(M_FTE,M_PARQUE,"N")', 0, 'IF(C=0,"OK","ERROR")', "Solo cuenta el parque principal.", "02_MAESTRO_CRITICIDAD"),
        ("Componentes y máquinas principales no se confunden", '=COUNTIFS(M_ESTADO,"Componente*",M_PARQUE,"S")', 0, 'IF(C=0,"OK","ERROR")', "Componentes fuera del recuento.", "02_MAESTRO_CRITICIDAD"),
        ("Ningún activo sin información suficiente clasificado como D", '=COUNTIFS(S_SUF,"N",S_CLASE,"D")', 0, 'IF(C=0,"OK","ERROR")', "Regla: la falta de datos nunca envía a D.", "03_SCORE_CRITICIDAD"),
        ("Todos los cuellos de botella del parque son A", '=COUNTIFS(S_CB,"S",S_CLASE,"<>A")', 0, 'IF(C=0,"OK","ERROR")', "", "03_SCORE_CRITICIDAD"),
        ("Todo el montaje principal es A", '=COUNTIFS(S_MONT,"S",S_CLASE,"<>A")', 0, 'IF(C=0,"OK","ERROR")', "", "03_SCORE_CRITICIDAD"),
        ("Todos los activos del parque tienen score", '=COUNT(S_SCORE)', "=ROWS(S_ID)", 'IF(C=D,"OK","ERROR")', "", "03_SCORE_CRITICIDAD"),
        ("Fórmulas sin error en score y clase", '=SUMPRODUCT(--ISERROR(S_SCORE))+SUMPRODUCT(--ISERROR(S_FTE))', 0, 'IF(C=0,"OK","ERROR")', "", "03_SCORE_CRITICIDAD"),
        ("Clases técnicas suman el parque técnico", '=COUNTIF(S_CLASE,"A")+COUNTIF(S_CLASE,"B")+COUNTIF(S_CLASE,"C")+COUNTIF(S_CLASE,"D")', "=ROWS(S_ID)", 'IF(C=D,"OK","ERROR")', "", "04_FTE_CRITICIDAD"),
        ("Parque fila a fila = parque oficial", "=K_TEC_ACTIVOS", "=P_PARQUE_OF", 'IF(C=D,"OK","REVISAR")', "Diferencia explicada en 17 D05.", "17_DIFERENCIAS"),
        ("Horas de avería asignadas + sin activo = total", '=SUMIF(AV_GRUPO,"<>SIN ACTIVO",AV_MIN)/60+SUMIF(AV_GRUPO,"SIN ACTIVO",AV_MIN)/60', "=K_AV_H", 'IF(ABS(C-D)<=0.1,"OK","ERROR")', "", "10_AVERIAS"),
        ("Regla jerárquica = depuración SAP previa", int((c.inv["CONTROL_JERARQUIA"] == "DIFERENTE").sum()), 0, 'IF(C=0,"OK","ERROR")', "Registros con grupo distinto al de LISTADO MAQUINAS SAP (calculado al generar).", "13_METODOLOGIA"),
        ("Todas las fuentes informadas", '=COUNTIF(T_DISP,"NO DISPONIBLE*")', 0, 'IF(C=0,"OK","REVISAR")', "Fuentes que faltan en el repositorio (14).", "14_FUENTES"),
        ("KPIs del PPT reproducidos sin diferencia (11)", '=COUNTIF(T_ESTADOS,"DIFERENCIA*")+COUNTIF(T_ESTADOS,"TRAZABILIDAD*")', 0, 'IF(C=0,"OK","REVISAR")', "Diferencias y trazabilidad incompleta documentadas en 11 y 17.", "11_TRAZABILIDAD"),
        ("Pendientes de validación abiertos (15)", '=COUNTIF(T_PEND,"PENDIENTE")', 0, 'IF(C=0,"OK","REVISAR")', "GAP orientativo mientras haya pendientes de impacto alto.", "15_PENDIENTES_VALIDACION"),
    ]
    for i, (lbl, res, esp, fx, crit, ira) in enumerate(C):
        r = 11 + i
        put(ws, f"A{r}", f"C{i + 1:02d}", F_BOLD, align=CENTER)
        put(ws, f"B{r}", lbl, align=WRAP)
        is_pct = "80 %" in lbl
        put(ws, f"C{r}", res, F_BOLD, fmt="pct" if is_pct else "#,##0.00")
        put(ws, f"D{r}", esp, fill_="oficial" if not str(esp).startswith("=") else None, fmt="pct" if is_pct else "#,##0.00")
        if esp == "VERDADERO":
            ws[f"D{r}"].value = True
        put(ws, f"E{r}", "=" + _fx(fx, r), F_BOLD)
        put(ws, f"F{r}", crit, align=WRAP)
        put(ws, f"G{r}", hl(f"#'{ira}'!A1", "→ " + ira), F_LINK)
    status_format(ws, f"E11:E{10 + len(C)}")
    ws.freeze_panes = "C11"


def _fx(fx: str, r: int) -> str:
    """Sustituye los marcadores C y D (resultado / esperado) por celdas de la fila r."""
    import re as _re
    return _re.sub(r"\b([CD])\b", lambda m: f"{m.group(1)}{r}", fx)


# ======================================================================================
# 01_RESUMEN_EJECUTIVO
# ======================================================================================
def hoja_resumen(c: Ctx):
    ws = c.ws["01_RESUMEN_EJECUTIVO"]
    title(ws, "01 · RESUMEN EJECUTIVO – Respaldo técnico del PowerPoint final",
          f"Resultado oficial: {O.PPT}. Cada cifra se calcula con fórmulas sobre 12_SUPUESTOS y se compara con la diapositiva.")
    widths(ws, {"A": 3, "B": 30, "C": 15, "D": 15, "E": 13, "F": 12, "G": 26, "H": 3, "I": 30, "J": 15, "K": 15, "L": 15})
    cards = [
        ("B", "Parque oficial", "=P_PARQUE_OF", "2.043", "#,##0", "diap. 4 · 8"),
        ("D", "Activos críticos (A)", "=P_N_A", "418", "#,##0", "diap. 4 · 8"),
        ("F", "FTE base criticidad", "=K_FTE_BASE", "85,37", "0.00", "diap. 4"),
        ("I", "Necesidad total (FTE)", '=TEXT(K_NEC_MIN,"0")&" – "&TEXT(K_NEC_MAX,"0")', "115 – 125", None, "diap. 7"),
        ("K", "Plantilla nominal", "=P_PLANTILLA", "~80", "#,##0", "diap. 5 · 7"),
    ]
    for col, lbl, f, ppt, fmt, diap in cards:
        c0 = ws[f"{col}5"]
        c0.value = lbl
        c0.font = F_SUB
        cv = ws[f"{col}6"]
        cv.value = f
        cv.font = F_BIG
        if fmt:
            cv.number_format = fmt
        cp = ws[f"{col}7"]
        cp.value = f"PPT: {ppt} ({diap})"
        cp.font = F_NOTE
        for rr in (5, 6, 7):
            ws[f"{col}{rr}"].fill = FILL["card"]
    ws.row_dimensions[6].height = 30
    put(ws, "B9", '="GAP orientativo: "&TEXT(K_GAP_MIN,"+0.0")&" a "&TEXT(K_GAP_MAX,"+0.0")&" FTE (PPT: +34 diap. 7 · +33 diap. 5) · Cobertura actual "&TEXT(K_COB_ACTUAL,"0%")',
        F_SECTION, border=False)

    section(ws, 11, "Waterfall de la necesidad (06_CAPACIDAD_PERSONAL)", 7, 2)
    header_row(ws, 12, 2, ["Capa", "Modelo (exacto)", "Redondeado", "PPT", "Diap.", "Estado"])
    capas = [("A · Base por criticidad", "=K_FTE_BASE", 85, "4 · 5"), ("B · Preventivo + avisos", "=K_PM", 10, "5"), ("C · Formación", "=K_FORM", 3, "5"),
             ("D · Sábados / libranzas", "=K_SAB", 2, "5"), ("E · Cobertura mín.", "=K_COB_MIN", 13, "5"), ("E · Cobertura máx.", "=K_COB_MAX", 21, "5"),
             ("F · Necesidad mín.", "=K_NEC_MIN", 115, "7"), ("F · Necesidad máx.", "=K_NEC_MAX", 125, "7"), ("G · Plantilla", "=K_PLANT_FTE", 80, "5 · 7"),
             ("H · GAP mín.", "=K_GAP_MIN", 34, "7"), ("H · GAP máx.", "=K_GAP_MAX", None, "—")]
    for i, (lbl, f, ppt, d) in enumerate(capas):
        r = 13 + i
        put(ws, f"B{r}", lbl, F_BOLD if lbl.startswith(("F", "H")) else F_BASE)
        put(ws, f"C{r}", f, F_BOLD, fmt="0.00")
        put(ws, f"D{r}", f"=ROUND(C{r},0)", fmt="0")
        put(ws, f"E{r}", ppt, fill_="oficial" if ppt is not None else "pendiente", fmt="0")
        put(ws, f"F{r}", d, align=CENTER)
        put(ws, f"G{r}", f'=IF(E{r}="","NO PRESENTADO",IF(ABS(C{r}-E{r})<=P_TOL_RED,"OK","REVISAR – ver 17"))', F_BOLD)
    status_format(ws, "G13:G23")

    section(ws, 11, "Distribución A/B/C/D oficial (04_FTE_CRITICIDAD)", 4, 9)
    header_row(ws, 12, 9, ["Clase", "Activos", "Ratio", "FTE"])
    for i, k in enumerate("ABCD"):
        r = 13 + i
        put(ws, f"I{r}", k, F_BOLD, k)
        if k == "A":
            ws[f"I{r}"].font = F_HEAD
        put(ws, f"J{r}", f"='04_FTE_CRITICIDAD'!B{7 + i}", fmt="int")
        put(ws, f"K{r}", f"='04_FTE_CRITICIDAD'!C{7 + i}", fmt="0")
        put(ws, f"L{r}", f"='04_FTE_CRITICIDAD'!D{7 + i}", fmt="0.00")
    put(ws, "I17", "TOTAL", F_BOLD, "total")
    put(ws, "J17", "=SUM(J13:J16)", F_BOLD, "total", "int")
    put(ws, "K17", "", fill_="total")
    put(ws, "L17", "=SUM(L13:L16)", F_BOLD, "total", "0.00")

    ch = BarChart()
    ch.type = "col"
    ch.title = "Activos por clase (oficial)"
    ch.y_axis.title = "Activos"
    ch.add_data(Reference(ws, min_col=10, min_row=12, max_row=16), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=9, min_row=13, max_row=16))
    ch.height, ch.width = 7, 12
    ch.legend = None
    ch.dataLabels = DataLabelList()
    ch.dataLabels.showVal = True
    from openpyxl.chart.series import DataPoint
    s = ch.series[0]
    for idx, colr in enumerate(["C00000", "F4B084", "FFE699", "A9D08E"]):
        pt = DataPoint(idx=idx)
        pt.graphicalProperties.solidFill = colr
        s.dPt.append(pt)
    ws.add_chart(ch, "I19")

    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "Capas de la necesidad (FTE, modelo)"
    ch2.add_data(Reference(ws, min_col=3, min_row=12, max_row=19), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=2, min_row=13, max_row=19))
    ch2.height, ch2.width = 8, 14
    ch2.legend = None
    ch2.dataLabels = DataLabelList()
    ch2.dataLabels.showVal = True
    ch2.series[0].graphicalProperties.solidFill = NAVY
    ws.add_chart(ch2, "B26")

    r = 43
    section(ws, r, "Mensajes clave (calculados)", 11, 2)
    msgs = [
        '="1. El parque oficial de "&TEXT(P_PARQUE_OF,"#,##0")&" activos exige "&TEXT(K_FTE_BASE,"0.0")&" FTE base solo por criticidad (A "&P_N_A&" · B "&P_N_B&" · C "&P_N_C&" · D "&P_N_D&")."',
        '="2. Con carga planificada, formación, libranzas y cobertura "&TEXT(P_COB_MIN,"0%")&"–"&TEXT(P_COB_MAX,"0%")&", la necesidad es de "&TEXT(K_NEC_MIN,"0")&"–"&TEXT(K_NEC_MAX,"0")&" FTE."',
        '="3. Frente a ~"&P_PLANTILLA&" personas, el GAP orientativo es de "&TEXT(K_GAP_MIN,"0")&"–"&TEXT(K_GAP_MAX,"0")&" FTE; la plantilla cubre el "&TEXT(K_COB_ACTUAL,"0%")&" de la necesidad antes de cobertura organizativa."',
        '="4. La clase A (418) y los ratios son trazables; el reparto B/C/D y el reparto por grupos MEB son declarados: TRAZABILIDAD INCOMPLETA (ver 17)."',
        '="5. El inventario fila a fila suma "&TEXT(K_TEC_ACTIVOS,"#,##0")&" activos ("&TEXT(K_TEC_PROV,"#,##0")&" provisionales) y "&TEXT(K_TEC_FTE,"0.0")&" FTE base: contraste técnico, no resultado oficial."',
        '="6. Controles: "&K_CTRL_OK&" OK · "&K_CTRL_REV&" REVISAR · "&K_CTRL_ERR&" ERROR · "&COUNTIF(T_DIF,"ABIERTA")&" diferencias abiertas · "&COUNTIFS(T_PEND,"PENDIENTE",T_PEND_IMP,"Alta")&" pendientes de impacto alto."',
    ]
    for i, m in enumerate(msgs):
        rr = r + 1 + i
        cl = ws.cell(rr, 2, m)
        cl.font = F_BASE
        cl.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=12)
        ws.row_dimensions[rr].height = 30


# ======================================================================================
# 00_INICIO
# ======================================================================================
def hoja_inicio(c: Ctx):
    ws = c.ws["00_INICIO"]
    title(ws, "PK-16 · MAESTRO DE CRITICIDAD Y DIMENSIONAMIENTO DE MANTENIMIENTO",
          "Respaldo técnico, trazable y auditable de la presentación final a Dirección.", home=False)
    widths(ws, {"A": 3, "B": 30, "C": 64, "D": 18, "E": 3, "F": 30, "G": 40})
    put(ws, "B4", "POWERPOINT FINAL = RESULTADO OFICIAL PRESENTADO", F_SECTION, "oficial", border=False)
    ws.merge_cells("B4:D4")
    put(ws, "B5", "EXCEL MAESTRO = DATOS + CÁLCULOS + CRITERIOS + METODOLOGÍA + TRAZABILIDAD", F_SECTION, "dato", border=False)
    ws.merge_cells("B5:D5")
    put(ws, "B6", "Inventario → Criticidad → Carga → Disponibilidad real → Necesidad → GAP", F_BOLD, border=False)
    put(ws, "B7", "Trazabilidad: POWERPOINT → RESULTADO OFICIAL → CÁLCULO → GRUPO → ACTIVO → CRITERIOS → DATO ORIGINAL → ARCHIVO FUENTE", F_BASE, border=False)
    header_row(ws, 9, 2, ["Hoja", "Contenido", "Tipo"])
    tipo = {"00": "Navegación", "01": "Resultado", "02": "Datos", "03": "Cálculo", "04": "Cálculo", "05": "Cálculo", "06": "Cálculo", "07": "Cálculo",
            "08": "Cálculo", "09": "Cálculo", "10": "Datos", "11": "Trazabilidad", "12": "Supuestos", "13": "Metodología", "14": "Fuentes",
            "15": "Pendientes", "16": "Controles", "17": "Diferencias", "18": "Datos", "19": "Datos", "20": "Datos"}
    fill_tipo = {"Datos": "dato", "Supuestos": "supuesto", "Pendientes": "pendiente", "Resultado": "oficial", "Diferencias": "pendiente"}
    for i, (name, desc) in enumerate(SHEETS):
        r = 10 + i
        put(ws, f"B{r}", hl(f"#'{name}'!A1", name), F_LINK)
        put(ws, f"C{r}", desc, align=WRAP)
        t = tipo[name[:2]]
        put(ws, f"D{r}", t, fill_=fill_tipo.get(t))
    put(ws, "F9", "Leyenda de colores", F_HEAD, "head")
    put(ws, "G9", "", fill_="head")
    leg = [("oficial", "Valor oficial presentado en el PPT"), ("dato", "Dato importado de una fuente"), ("supuesto", "Supuesto / parámetro editable (12)"),
           ("pendiente", "Pendiente / no disponible / excluido"), ("A", "Criticidad A"), ("B", "Criticidad B"), ("C", "Criticidad C"), ("D", "Criticidad D"),
           ("OK", "OK"), ("REVISAR", "REVISAR / provisional"), ("ERROR", "ERROR / diferencia")]
    for i, (f, t) in enumerate(leg):
        put(ws, f"F{10 + i}", "", fill_=f)
        put(ws, f"G{10 + i}", t)
    put(ws, "F23", "Cómo usar el libro", F_HEAD, "head")
    put(ws, "G23", "", fill_="head")
    uso = ["1. Cambie parámetros solo en 12_SUPUESTOS (celdas ámbar).",
           "2. Revise 16_CONTROLES y 17_DIFERENCIAS tras cada cambio.",
           "3. 11_TRAZABILIDAD explica cada cifra del PPT.",
           "4. 02 y 03 permiten filtrar cualquier activo y ver su motivo.",
           "5. El libro se regenera con modelo/build.py desde las fuentes."]
    for i, t in enumerate(uso):
        put(ws, f"F{24 + i}", t, border=False)
        ws.merge_cells(f"F{24 + i}:G{24 + i}")
    put(ws, "F30", "Estado del estudio", F_HEAD, "head")
    put(ws, "G30", "", fill_="head")
    est = [("Controles OK", "=K_CTRL_OK"), ("Controles REVISAR", "=K_CTRL_REV"), ("Controles ERROR", "=K_CTRL_ERR"),
           ("Diferencias abiertas", '=COUNTIF(T_DIF,"ABIERTA")'), ("Pendientes de validación", '=COUNTIF(T_PEND,"PENDIENTE")'),
           ("Generado", dt.datetime.now().strftime("%d/%m/%Y %H:%M"))]
    for i, (a, b) in enumerate(est):
        put(ws, f"F{31 + i}", a)
        put(ws, f"G{31 + i}", b, F_BOLD)


# ======================================================================================
def construir_excel(M, S, path):
    c = Ctx(M, S)
    hoja_supuestos(c)
    hoja_score(c)
    hoja_maestro(c)
    hoja_fte_criticidad(c)
    hoja_fte_grupos(c)
    hoja_sabados_formacion(c)
    hoja_pm(c)
    hoja_capacidad(c)
    hoja_meb(c)
    hoja_reg_averias(c)
    hoja_cruce_montaje(c)
    hoja_cruce_cuellos(c)
    hoja_averias(c)
    hoja_trazabilidad(c)
    hoja_diferencias(c)
    hoja_metodologia(c)
    hoja_fuentes(c)
    hoja_pendientes(c)
    hoja_controles(c)
    hoja_resumen(c)
    hoja_inicio(c)
    tabs = {"00": NAVY, "01": "7030A0", "12": "FFC000", "15": "A6A6A6", "16": "00B050", "17": "C00000"}
    for name, ws in c.ws.items():
        ws.sheet_properties.tabColor = tabs.get(name[:2], "8EA9DB")
        for row in ws.iter_rows():
            for cell in row:
                if cell.font is None or cell.font.name != "Arial":
                    f = cell.font
                    cell.font = f.copy(name="Arial") if f else F_BASE
    c.wb.active = 0
    c.wb.calculation.fullCalcOnLoad = True  # Excel recalcula al abrir
    c.wb.save(path)
    return c
