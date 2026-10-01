# -*- coding: utf-8 -*-
"""Estilos y utilidades de escritura para el Excel maestro PK-16."""
from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

FONT = "Arial"
NAVY = "1F3864"

F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
F_SECTION = Font(name=FONT, size=12, bold=True, color=NAVY)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_BIG = Font(name=FONT, size=18, bold=True, color=NAVY)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


FILL = {
    "head": fill(NAVY),
    "section": fill("D9E1F2"),
    "dato": fill("DDEBF7"),       # dato importado de fuente (azul)
    "supuesto": fill("FFF2CC"),   # supuesto editable (ámbar)
    "pendiente": fill("E7E6E6"),  # pendiente / no disponible (gris)
    "oficial": fill("E4DFEC"),    # valor oficial presentado en PPT (lila)
    "formula": fill("FFFFFF"),
    "total": fill("F2F2F2"),
    "card": fill("F3F6FB"),
    "A": fill("C00000"), "B": fill("F4B084"), "C": fill("FFE699"), "D": fill("A9D08E"),
    "OK": fill("C6EFCE"), "REVISAR": fill("FFEB9C"), "ERROR": fill("FFC7CE"),
}
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

NUM = {"int": "#,##0", "fte": "0.00", "fte1": "0.0", "pct": "0.0%", "h": "#,##0.0", "dec": "0.00", "x": "0.00"}


def put(ws, ref, value, font=None, fill_=None, fmt=None, align=None, border=True):
    c = ws[ref]
    c.value = value
    if isinstance(value, str) and (value.startswith("= ") or value.startswith("=\u03a3")):
        c.data_type = "s"  # texto explicativo de fórmula, no fórmula
    c.font = font or F_BASE
    if fill_ is not None:
        c.fill = FILL[fill_] if isinstance(fill_, str) else fill_
    if fmt:
        c.number_format = NUM.get(fmt, fmt)
    if align is not None:
        c.alignment = align
    if border:
        c.border = BORDER
    return c


def hl(href: str, text: str) -> str:
    """Hipervínculo interno como fórmula HYPERLINK (sobrevive al recálculo y al renombrado del fichero)."""
    t = str(text).replace('"', "'")
    return f'=HYPERLINK("{href}","{t}")'


def header_row(ws, row, col, headers, widths=None, height=30):
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font = F_HEAD
        c.fill = FILL["head"]
        c.alignment = CENTER
        c.border = BORDER
    ws.row_dimensions[row].height = height
    if widths:
        for j, w in enumerate(widths):
            if w:
                ws.column_dimensions[get_column_letter(col + j)].width = w


def title(ws, text, sub=None, home=True):
    ws["A1"].value = text
    ws["A1"].font = F_TITLE
    ws.row_dimensions[1].height = 26
    if sub:
        ws["A2"].value = sub
        ws["A2"].font = F_SUB
    if home:
        c = ws["A3"]
        c.value = hl("#'00_INICIO'!A1", "⌂ Volver a 00_INICIO")
        c.font = F_LINK
    ws.sheet_view.showGridLines = False


def section(ws, row, text, ncols=8, col=1):
    for j in range(ncols):
        c = ws.cell(row=row, column=col + j)
        c.fill = FILL["section"]
    c = ws.cell(row=row, column=col, value=text)
    c.font = F_SECTION
    ws.row_dimensions[row].height = 20


def note(ws, ref, text):
    c = ws[ref]
    c.value = text
    c.font = F_NOTE
    c.alignment = Alignment(wrap_text=False, vertical="top")


def define(wb, name, sheet, ref):
    dn = DefinedName(name, attr_text=f"'{sheet}'!{ref}")
    wb.defined_names[name] = dn


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


STATUS_RULES = [
    ("OK", "C6EFCE", "006100"),
    ("REVISAR", "FFEB9C", "9C5700"),
    ("ERROR", "FFC7CE", "9C0006"),
    ("DIFERENCIA", "FFC7CE", "9C0006"),
    ("TRAZABILIDAD INCOMPLETA", "FCE4D6", "833C0B"),
    ("PENDIENTE", "E7E6E6", "404040"),
    ("PROVISIONAL", "FFF2CC", "7F6000"),
    ("NO DISPONIBLE", "E7E6E6", "404040"),
    ("EXCLUIDO", "F2F2F2", "7F7F7F"),
    ("ALERTA", "FFEB9C", "9C5700"),
    ("RECONSTRUIDO", "DDEBF7", "1F3864"),
]


def status_format(ws, rng):
    """Formato condicional por texto de estado (OK / REVISAR / ERROR / PENDIENTE / ...)."""
    first = rng.split(":")[0].replace("$", "")
    for txt, bg, fg in STATUS_RULES:
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'ISNUMBER(SEARCH("{txt}",{first}))'], fill=fill(bg),
                             font=Font(name=FONT, size=10, bold=True, color=fg), stopIfTrue=True))


def class_format(ws, rng):
    """A rojo, B naranja, C amarillo, D verde (texto que empieza por la letra de la clase)."""
    first = rng.split(":")[0].replace("$", "")
    for k, fg in (("A", "FFFFFF"), ("B", "000000"), ("C", "000000"), ("D", "000000")):
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'LEFT({first},1)="{k}"'], fill=FILL[k],
                             font=Font(name=FONT, size=10, bold=True, color=fg), stopIfTrue=True))
